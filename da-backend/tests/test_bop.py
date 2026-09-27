"""BOP port checks that need no network.

The emitter is exercised against a blank python-docx document rather than the real SOP
template, because the named styles it relies on ("Heading 1", "Normal",
"List Paragraph") exist in the default template too. That keeps these fast and
runnable without S3 credentials.
"""

from __future__ import annotations

import io

import pytest

from da_platform.settings import REPO_ROOT
from da_platform.silo_registry import _load_module

BOP_DIR = REPO_ROOT / "silos" / "bop"

pytestmark = pytest.mark.skipif(
    not (BOP_DIR / "silo.py").is_file(), reason="the BOP silo is not present"
)


@pytest.fixture(scope="module")
def bop():
    """Load the real BOP silo, registering its package so submodules import."""
    module = _load_module("bop", BOP_DIR / "silo.py")
    from da_silos.bop import editing, emit, generate, schema

    return {
        "silo": module,
        "editing": editing,
        "emit": emit,
        "generate": generate,
        "schema": schema,
    }


def blank_template() -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    Document().save(buffer)
    return buffer.getvalue()


# ── contract ──────────────────────────────────────────────────────────────────


def test_silo_declares_the_expected_contract(bop):
    silo = bop["silo"]
    assert silo.LABEL
    assert silo.ACCEPTS == [".pdf", ".docx"]
    assert silo.STAGES == ["extract", "generate", "review", "await_review", "build"]
    assert silo.STORAGE_PREFIX == "BOP"
    for stage in silo.STAGES:
        assert callable(getattr(silo, stage)), stage


# ── schema ────────────────────────────────────────────────────────────────────


def test_dotted_set_preserves_list_slots(bop):
    """Matches the source: writing a scalar into a list slot wraps it in a list."""
    schema = bop["schema"]
    document = schema.empty_document()
    schema.dotted_set(document, "description.tools", "a single tool")
    assert document["description"]["tools"] == ["a single tool"]

    schema.dotted_set(document, "safety.hazards", "<p>Pressurised</p>")
    assert document["safety"]["hazards"] == "<p>Pressurised</p>"


def test_dotted_get_returns_the_default_when_absent(bop):
    schema = bop["schema"]
    assert schema.dotted_get({}, "a.b.c", "fallback") == "fallback"


def test_empty_document_is_a_fresh_copy(bop):
    schema = bop["schema"]
    first = schema.empty_document()
    first["description"]["tools"].append("mutated")
    assert schema.empty_document()["description"]["tools"] == []


# ── response parsing ──────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "raw,expected",
    [
        ('{"purpose": "x"}', {"purpose": "x"}),
        ('```json\n{"purpose": "x"}\n```', {"purpose": "x"}),
        ('```\n{"purpose": "x"}\n```', {"purpose": "x"}),
        ('Here you go: {"purpose": "x"} — hope that helps', {"purpose": "x"}),
    ],
)
def test_json_parsing_tolerates_fences_and_chatter(bop, raw, expected):
    assert bop["generate"].parse_json_response(raw) == expected


def test_json_parsing_rejects_empty(bop):
    with pytest.raises(ValueError):
        bop["generate"].parse_json_response("   ")


@pytest.mark.parametrize(
    "text,truncated",
    [('{"a": 1}', False), ("[1, 2]", False), ('{"a": "unterminat', True), ("", True)],
)
def test_truncation_heuristic(bop, text, truncated):
    assert bop["generate"].looks_truncated(text) is truncated


@pytest.mark.parametrize(
    "value,expected",
    [
        ({"equipment": ["a"], "software": ["b"]}, {"equipment": ["a"], "software": ["b"]}),
        (["a", "b"], {"equipment": ["a", "b"], "software": []}),
        ("one step", {"equipment": ["one step"], "software": []}),
        (None, {"equipment": [], "software": []}),
        ({"equipment": "a"}, {"equipment": ["a"], "software": []}),
    ],
)
def test_op_phase_normalisation(bop, value, expected):
    assert bop["generate"].normalize_op_phase("setup", value) == expected


def test_troubleshooting_normalises_to_a_list(bop):
    normalize = bop["generate"].normalize_op_phase
    assert normalize("troubleshooting", [{"issue": "i", "resolution": "r"}]) == [
        {"issue": "i", "resolution": "r"}
    ]
    assert normalize("troubleshooting", "not a list") == []


# ── editing round trip ────────────────────────────────────────────────────────


def test_flatten_exposes_prose_and_list_leaves(bop):
    editing, schema = bop["editing"], bop["schema"]
    document = schema.empty_document()
    document["purpose"] = "<p>Purpose</p>"
    document["description"]["tools"] = ["Wrench", "Driver"]

    flat = editing.flatten(document)
    assert flat["purpose"] == "<p>Purpose</p>"
    assert flat["description.tools"] == "<ul><li>Wrench</li><li>Driver</li></ul>"
    # Every paired phase contributes equipment and software slots.
    assert "operating_procedure.setup.equipment" in flat
    assert "operating_procedure.setup.software" in flat


def test_edits_preserve_list_types(bop):
    """The source writes the edited HTML string straight into the slot, which turns a
    list into a string and makes the emitter drop the section. Types survive here."""
    editing, schema = bop["editing"], bop["schema"]
    document = schema.empty_document()
    document["description"]["tools"] = ["Wrench"]

    result = editing.apply_edits(
        document, [("description.tools", "<ul><li>Wrench</li><li>Driver</li></ul>", 2)]
    )
    assert result["description"]["tools"] == ["Wrench", "Driver"]


def test_untouched_sections_are_not_written_back(bop):
    """revision == 1 means "exactly as generated", so it must be left alone."""
    editing, schema = bop["editing"], bop["schema"]
    document = schema.empty_document()
    document["description"]["tools"] = ["Wrench"]
    flat = editing.flatten(document)

    result = editing.apply_edits(document, [("description.tools", flat["description.tools"], 1)])
    assert result["description"]["tools"] == ["Wrench"]


def test_prose_edits_are_stored_as_html(bop):
    editing, schema = bop["editing"], bop["schema"]
    document = schema.empty_document()
    result = editing.apply_edits(document, [("safety.ppe", "<p>Gloves</p>", 3)])
    assert result["safety"]["ppe"] == "<p>Gloves</p>"


def test_html_to_list_handles_plain_lines(bop):
    """A user who deletes the markup should not lose their entries."""
    assert bop["editing"].html_to_list("first\nsecond") == ["first", "second"]


# ── emitter ───────────────────────────────────────────────────────────────────


def test_emitter_produces_every_section_in_order(bop):
    from docx import Document

    emit, schema = bop["emit"], bop["schema"]
    document = schema.empty_document()
    document["purpose"] = "<p>Lead</p><ul><li>Capability</li></ul>"
    document["abbreviations"] = [["SFE", "Supercritical Fluid Extraction"]]
    document["operating_procedure"]["troubleshooting"] = [
        {"issue": "No pressure", "resolution": "Check the seal"}
    ]

    result = emit.build_docx(document, blank_template(), source_name="Helix_Manual.pdf")
    assert result[:4] == b"PK\x03\x04"

    parsed = Document(io.BytesIO(result))
    headings = [
        paragraph.text
        for paragraph in parsed.paragraphs
        if paragraph.style.name.startswith("Heading 1")
    ]
    assert headings == [
        "PURPOSE",
        "SCOPE",
        "DESCRIPTION",
        "SAFETY",
        "OPERATING PROCEDURE",
        "ABBREVIATIONS AND DEFINITIONS",
        "RELATED DOCUMENTS",
    ]
    # Troubleshooting and abbreviations each render as a table.
    assert len(parsed.tables) == 2


def test_emitter_renders_html_lists_as_prefixed_runs(bop):
    from docx import Document

    emit, schema = bop["emit"], bop["schema"]
    document = schema.empty_document()
    document["purpose"] = "<ul><li>First</li><li>Second</li></ul>"

    parsed = Document(io.BytesIO(emit.build_docx(document, blank_template())))
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "• First" in text
    assert "• Second" in text


def test_emitter_honours_data_list_ordered(bop):
    """Editors that emit every list as <ol> mark bullets with data-list."""
    from docx import Document

    emit, schema = bop["emit"], bop["schema"]
    document = schema.empty_document()
    document["purpose"] = '<ol><li data-list="ordered">Step</li></ol>'

    parsed = Document(io.BytesIO(emit.build_docx(document, blank_template())))
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "1. Step" in text


def test_emitter_skips_empty_phase_groups(bop):
    """A phase with no software steps must not emit an empty Software heading."""
    from docx import Document

    emit, schema = bop["emit"], bop["schema"]
    document = schema.empty_document()
    document["operating_procedure"]["setup"] = {
        "equipment": ["Insert the wool"],
        "software": [],
    }

    parsed = Document(io.BytesIO(emit.build_docx(document, blank_template())))
    level_three = [
        paragraph.text
        for paragraph in parsed.paragraphs
        if paragraph.style.name.startswith("Heading 3")
    ]
    assert "Equipment" in level_three
    assert "Software" not in level_three


def test_emitter_tolerates_a_completely_empty_document(bop):
    """Every generation call can fail; the build must still produce a valid file."""
    emit, schema = bop["emit"], bop["schema"]
    result = emit.build_docx(schema.empty_document(), blank_template())
    assert result[:4] == b"PK\x03\x04"
