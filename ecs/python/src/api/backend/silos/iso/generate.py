"""Building the assessment document. Ported from backend.py:5158-5388, plus the output
filename from 5566-5567 which lived in the Flask route rather than in `generate_docx`.

The report is a two-column table — section number, description — preceded by a title
page, a Project Information placeholder and a Purpose/Overview paragraph, and followed by
an end marker. Figures, tables and formulas are inserted inline in the description cell,
with a `keepWithNext` on the text paragraph so Word does not split a caption from its
image across a page break.

**Two different titles, and they must not be swapped.** The title printed *inside* the
document comes from the PDF's own metadata, falling back to the file's basename. The
RAG-derived title — the one the identification step produced — names the output *file*
only. They are separate values in the source, and conflating them is exactly the kind of
silent output drift this port exists to avoid.

Seam changes only: `pdf_path` now comes from the workspace, and the `textract`/`llm`
capabilities are passed through to the pre-scan.
"""

from __future__ import annotations

import io
import logging
import os
import re

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .docxstyle import (
    _COL_W,
    _FONT_SIZE,
    _FONT_SMALL,
    _HDR_SHADE,
    _SEC_SHADE,
    _add_image_to_cell,
    _apply_fixed_cols,
    _cell_border,
    _cell_shade,
    _cell_width,
    _row_min_height,
)
from .patterns import _xml_safe
from .prescan import prescan_media
from .rows import _write_text_highlighted, extract_all_rows
from .sections import _section_num

logger = logging.getLogger(__name__)

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def generate_docx(ws, toc, start_idx, end_idx, llm_pages=None, textract=None, llm=None):
    """Build and return the assessment DOCX as bytes.

    llm_pages: optional ordered list of page dicts from the vision extraction
               (vision path for garbled PDFs).
    """
    pdf_path = ws.pdf_path

    # Metadata
    with fitz.open(pdf_path) as pdf_doc:
        meta  = pdf_doc.metadata
        # NOTE: this is the title printed INSIDE the document. It is deliberately not
        # the RAG-derived title, which names the output file instead.
        doc_title = (meta.get('title', '') or '').strip() \
                    or os.path.basename(pdf_path).replace('.pdf', '')

    # Build llm_texts lookup from llm_pages list (vision path)
    llm_texts = None
    if llm_pages:
        llm_texts = {
            pd["page_idx"]: pd["text"]
            for pd in llm_pages
            if pd and pd.get("text")
        }

    # Pre-scan document for figures / tables / formulas
    # When a workspace is provided, images are persisted to the run's media
    # directory and the returned dicts hold absolute file paths; otherwise they
    # hold raw PNG bytes (in-memory, used for tests).
    # *toc* is passed so the unnumbered-formula detector can (a) restrict
    # detection to the selected page range and (b) map each formula to its
    # section for anchor-based DOCX insertion.
    prescan_result = prescan_media(
        pdf_path, llm_pages=llm_pages, ws=ws,
        toc=toc, toc_start_idx=start_idx, toc_end_idx=end_idx,
        textract=textract, llm=llm,
    )
    # Only lengths 3 and 8 are reachable against the current prescan_media; the
    # intermediate branches are dead. Ported as-is rather than trimmed, but note the
    # hazard: this ladder would silently absorb an arity mistake as an empty
    # media_first_page and a subtly wrong document, instead of raising.
    suppress_texts_by_page = {}
    if len(prescan_result) == 9:
        (figures, tables, formulas,
         unnumbered_formulas, table_rects_by_page, fig_rects_by_page,
         unlabeled_tables_by_page, media_first_page,
         suppress_texts_by_page) = prescan_result
    elif len(prescan_result) == 8:
        (figures, tables, formulas,
         unnumbered_formulas, table_rects_by_page, fig_rects_by_page,
         unlabeled_tables_by_page, media_first_page) = prescan_result
    elif len(prescan_result) == 7:
        (figures, tables, formulas,
         unnumbered_formulas, table_rects_by_page, unlabeled_tables_by_page,
         media_first_page) = prescan_result
        fig_rects_by_page = {}
    elif len(prescan_result) == 6:
        (figures, tables, formulas,
         unnumbered_formulas, table_rects_by_page, unlabeled_tables_by_page) = prescan_result
        fig_rects_by_page = {}
        media_first_page = {}
    elif len(prescan_result) == 5:
        (figures, tables, formulas,
         unnumbered_formulas, table_rects_by_page) = prescan_result
        fig_rects_by_page = {}
        unlabeled_tables_by_page = {}
        media_first_page = {}
    elif len(prescan_result) == 4:
        figures, tables, formulas, unnumbered_formulas = prescan_result
        table_rects_by_page = {}
        fig_rects_by_page = {}
        unlabeled_tables_by_page = {}
        media_first_page = {}
    else:
        figures, tables, formulas = prescan_result
        unnumbered_formulas = []
        table_rects_by_page = {}
        fig_rects_by_page = {}
        unlabeled_tables_by_page = {}
        media_first_page = {}

    # Extract content rows
    rows = extract_all_rows(pdf_path, toc, start_idx, end_idx,
                            figures, tables, formulas, llm_texts=llm_texts,
                            unnumbered_formulas=unnumbered_formulas,
                            table_rects_by_page=table_rects_by_page,
                            fig_rects_by_page=fig_rects_by_page,
                            unlabeled_tables_by_page=unlabeled_tables_by_page,
                            media_first_page=media_first_page,
                            suppress_texts_by_page=suppress_texts_by_page)

    # ── Build Document ──────────────────────────────────────────────────────
    doc = Document()

    # A4 page
    for section in doc.sections:
        section.page_width  = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin  = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.top_margin   = Cm(2.54)
        section.bottom_margin = Cm(2.54)

    # Title
    title_p = doc.add_paragraph()
    try:
        title_p.style = doc.styles['Title']
    except Exception:
        pass
    run = title_p.add_run(_xml_safe(f'{doc_title}'))
    run.font.size = Pt(18)
    run.font.bold = True

    sub_run = title_p.add_run('\nApplicability Assessment Report')
    sub_run.font.size = Pt(13)
    sub_run.font.bold = False

    doc.add_paragraph()

    # Project Information
    doc.add_heading('Project Information', level=1)
    p = doc.add_paragraph('<< Add a description of the specific project here. >>')
    if p.runs:
        p.runs[0].font.italic = True
        p.runs[0].font.color.rgb = RGBColor(0x00, 0x70, 0xC0)

    doc.add_paragraph()

    # Purpose / Overview
    doc.add_heading('Purpose / Overview', level=1)
    start_num = _section_num(toc[start_idx]['title'])
    end_num   = _section_num(toc[end_idx]['title'])
    doc.add_paragraph(
        f'This document reviews the requirements of {doc_title} to determine the '
        f'applicability of each clause to the subject medical device. Each requirement '
        f'is assessed for compliance. The assessment covers sections '
        f'{start_num} through {end_num}.'
    )

    doc.add_paragraph()

    # Assessment heading
    doc.add_heading(f'Assessment of {doc_title}', level=1)
    note_p = doc.add_paragraph(
        'NOTE: Where differences exist between British English and American English '
        'spellings, the American English spelling has been used. '
        'Blue italic text represents guidance and should be removed prior to routing.'
    )
    if note_p.runs:
        note_p.runs[0].font.italic = True
        note_p.runs[0].font.size   = _FONT_SMALL

    doc.add_paragraph()

    # ── Assessment Table ────────────────────────────────────────────────────
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'

    _apply_fixed_cols(tbl._tbl, _COL_W)

    # Header row
    hdr_cells = tbl.rows[0].cells
    hdr_labels = [
        'Section\nNumber',
        'Description',
    ]
    for i, (cell, label) in enumerate(zip(hdr_cells, hdr_labels)):
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.font.bold = True
        r.font.size = _FONT_SIZE
        _cell_shade(cell, _HDR_SHADE)
        _cell_border(cell)
        _cell_width(cell, _COL_W[i])
    _row_min_height(tbl.rows[0], 500)

    # ── Content rows ─────────────────────────────────────────────────────────
    for row_data in rows:
        row = tbl.add_row()
        cells = row.cells

        for i, w in enumerate(_COL_W):
            _cell_width(cells[i], w)

        if row_data['type'] == 'header':
            cells[0].text = ''
            p0 = cells[0].paragraphs[0]
            r0 = p0.add_run(_xml_safe(row_data['section_num']))
            r0.font.bold = True
            r0.font.size = _FONT_SIZE

            cells[1].text = ''
            p1 = cells[1].paragraphs[0]
            r1 = p1.add_run(_xml_safe(row_data['text']))
            r1.font.bold = True
            r1.font.size = _FONT_SIZE

            _cell_shade(cells[0], _SEC_SHADE)
            _cell_shade(cells[1], _SEC_SHADE)
            _cell_border(cells[0])
            _cell_border(cells[1])
            _row_min_height(row, 340)

        else:
            # Content row
            # Col 0 — Section number
            cells[0].text = ''
            cells[0].paragraphs[0].add_run(_xml_safe(row_data['section_num'])).font.size = _FONT_SIZE

            # Col 1 — Description (text + inline images)
            # References to Figure/Table/Formula are highlighted in yellow.
            cells[1].text = ''
            _write_text_highlighted(cells[1].paragraphs[0], row_data['text'], _FONT_SIZE)
            # When the cell contains an image keep the caption text on the same
            # page as the image — prevents Word splitting them across pages.
            if row_data.get('inline_images'):
                _pPr = cells[1].paragraphs[0]._p.get_or_add_pPr()
                _kwn = OxmlElement('w:keepWithNext')
                _kwn.set(qn('w:val'), '1')
                _pPr.append(_kwn)
            for img in row_data.get('inline_images', []):
                if 'data' in img:
                    _add_image_to_cell(
                        cells[1], img['data'], img.get('caption', ''),
                        alt_text=img.get('latex', ''),
                    )

            for cell in cells:
                _cell_border(cell)
            _row_min_height(row, 280)

    # End marker
    doc.add_paragraph()
    end_p = doc.add_paragraph('— END OF DOCUMENT —')
    end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if end_p.runs:
        end_p.runs[0].font.bold = True
        end_p.runs[0].font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_document(
    ctx,
    ws,
    *,
    toc,
    doc_title,
    start_idx,
    end_idx,
    use_vision,
    llm_pages,
    total_pages,
):
    """The stage-facing wrapper.

    `generate_docx` is the ported function; this adds only the output filename, which
    lived in the Flask `/generate` route (backend.py:5566-5567).

    `doc_title` here is the RAG-derived title — the value the identification step
    produced. It names the FILE and nothing else; the title printed inside the document
    is derived independently by `generate_docx` from the PDF's own metadata. They are
    different values in the source, and passing this one through would be silent drift.
    """
    ctx.progress("Building the assessment", pct=85)

    docx_bytes = generate_docx(
        ws,
        toc,
        start_idx,
        end_idx,
        llm_pages=llm_pages,
        # Textract is charged per page, so the flag gates the capability here rather
        # than being consulted somewhere deeper.
        textract=ctx.textract if use_vision else None,
        llm=ctx.llm,
    )

    safe = re.sub(r'[^\w\s\-]', '', doc_title or '')[:40].strip().replace(' ', '_')
    filename = f'{safe}_Assessment.docx' if safe else 'Assessment_Report.docx'
    logger.info(
        "Built %s from sections %d-%d of a %d-page document",
        filename,
        start_idx,
        end_idx,
        total_pages,
    )

    return {
        "filename": filename,
        "docx": docx_bytes,
        "content_type": DOCX_CONTENT_TYPE,
        # ISO ships no editable sections: the app being replaced has no section editor,
        # so producing them would be an addition rather than a port.
        "sections": None,
    }
