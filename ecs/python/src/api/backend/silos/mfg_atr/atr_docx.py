# Ported from src/atr_docx.py.
#
# Verbatim: this module imports only third-party libraries and its siblings, so it
# needed no platform seam changes. It writes to a filesystem path, so the calling
# stage materialises a scratch directory and reads the bytes back.
"""Render an ATR to a native Word (.docx) document.

A sibling of ``pdf_builder`` for the "Download as Word" path: instead of scraping
the generated PDF (pdf2docx cannot read AcroForm widget values), this builds an
editable .docx straight from the shaped ``AtrReport`` so the layout and values match
the PDF — header, Regulatory Report, HQC Assessment, §1 Scope, §2 Summary, §3 Methods
table, and §4 Results (per-batch matrices, with the samples-as-columns fallback).

Formatting is aligned to the reportlab PDF via ``docx_common`` (Arial base font, navy
headings, shaded table headers, fixed column widths, zebra §4 rows, landscape when the
results are wide, and the same running header/footer).
"""
from __future__ import annotations

from pathlib import Path

from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt

from . import docx_common as dc
from .atr_report import AtrReport
from .atr_pdf import HQC_OPTIONS, TEMPLATE_ID, ManualDefaults

_HF_HEADER = "OneV-1327075 v 5.0  Effective on 09 Feb 2026  AbbVie Confidential General"


def _fit(widths_mm: list[float], avail_mm: float) -> list[float]:
    """Scale widths down proportionally so the table never exceeds the frame."""
    total = sum(widths_mm)
    if avail_mm > 0 and total > avail_mm:
        k = avail_mm / total
        return [w * k for w in widths_mm]
    return widths_mm


def _effective_cols(bm, is_stability: bool) -> int:
    """Column count a batch matrix will actually render with — after the transpose rule.
    A transposed dev-sample table is Test + Parameter + one column per sample."""
    if not is_stability and len(bm.sub_headers) > len(bm.rows):
        return 2 + len(bm.rows)
    return len(bm.lead_headers) + len(bm.sub_headers)


# --- automated blocks -------------------------------------------------------
def _header_block(doc, report: AtrReport) -> None:
    t = dc.grid(doc, 4, 2)
    t.cell(0, 0).merge(t.cell(0, 1))
    dc.set_cell(t.cell(0, 0), "ANALYTICAL REPORT", bold=True, size=10, fill="F5F5F5")
    for r, (label, value) in enumerate((
        ("CMC Request ID:", report.header["cmc_request_id"]),
        ("Project Code:", report.header["project_code"]),
        ("Batch/Sample ID:", report.header["batch_sample_id"]),
    ), start=1):
        dc.set_cell(t.cell(r, 0), label, bold=True, size=9, fill="F5F5F5")
        dc.set_cell(t.cell(r, 1), value, size=9)
    dc.set_col_widths(t, [45, 130])


def _cell_para(cell):
    """Empty a cell and return its paragraph, ready for custom runs / a checkbox."""
    cell.text = ""
    return cell.paragraphs[0]


def _sized(run, *, bold: bool = False, size: float = 9):
    run.bold = bold
    run.font.size = Pt(size)
    return run


def _regulatory_block(doc, d: ManualDefaults) -> None:
    """Bordered 2-column table (matches the PDF): the statement on the left, the Yes/No
    boxes on the right, all enclosed in cell borders."""
    t = dc.grid(doc, 1, 2)
    p = _cell_para(t.cell(0, 0))
    _sized(p.add_run("Regulatory report intended to support regulatory submissions, "
                     "pre-approval inspections or GMP activities. "), bold=True)
    _sized(p.add_run("(only required for PDS&T)"))
    p = _cell_para(t.cell(0, 1))
    dc.add_checkbox(p, checked=d.regulatory == "Yes"); _sized(p.add_run(" Yes        "))
    dc.add_checkbox(p, checked=d.regulatory == "No"); _sized(p.add_run(" No"))
    dc.set_col_widths(t, [137, 38])


def _hqc_block(doc, d: ManualDefaults) -> None:
    """Bordered table (matches the PDF): a spanned/shaded heading row, the Specification ID
    row, one spanned row per HQC option, and a spanned Remarks row — all with cell borders."""
    t = dc.grid(doc, 3 + len(HQC_OPTIONS), 2)   # heading + spec-id + options + remarks

    t.cell(0, 0).merge(t.cell(0, 1))            # heading spans both columns, shaded
    p = _cell_para(t.cell(0, 0))
    _sized(p.add_run("HQC Assessment "), bold=True)
    _sized(p.add_run("(only required if clinical material with an associated specification "
                     "was analyzed)"))
    dc.shade_cell(t.cell(0, 0), "F5F5F5")

    dc.set_cell(t.cell(1, 0), "Specification ID:", bold=True, size=9)
    dc.set_cell(t.cell(1, 1), d.spec_id or "", size=9)

    row = 2
    for key, text in HQC_OPTIONS:               # each option spans both columns
        t.cell(row, 0).merge(t.cell(row, 1))
        p = _cell_para(t.cell(row, 0))
        dc.add_checkbox(p, checked=d.hqc == key)
        _sized(p.add_run(f" {text}"))
        row += 1

    t.cell(row, 0).merge(t.cell(row, 1))        # remarks spans both columns
    p = _cell_para(t.cell(row, 0))
    _sized(p.add_run("Remarks: "), bold=True)
    dc.add_checkbox(p, checked=d.remarks_na)
    _sized(p.add_run(f" N/A        {'' if d.remarks_na else (d.remarks or '')}"))
    dc.set_col_widths(t, [45, 130])


def _methods_table(doc, report: AtrReport, section_no: int = 3) -> None:
    headers = ("Test no.", "Test description", "Test method reference", "ELN Unique ID")
    t = dc.grid(doc, 1 + len(report.methods), len(headers))
    for c, h in enumerate(headers):
        dc.set_cell(t.cell(0, c), h, bold=True, size=9)
    for i, m in enumerate(report.methods, start=1):
        dc.set_cell(t.cell(i, 0), f"{section_no}.{i}", size=9)
        dc.set_cell(t.cell(i, 1), m.technique, size=9)
        dc.set_cell(t.cell(i, 2), m.method_reference or "—", size=9)
        dc.set_cell(t.cell(i, 3), m.eln_unique_id or "", size=9)
    dc.shade_header_row(t, dc.GREY_FILL)
    dc.set_col_widths(t, [16, 80, 50, 29])


def _matrix_table(doc, bm, frame_mm: float) -> None:
    n_lead, n_sub = len(bm.lead_headers), len(bm.sub_headers)
    t = dc.grid(doc, 2 + len(bm.rows), n_lead + n_sub)
    # lead headers span both header rows (vertical merge)
    for c in range(n_lead):
        t.cell(0, c).merge(t.cell(1, c))
        dc.set_cell(t.cell(0, c), bm.lead_headers[c], bold=True, size=7)
    # group headers span their sub-columns (horizontal merge on row 0)
    col = n_lead
    for disp, span in bm.group_headers:
        if span > 1:
            t.cell(0, col).merge(t.cell(0, col + span - 1))
        dc.set_cell(t.cell(0, col), disp, bold=True, size=7)
        col += span
    for j, sub in enumerate(bm.sub_headers):
        dc.set_cell(t.cell(1, n_lead + j), sub, bold=True, size=7)
    for r, rowvals in enumerate(bm.rows, start=2):
        for c, val in enumerate(rowvals):
            dc.set_cell(t.cell(r, c), val, size=7)
    # blue two-level header + zebra data rows (matches the PDF)
    for c in range(n_lead + n_sub):
        dc.shade_cell(t.cell(0, c), dc.BLUE_FILL)
        dc.shade_cell(t.cell(1, c), dc.BLUE_FILL)
    dc.zebra_rows(t, start=2)
    # column widths: Sample wide, Timepoint/Storage narrow, rest split the frame
    lead_def = {"Sample": 34.0, "Timepoint": 16.0, "Storage": 14.0}
    lead_w = [lead_def.get(h, 22.0) for h in bm.lead_headers]
    sub_w = max(9.0, (frame_mm - sum(lead_w)) / max(n_sub, 1))
    dc.set_col_widths(t, _fit(lead_w + [sub_w] * n_sub, frame_mm))


def _transposed_matrix_table(doc, bm, frame_mm: float) -> None:
    """Measurements-as-rows layout (wide dev-sample ATRs), matching the PDF: Test /
    Parameter down the left (Test merged per technique), one column per sample."""
    n_lead = len(bm.lead_headers)
    sample_labels = []
    for r in bm.rows:
        leads = [str(r[c]) for c in range(n_lead) if str(r[c]).strip()]
        sample_labels.append(" / ".join(leads) or "Sample")
    n_samp = len(sample_labels)

    # technique for each measurement row; keep the label only on each group's first row
    tech_at = [""] * len(bm.sub_headers)
    off = 0
    for disp, span in bm.group_headers:
        if off < len(tech_at):
            tech_at[off] = disp
        off += span

    t = dc.grid(doc, 1 + len(bm.sub_headers), 2 + n_samp)
    dc.set_cell(t.cell(0, 0), "Test", bold=True, size=7)
    dc.set_cell(t.cell(0, 1), "Parameter", bold=True, size=7)
    for si, s in enumerate(sample_labels):
        dc.set_cell(t.cell(0, 2 + si), s, bold=True, size=7)
    for c in range(2 + n_samp):
        dc.shade_cell(t.cell(0, c), dc.BLUE_FILL)

    for j, sub in enumerate(bm.sub_headers):
        r = 1 + j
        zebra = dc.ZEBRA_FILL if j % 2 == 1 else None
        dc.set_cell(t.cell(r, 0), tech_at[j], bold=True, size=7, fill=zebra)
        dc.set_cell(t.cell(r, 1), str(sub), size=7, fill=zebra)
        for si in range(n_samp):
            actual = bm.actual_flags[si][n_lead + j]
            fill = dc.AMBER_FILL if actual else zebra    # amber wins over zebra
            dc.set_cell(t.cell(r, 2 + si), bm.rows[si][n_lead + j], size=7, fill=fill)

    # merge the Test column over consecutive rows of the same technique group
    r = 1
    for _disp, span in bm.group_headers:
        if span > 1:
            t.cell(r, 0).merge(t.cell(r + span - 1, 0))
        r += span

    test_w, param_w = 30.0, 55.0
    val_w = max(14.0, (frame_mm - test_w - param_w) / max(n_samp, 1))
    dc.set_col_widths(t, _fit([test_w, param_w] + [val_w] * n_samp, frame_mm))


def _results(doc, report: AtrReport, frame_mm: float, section_no: int = 4) -> None:
    if report.batch_matrices:
        for i, bm in enumerate(report.batch_matrices, start=1):
            dc.heading(doc, f"{section_no}.{i}  Batch {bm.batch_id}", size=11, keep_with_next=True)
            # Transpose wide dev-sample tables (measurements > samples) so they read
            # top-to-bottom like the PDF; stability ATRs keep samples as rows.
            if not report.is_stability and len(bm.sub_headers) > len(bm.rows):
                _transposed_matrix_table(doc, bm, frame_mm)
            else:
                _matrix_table(doc, bm, frame_mm)
        return
    if not report.results:
        doc.add_paragraph("No automated results available.")
        return
    samples = report.sample_columns
    t = dc.grid(doc, 1, 2 + len(samples))
    dc.set_cell(t.cell(0, 0), "Test", bold=True, size=8)
    dc.set_cell(t.cell(0, 1), "", bold=True, size=8)
    for c, s in enumerate(samples, start=2):
        dc.set_cell(t.cell(0, c), s, bold=True, size=8)
    for tech in report.results:
        for k, p in enumerate(tech.params):
            row = t.add_row().cells
            dc.set_cell(row[0], tech.technique if k == 0 else "", bold=True, size=8)
            dc.set_cell(row[1], p.label, size=8)
            for c, s in enumerate(samples, start=2):
                dc.set_cell(row[c], str(p.values.get(s, "")), size=8)
    dc.shade_header_row(t, dc.GREY_FILL)


def generate_atr_docx(report: AtrReport, out_path: str | Path, *,
                      defaults: ManualDefaults | None = None,
                      generated_at: str | None = None) -> Path:
    """Render the ATR as an editable .docx to ``out_path`` and return the path."""
    d = defaults or ManualDefaults()
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    # Match generate_atr_pdf: wide results → landscape. Count columns *after* the
    # transpose rule, so a dev-sample table that transposes to a few columns stays
    # portrait instead of being forced landscape by its raw measurement count.
    matrix_cols = max((_effective_cols(bm, report.is_stability)
                       for bm in report.batch_matrices), default=0)
    wide = report.is_stability or len(report.sample_columns) > 8 or matrix_cols > 8

    doc = dc.new_document(font="Arial", size=10)
    dc.configure_page(doc, left=18, right=12, top=18, bottom=16)
    if wide:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
    page_w = 297 if wide else 210
    frame_mm = page_w - 18 - 12

    dc.set_centered_header(doc, _HF_HEADER)
    dc.set_centered_footer(doc, f"Template Name/ID: {TEMPLATE_ID}   |   Pg ", page_suffix="")

    doc.add_paragraph(
        f"Document Title: NBE_AR_Analytical Report - {report.header['project_code']} - "
        f"{report.request_id_display}")

    _header_block(doc, report)
    doc.add_paragraph()
    _regulatory_block(doc, d)
    doc.add_paragraph()   # separate the two tables — adjacent Word tables merge into one
    _hqc_block(doc, d)

    # §1 Scope and §2 Summary and Conclusion are omitted when the warehouse holds no
    # entry (they previously rendered an empty "—"). Remaining sections renumber so the
    # sequence never has a gap; Methods and Results are always present.
    n = 0
    if report.scope:
        n += 1
        dc.heading(doc, f"{n} SCOPE")
        doc.add_paragraph(report.scope)
    if report.summary_conclusion:
        n += 1
        dc.heading(doc, f"{n} SUMMARY AND CONCLUSION")
        doc.add_paragraph(report.summary_conclusion)
    n += 1
    dc.heading(doc, f"{n} ANALYTICAL METHODS")
    _methods_table(doc, report, n)
    n += 1
    dc.heading(doc, f"{n} RESULTS")
    _results(doc, report, frame_mm, n)

    doc.save(str(out_path))
    return out_path
