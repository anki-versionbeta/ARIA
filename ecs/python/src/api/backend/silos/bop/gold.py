"""Length and style bounds derived from the gold reference BOP.

This is what keeps generated sections as short as the human-authored document. The
gold doc is measured (characters, sentences, list items) and those measurements are
turned into per-section rules that go into every generation prompt.

Ported verbatim. Two notes on fidelity:

* The source defines a `_format_procedure_fragment` helper that nothing calls —
  `operating_procedure` uses `OP_BOUND` directly. It is omitted here because it is
  provably unreachable; no behaviour depends on it.
* The gold object is ~25 MB, so it is read with caching disabled: only the *derived*
  bounds are cached, exactly as the source caches `_GOLD_BOUNDS_CACHE`. Holding the
  document itself would cost 25 MB per worker process for no benefit.
"""

from __future__ import annotations

import io
import logging
import re
from typing import Any

from .schema import OP_BOUND, SECTION_BOUNDS

logger = logging.getLogger(__name__)

GOLD_ASSET = "gold/gold_bop.docx"

_bounds_cache: dict[str, str] | None = None

# Leaf BOP key -> where to find it in the gold document.
# `h1` and `label` match case-insensitively by STARTSWITH, so minor formatting
# drift in the source document does not break the lookup.
GOLD_SECTION_MAP: dict[str, dict[str, str]] = {
    "purpose": {"h1": "1.0", "label": "1.1 Purpose"},
    "scope": {"h1": "1.0", "label": "1.2 Scope"},
    # abbreviations is intentionally NOT mapped: the gold's 1.3 Definitions cell is
    # nearly empty and would derive a "1 sentence" bound that contradicts the
    # [acronym, definition] pair schema. The static bound is correct there.
    "description": {"h1": "1.0", "label": "1.8 Equipment"},
    "safety.hazards": {"h1": "1.0", "label": "1.7 Safety", "sub": "Hazards"},
    "safety.engineering_controls": {
        "h1": "1.0",
        "label": "1.7 Safety",
        "sub": "Engineering Controls",
    },
    "safety.ppe": {"h1": "1.0", "label": "1.7 Safety", "sub": "PPE"},
    "operating_procedure.setup": {"h1": "3.0", "label": "3.1 Set"},
    "operating_procedure.startup": {"h1": "3.0", "label": "3.2 Start", "sub": "Start up"},
    "operating_procedure.operation": {
        "h1": "3.0",
        "label": "3.2 Start",
        "sub": "Operation",
    },
    "operating_procedure.shutdown": {
        "h1": "3.0",
        "label": "3.3 Shut",
        "sub": "Normal Shut",
    },
    "operating_procedure.emergency_shutdown": {
        "h1": "3.0",
        "label": "3.3 Shut",
        "sub": "Emergency",
    },
    "operating_procedure.maintenance": {"h1": "3.0", "label": "3.4 Maintenance"},
    # Whole-H1 case: everything under 4.0 joined.
    "related_documents": {"h1": "4.0", "label": ""},
}

# Which leaf keys belong to which top-level section, in the order they should appear
# in the aggregated bound.
TOP_TO_LEAVES: dict[str, list[str]] = {
    "description": [
        "description.tools",
        "description.spare_parts",
        "description.consumables",
        "description.materials",
    ],
    "safety": [
        "safety.hazards",
        "safety.engineering_controls",
        "safety.ppe",
    ],
    "operating_procedure": [
        "operating_procedure.setup",
        "operating_procedure.startup",
        "operating_procedure.operation",
        "operating_procedure.shutdown",
        "operating_procedure.emergency_shutdown",
        "operating_procedure.maintenance",
    ],
}

LEAF_LABEL_OVERRIDES: dict[str, str] = {
    "safety.ppe": "PPE",
    "safety.hazards": "Hazards",
    "safety.engineering_controls": "Engineering Controls",
}


def load_gold_doc(ctx) -> dict[str, dict[str, str]] | None:
    """Parse the gold BOP into `{H1 heading: {row label or H2: cell text}}`.

    Its content lives mostly in two-column tables (label || body), so paragraphs and
    tables are walked in document order and every table row is grouped under the most
    recent H1. Returns None on any failure; the caller falls back to static bounds.
    """
    try:
        data = ctx.assets.read(GOLD_ASSET, cache=False)
        logger.info("Loaded the gold document (%d bytes)", len(data))
    except Exception as exc:
        logger.info("Gold document unavailable, using static bounds: %s", exc)
        return None

    try:
        from docx import Document

        document = Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("Gold document failed to parse: %s", exc)
        return None

    result: dict[str, dict[str, str]] = {}
    current_h1: str | None = None
    paragraphs = iter(document.paragraphs)
    tables = iter(document.tables)

    # The dual iterator mirrors the manual extraction walker. It can mis-align if the
    # document wraps paragraphs in <w:sdt> content controls, which python-docx hides
    # from .paragraphs but exposes in the body XML. The current gold has none.
    for child in document.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            try:
                paragraph = next(paragraphs)
            except StopIteration:
                continue
            style = paragraph.style.name if paragraph.style else ""
            text = (paragraph.text or "").strip()
            if "Heading 1" in style and text:
                current_h1 = text
                result.setdefault(current_h1, {})
            elif "Heading 2" in style and text and current_h1:
                # H2s are sub-section markers; treat them as labelled empty rows.
                result[current_h1].setdefault(text, "")
        elif tag == "tbl":
            try:
                table = next(tables)
            except StopIteration:
                continue
            if current_h1 is None:
                continue
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if not cells:
                    continue
                if len(cells) >= 2 and cells[0]:
                    body = " | ".join(cell for cell in cells[1:] if cell)
                    if body:
                        result[current_h1][cells[0]] = body
    return result


def split_inline_subblocks(cell_text: str, headers: list[str]) -> dict[str, str]:
    """Split a pipe-delimited cell on inline sub-block headers.

    The gold puts several sub-sections in one cell with inline headers ("Hazards",
    "Engineering Controls", "PPE"). Missing headers map to "".
    """
    parts = [part.strip() for part in cell_text.split("|")]
    out: dict[str, str] = {header: "" for header in headers}
    current: str | None = None
    buffer: list[str] = []

    def match_header(part: str) -> str | None:
        lowered = part.lower()
        for header in headers:
            # STARTSWITH tolerates trailing parentheticals such as
            # "Hazards (engineering controls and/or PPE …)".
            if lowered.startswith(header.lower()):
                return header
        return None

    for part in parts:
        if not part:
            continue
        matched = match_header(part)
        if matched is not None:
            if current is not None:
                out[current] = " | ".join(buffer).strip()
            current = matched
            buffer = []
        elif current is not None:
            buffer.append(part)
    if current is not None:
        out[current] = " | ".join(buffer).strip()
    return out


def _starts_with_ci(value: str, prefix: str) -> bool:
    return value.lower().startswith(prefix.lower())


def resolve_gold_section(gold: dict[str, Any] | None, leaf_key: str) -> str:
    """Gold text for a leaf key, or "" when unmapped, absent or empty."""
    if gold is None:
        return ""
    spec = GOLD_SECTION_MAP.get(leaf_key)
    if not spec:
        return ""

    h1_dict = next(
        (value for key, value in gold.items() if _starts_with_ci(key, spec["h1"])),
        None,
    )
    if h1_dict is None:
        return ""

    label_prefix = spec.get("label", "")
    if not label_prefix:
        # Whole-H1 case, e.g. related_documents pulls everything under 4.0.
        return " | ".join(
            value for value in h1_dict.values() if isinstance(value, str) and value
        ).strip()

    cell = next(
        (
            value
            for key, value in h1_dict.items()
            if isinstance(value, str) and _starts_with_ci(key, label_prefix)
        ),
        "",
    )
    if not cell:
        return ""

    sub = spec.get("sub", "")
    if not sub:
        return cell

    # Normalise newline-delimited cells (the real gold) into pipe-delimited form so
    # the splitter works uniformly on both shapes.
    if "\n" in cell and "|" not in cell:
        cell = " | ".join(line.strip() for line in cell.splitlines() if line.strip())

    safety_headers = ["Hazards", "Engineering Controls", "PPE"]
    procedure_headers = ["Start up", "Operation"]
    shutdown_headers = ["Normal Shut down", "Emergency shut down"]
    if sub in safety_headers:
        headers = safety_headers
    elif sub in procedure_headers:
        headers = procedure_headers
    elif sub in shutdown_headers:
        headers = shutdown_headers
    else:
        headers = [sub]
    return split_inline_subblocks(cell, headers).get(sub, "")


def count_chars_sentences_items(text: str) -> tuple[int, int, int]:
    """(characters, sentences, list items) for a gold section body."""
    if not text:
        return (0, 0, 0)
    chars = len(text)
    pieces = [piece.strip() for piece in re.split(r"[.!?]\s+", text) if piece.strip()]
    sentences = max(1, len(pieces))
    if "|" in text:
        item_pieces = [piece.strip() for piece in text.split("|") if piece.strip()]
    else:
        item_pieces = [piece.strip() for piece in text.splitlines() if piece.strip()]
    return (chars, sentences, max(1, len(item_pieces)))


def _string_fragment(label: str, chars: int, sentences: int) -> str:
    sentence_phrase = (
        "1 sentence"
        if sentences <= 1
        else "1–2 sentences"
        if sentences == 2
        else f"{sentences} sentences"
    )
    return (
        f"{label}: approximately {chars} characters / {sentence_phrase}. "
        f"State the fact concisely; do NOT pad with applications, "
        f"configurations, or operating ranges that belong elsewhere."
    )


def _list_fragment(label: str, items: int) -> str:
    return (
        f"{label}: concise list of approximately {items}±2 entries. "
        f"Preserve part numbers, model numbers, catalog numbers, and "
        f"serial numbers EXACTLY as they appear in the source manual "
        f"(e.g. 'Wrench set cat. # 7750', 'ASI model 74090', "
        f"'serial # 710200067150'). Do NOT genericize."
    )


def _safety_fragment(label: str) -> str:
    """`chars` is deliberately not emitted — the "one short paragraph per category"
    rule is the actual calibration."""
    return (
        f"{label}: one short paragraph per category covered (e.g. "
        f"Pressurization, Thermal, Electrical for Hazards). Do NOT "
        f"expand any single category into multiple paragraphs."
    )


def _leaf_label(leaf: str) -> str:
    if leaf in LEAF_LABEL_OVERRIDES:
        return LEAF_LABEL_OVERRIDES[leaf]
    return leaf.split(".")[-1].replace("_", " ").title()


def compute_gold_bounds(gold: dict[str, Any] | None) -> dict[str, str]:
    """Aggregate resolved leaves into the top-level bound strings."""
    if gold is None:
        return {}

    out: dict[str, str] = {}

    for top_key, label in [
        ("purpose", "Purpose"),
        ("scope", "Scope"),
        ("abbreviations", "Definitions / Abbreviations"),
        ("related_documents", "Related Documents"),
    ]:
        text = resolve_gold_section(gold, top_key)
        if not text:
            continue
        chars, sentences, _ = count_chars_sentences_items(text)
        out[top_key] = _string_fragment(label, chars, sentences)

    description_fragments: list[str] = []
    for leaf in TOP_TO_LEAVES["description"]:
        text = resolve_gold_section(gold, leaf)
        if not text:
            continue
        _, _, items = count_chars_sentences_items(text)
        description_fragments.append(_list_fragment(_leaf_label(leaf), items))
    if not description_fragments:
        # Fall back to the bare "description" key (the 1.8 Equipment cell) so the
        # aggregated bound still carries a Tools fragment.
        text = resolve_gold_section(gold, "description")
        if text:
            _, _, items = count_chars_sentences_items(text)
            description_fragments.append(_list_fragment("Tools", items))
    if description_fragments:
        description_fragments.append(
            "For Components, Specifications, and Utilities fields: if "
            "the source manual gives a model / cat / catalog / serial "
            "number for the item, include it verbatim alongside the "
            "field's prose."
        )
        out["description"] = "\n\n".join(description_fragments)

    safety_fragments: list[str] = []
    for leaf in TOP_TO_LEAVES["safety"]:
        if resolve_gold_section(gold, leaf):
            safety_fragments.append(_safety_fragment(_leaf_label(leaf)))
    if safety_fragments:
        out["safety"] = "\n\n".join(safety_fragments)

    # The narrative-step rule is emitted ONCE. Emitting it per leaf bloated the
    # prompt without adding signal, and an explicit character target here made the
    # model produce ~30 KB of JSON that truncated at max_tokens and fell through to
    # the empty fallback.
    covered: list[str] = []
    skipped: list[str] = []
    for leaf in TOP_TO_LEAVES["operating_procedure"]:
        if resolve_gold_section(gold, leaf):
            covered.append(_leaf_label(leaf))
        else:
            skipped.append(leaf)
    if covered:
        out["operating_procedure"] = OP_BOUND
    if skipped:
        # Logged so a future gold swap that fills these in is observable.
        logger.info(
            "No gold text for %s; those fall through to the same narrative rule",
            skipped,
        )
    return out


def get_section_bounds(ctx) -> dict[str, str]:
    """Cached bounds: gold-derived where available, static elsewhere.

    The merge order means keys missing from the gold inherit the static text.
    """
    global _bounds_cache
    if _bounds_cache is not None:
        return _bounds_cache

    gold = load_gold_doc(ctx)
    if gold is None:
        _bounds_cache = dict(SECTION_BOUNDS)
        logger.info("Using static SECTION_BOUNDS (no gold available)")
        return _bounds_cache

    derived = compute_gold_bounds(gold)
    merged = dict(SECTION_BOUNDS)
    merged.update(derived)
    _bounds_cache = merged
    logger.info(
        "Derived %d bound(s) from gold; %d retained static (derived=%s)",
        len(derived),
        len(set(SECTION_BOUNDS) - set(derived)),
        sorted(derived),
    )
    return _bounds_cache


def reset_cache() -> None:
    global _bounds_cache
    _bounds_cache = None
