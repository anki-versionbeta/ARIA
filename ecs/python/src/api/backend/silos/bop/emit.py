"""Render the BOP document to .docx.

Ported verbatim: the template is opened for its named styles, the body stripped, and
content appended with explicit spacing and indent overrides. Section order, headings,
labels, indents, font sizes and page breaks are all as in the app being replaced —
this is the deliverable users compare against, so nothing here is tidied.
"""

from __future__ import annotations

import io
import os
from html.parser import HTMLParser
from typing import Any

TEMPLATE_ASSET = "templates/SOP template.docx"

INLINE_FLAGS = {
    "b": "bold",
    "strong": "bold",
    "i": "italic",
    "em": "italic",
    "u": "underline",
}

# Body indent depends on the depth of the heading immediately above it, so the
# visual hierarchy mirrors the heading nesting.
BODY_INDENT_INCHES = {1: 0.25, 2: 0.4, 3: 0.6}


class _DocxEmitter(HTMLParser):
    """Walks HTML and appends runs to a paragraph.

    Only bold/italic/underline are set; fonts come from the paragraph's named style.
    """

    def __init__(self, paragraph) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraph = paragraph
        self._stack: list[dict[str, bool]] = [{}]
        self._list_depth = 0
        self._list_counters: list[int] = []
        self._ordered_stack: list[bool] = []
        self._pending_prefix: str | None = None
        self._first_output = True

    def _fmt(self) -> dict[str, bool]:
        merged: dict[str, bool] = {}
        for frame in self._stack:
            merged.update(frame)
        return merged

    def _emit(self, text: str) -> None:
        if not text:
            return
        if self._pending_prefix is not None:
            text = self._pending_prefix + text
            self._pending_prefix = None
        fmt = self._fmt()
        run = self.paragraph.add_run(text)
        if fmt.get("bold"):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("underline"):
            run.underline = True
        self._first_output = False

    def _break(self) -> None:
        if self._first_output:
            return
        self.paragraph.add_run().add_break()

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in INLINE_FLAGS:
            self._stack.append({INLINE_FLAGS[tag]: True})
        elif tag == "br":
            self._break()
        elif tag in ("p", "div"):
            if not self._first_output:
                self._break()
            self._stack.append({})
        elif tag in ("ul", "ol"):
            self._list_depth += 1
            self._ordered_stack.append(tag == "ol")
            self._list_counters.append(0)
            self._stack.append({})
        elif tag == "li":
            if not self._first_output:
                self._break()
            # Some editors emit every list as <ol> and distinguish bullet from
            # ordered with `data-list` on each <li>. Honouring it stops an edited
            # bullet list rendering as "1. 2. 3." in the docx.
            data_list = dict(attrs or []).get("data-list", "")
            if data_list == "bullet":
                is_ordered = False
            elif data_list == "ordered":
                is_ordered = True
            else:
                is_ordered = bool(self._ordered_stack and self._ordered_stack[-1])
            indent = "  " * max(self._list_depth - 1, 0)
            if is_ordered:
                if self._list_counters:
                    self._list_counters[-1] += 1
                    number = self._list_counters[-1]
                else:
                    number = 1
                self._pending_prefix = f"{indent}{number}. "
            else:
                self._pending_prefix = f"{indent}• "
            self._stack.append({})
        else:
            self._stack.append({})

    def handle_endtag(self, tag):
        if self._stack:
            self._stack.pop()
        if tag.lower() in ("ul", "ol"):
            self._list_depth = max(0, self._list_depth - 1)
            if self._ordered_stack:
                self._ordered_stack.pop()
            if self._list_counters:
                self._list_counters.pop()

    def handle_data(self, data):
        self._emit(data)


def html_to_runs(paragraph, html: str | None) -> None:
    if html is None:
        return
    emitter = _DocxEmitter(paragraph)
    emitter.feed(html)
    emitter.close()


def _is_html(value: Any) -> bool:
    return isinstance(value, str) and "<" in value and ">" in value


def _strip_body(document) -> None:
    body = document.element.body
    for child in list(body):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in ("p", "tbl"):
            body.remove(child)


def _setup_page_layout(document, title: str = "Basis of Procedure") -> None:
    """1-inch margins, a centred bold header, and a live "Page X of Y" footer.

    Section properties live in <w:sectPr>, which `_strip_body` preserves.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.text = ""
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.add_run(title).bold = True

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.text = ""
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.add_run("Page ")

    def add_field(run, instruction: str) -> None:
        # python-docx has no helper for field codes, so drop to OOXML. Fields are
        # used so the page numbers update when the document is opened.
        for command, value in (("begin", None), (None, instruction), ("end", None)):
            if command in ("begin", "end"):
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), command)
            else:
                element = OxmlElement("w:instrText")
                element.text = value
            run._r.append(element)

    add_field(footer_paragraph.add_run(), "PAGE")
    footer_paragraph.add_run(" of ")
    add_field(footer_paragraph.add_run(), "NUMPAGES")


def _add_heading(document, text: str, level: int, page_break_before: bool = False):
    """Heading with explicit spacing — the template's styles inherit very little."""
    from docx.shared import Pt

    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    # Run-level emphasis so headings render bold and underlined at a fixed size
    # regardless of what the template's heading styles inherit. Family and colour
    # still come from the named style.
    run.bold = True
    run.underline = True
    run.font.size = Pt({1: 16, 2: 14, 3: 12}.get(level, 11))
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt({1: 18, 2: 12, 3: 6}.get(level, 4))
    fmt.space_after = Pt({1: 8, 2: 6, 3: 3}.get(level, 3))
    if page_break_before:
        fmt.page_break_before = True
    return paragraph


def _add_body(document, text_or_html: str, parent_level: int = 1):
    from docx.shared import Inches, Pt

    if not text_or_html:
        return None
    paragraph = document.add_paragraph(style="Normal")
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(BODY_INDENT_INCHES.get(parent_level, 0.25))
    fmt.space_after = Pt(6)
    if _is_html(text_or_html):
        html_to_runs(paragraph, text_or_html)
    else:
        paragraph.add_run(text_or_html)
    return paragraph


def _add_bullets(document, items: list[Any], parent_level: int = 2) -> None:
    from docx.shared import Inches, Pt

    indent = Inches(BODY_INDENT_INCHES.get(parent_level, 0.4))
    for item in items:
        if item is None:
            continue
        text = item if isinstance(item, str) else str(item)
        if not text.strip():
            continue
        paragraph = document.add_paragraph(style="List Paragraph")
        fmt = paragraph.paragraph_format
        fmt.left_indent = indent
        fmt.space_after = Pt(3)
        if _is_html(text):
            html_to_runs(paragraph, text)
        else:
            paragraph.add_run(text)


def _add_table(document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            text = value if isinstance(value, str) else ("" if value is None else str(value))
            if _is_html(text):
                html_to_runs(paragraph, text)
            else:
                paragraph.add_run(text)


def _phase_groups(phase: Any) -> tuple[list[Any], list[Any]]:
    """Normalise a phase into (equipment, software) lists for rendering."""
    if isinstance(phase, dict):
        equipment = phase.get("equipment") or []
        software = phase.get("software") or []
        return (
            equipment if isinstance(equipment, list) else [equipment],
            software if isinstance(software, list) else [software],
        )
    if isinstance(phase, list):
        return (phase, [])
    if isinstance(phase, str) and phase.strip():
        return ([phase], [])
    return ([], [])


def _labeled_items(document, mapping: dict[str, str], labels: list[tuple[str, str]]) -> None:
    for label_text, key in labels:
        _add_heading(document, label_text, 3)
        _add_body(document, mapping.get(key, ""), parent_level=3)


def build_docx(
    bop_json: dict[str, Any], template_bytes: bytes, source_name: str | None = None
) -> bytes:
    from docx import Document

    document = Document(io.BytesIO(template_bytes))
    _strip_body(document)

    title = "Basis of Procedure"
    if source_name:
        stem = os.path.splitext(source_name)[0] or source_name
        title = f"BOP — {stem}"
    _setup_page_layout(document, title)

    _add_heading(document, "PURPOSE", 1)  # first H1, no page break needed
    _add_body(document, bop_json.get("purpose", ""), parent_level=1)

    _add_heading(document, "SCOPE", 1)  # short — shares the page with PURPOSE
    _add_body(document, bop_json.get("scope", ""), parent_level=1)

    _add_heading(document, "DESCRIPTION", 1, page_break_before=True)
    description = bop_json.get("description", {}) or {}

    _add_heading(document, "Components", 2)
    _labeled_items(
        document,
        description.get("components", {}) or {},
        [
            ("Vessel:", "vessel"),
            ("Pump:", "pump"),
            ("Agitation or Motor:", "agitation"),
            ("Instrumentation:", "instrumentation"),
            ("Computer and Software:", "software"),
        ],
    )

    _add_heading(document, "Specifications", 2)
    _labeled_items(
        document,
        description.get("specifications", {}) or {},
        [
            ("Dimensions:", "dimensions"),
            ("Capacity:", "capacity"),
            ("Temperature Operating Range:", "temperature_range"),
            ("Pressure Operating Range:", "pressure_range"),
            ("Wetted Materials of Construction:", "wetted_materials"),
        ],
    )

    _add_heading(document, "Utility Requirements", 2)
    _labeled_items(
        document,
        description.get("utilities", {}) or {},
        [
            ("Electrical:", "electrical"),
            ("Heating & Cooling:", "heating_cooling"),
            ("Water & Steam:", "water_steam"),
            ("Air/Nitrogen/Gases:", "gases"),
        ],
    )

    _add_heading(document, "Tools", 2)
    _add_bullets(document, description.get("tools", []) or [])
    _add_heading(document, "Durable Spare Parts", 2)
    _add_bullets(document, description.get("spare_parts", []) or [])
    _add_heading(document, "Consumable Replacement Parts", 2)
    _add_bullets(document, description.get("consumables", []) or [])
    _add_heading(document, "Materials", 2)
    _add_bullets(document, description.get("materials", []) or [])

    _add_heading(document, "SAFETY", 1, page_break_before=True)
    safety = bop_json.get("safety", {}) or {}
    _add_heading(document, "Hazards", 2)
    _add_body(document, safety.get("hazards", ""), parent_level=2)
    _add_heading(document, "Engineering Controls", 2)
    _add_body(document, safety.get("engineering_controls", ""), parent_level=2)
    _add_heading(document, "PPE", 2)
    _add_body(document, safety.get("ppe", ""), parent_level=2)

    _add_heading(document, "OPERATING PROCEDURE", 1, page_break_before=True)
    procedure = bop_json.get("operating_procedure", {}) or {}
    for label, key in [
        ("Setup", "setup"),
        ("Startup", "startup"),
        ("Routine Operation", "operation"),
        ("Shutdown", "shutdown"),
        ("Emergency Shutdown", "emergency_shutdown"),
        ("Cleaning", "cleaning"),
        ("Storage", "storage"),
        ("Maintenance", "maintenance"),
    ]:
        _add_heading(document, label, 2)
        equipment, software = _phase_groups(procedure.get(key))
        if equipment:
            _add_heading(document, "Equipment", 3)
            _add_bullets(document, equipment, parent_level=3)
        if software:
            _add_heading(document, "Software", 3)
            _add_bullets(document, software, parent_level=3)

    _add_heading(document, "Troubleshooting", 2)
    trouble_rows: list[list[str]] = []
    for entry in procedure.get("troubleshooting", []) or []:
        if isinstance(entry, dict):
            trouble_rows.append([entry.get("issue", ""), entry.get("resolution", "")])
        elif isinstance(entry, (list, tuple)) and len(entry) >= 2:
            trouble_rows.append([str(entry[0]), str(entry[1])])
    _add_table(document, ["Issue", "Resolution"], trouble_rows)

    _add_heading(document, "ABBREVIATIONS AND DEFINITIONS", 1, page_break_before=True)
    abbreviation_rows: list[list[str]] = []
    for entry in bop_json.get("abbreviations", []) or []:
        if isinstance(entry, (list, tuple)) and len(entry) >= 2:
            abbreviation_rows.append([str(entry[0]), str(entry[1])])
        elif isinstance(entry, dict):
            abbreviation_rows.append(
                [str(entry.get("acronym", "")), str(entry.get("definition", ""))]
            )
    _add_table(document, ["Acronym", "Definition"], abbreviation_rows)

    _add_heading(document, "RELATED DOCUMENTS", 1, page_break_before=True)
    _add_bullets(document, bop_json.get("related_documents", []) or [], parent_level=1)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
