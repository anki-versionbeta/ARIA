"""Ingest the NBE (Vials) Smartsheet into the canonical DB:
product / site / vendor / cap_color / product_site + field_provenance.
Grain = product-presentation (one row per spreadsheet data row).

This is the FALLBACK reader, for when the live Smartsheet API is not configured. Two other modules import
it for its column contract and parsing helpers (`COLS`, `PROV_FIELDS`, `clean`, `parse_color_code`,
`split_sites`), so it is on the module-scope import path of `router.py` and of all three stages — which is
why openpyxl is imported inside `main()` and not here.

That is not tidiness. `openpyxl` is absent from `da-backend/requirements.txt`, so a module-scope import made
it a LOAD-TIME requirement of this silo's router on any host provisioned from that file, and
`silo_registry._load_router` catches an import failure, logs it and returns None — so PSA would have
appeared in `GET /api/silos` with correct label and stages while every endpoint was silently absent and the
React feature 404'd. Under ARIA this reader is unreachable anyway (`aria.py` sets `ingest_source="api"` and
`smartsheet_credentials()` raises when the token is missing), so nothing should pay for it at import.
"""
import os, re, sqlite3

from . import paths


def xlsx_path():
    """The stale .xlsx export — the fallback source when the live Smartsheet is not configured."""
    return os.path.join(paths.root(), "Input_Data_Sources", "Smartsheets",
                        "NBE (Vials) Product Similarity Database.xlsx")

# product attribute  ->  0-based Smartsheet column index
COLS = {
    "program_no": 0, "program_name": 1, "route_of_admin": 2, "strength": 3,
    "target_fill_volume_ml": 4, "vial_container_size": 5, "product_color": 6,
    "modality": 7, "contact": 8, "status": 9, "list_number": 10,
    "batch_type": 11, "_cap_color": 12, "_dp_mfr": 13, "_dp_pkging": 14,
    "launch_year": 15, "form": 16, "_cap_vendor": 17, "shape": 18, "marking": 19,
    "picture1": 20, "picture2": 21, "last_modified": 22, "last_updated_by": 23,
    "verified_flag": 24, "link_dpsa": 25, "link_pdd": 26, "link_tpp": 27,
}
# scalar product columns that get a provenance row when non-null
PROV_FIELDS = ["program_no", "program_name", "route_of_admin", "strength",
               "target_fill_volume_ml", "vial_container_size", "product_color",
               "modality", "contact", "status", "list_number", "batch_type",
               "launch_year", "form", "shape", "marking", "verified_flag"]

SKIP_SITE = {"", "TBD", "TBC", "N/A", "TO BE ADDED"}
SHEET = "NBE (Vials) Product Similarity "

# Address cues used to tell an AbbVie *site code* cell (e.g. "ABB", "AP16\nLU",
# "TPM (Almac)") from a free-text postal *address* entered in the site column
# (e.g. BoNT/E's manufacturing site "...Castlebar Road, Westport, County Mayo, Ireland").
# A postal address names ONE physical (manufacturing) site, so split_sites collapses it to
# a SINGLE site node — it must still appear as its own Part E block (QPP §3: one Part E per
# manufacturing site, incl. third-party contract) — instead of being shredded into one
# phantom site per comma fragment. The raw cell is also kept verbatim (dp_mfr_site_raw) for Part B.
_ADDR_STREET = re.compile(
    r"\b(road|rd|street|st|avenue|ave|lane|ln|boulevard|blvd|drive|dr|highway|hwy)\b", re.I)
_ADDR_GEO = re.compile(
    r"\b(ireland|italy|germany|france|spain|switzerland|netherlands|belgium|"
    r"puerto rico|united states|u\.?s\.?a|county|province|postal|zip)\b", re.I)


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _looks_like_address(raw):
    """True when a site cell is really a postal address, not one or more site codes.
    Conservative: only fires on clear address cues (a street type, a country/geography
    term, four+ comma/newline parts, or a part of three+ words) so legitimate multi-word
    site names like 'TPM (Almac)' are NOT misclassified."""
    if not raw:
        return False
    s = str(raw)
    if _ADDR_STREET.search(s) or _ADDR_GEO.search(s):
        return True
    parts = [t.strip() for t in re.split(r"[\n,]+", s) if t.strip()]
    if len(parts) >= 4:
        return True
    # a three-plus-word part is prose (an address), unless it is a parenthesised site
    # name such as "TPM (UPS VDL)" — those are legitimate multi-token site codes.
    return any(len(p.split()) >= 3 and "(" not in p for p in parts)


def split_sites(raw):
    if not raw:
        return []
    if _looks_like_address(raw):
        # One postal address = one physical site (e.g. BoNT/E's Westport manufacturing
        # site). Collapse to a SINGLE site node (newlines -> ", ") so it still yields one
        # Part E block, rather than splitting the address into phantom per-fragment sites.
        one = re.sub(r"\s+", " ", str(raw).replace("\n", ", ")).strip().strip(",").strip()
        return [one] if one and one.upper() not in SKIP_SITE else []
    out = []
    for tok in re.split(r"[\n,]+", str(raw)):
        t = tok.strip()
        if t and t.upper() not in SKIP_SITE and not t.upper().startswith("N/A"):
            out.append(t)
    return out


def parse_color_code(name):
    if not name:
        return None
    m = re.findall(r"(L\d+|\b[A-Z]?\d{3,4}[A-Z]?\b)", name)
    return m[-1] if m else None


def main():
    # Imported here, not at module scope — see the module docstring. This is the only function that needs
    # openpyxl, and it only runs when the live Smartsheet API is unavailable.
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(xlsx_path(), data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))

    con = sqlite3.connect(paths.db_path())
    con.execute("PRAGMA foreign_keys = ON")
    cur = con.cursor()
    site_id, vendor_id = {}, {}

    def get_site(code):
        if code not in site_id:
            cur.execute("INSERT INTO site (site_code, site_name) VALUES (?,?)", (code, code))
            site_id[code] = cur.lastrowid
        return site_id[code]

    def get_vendor(name):
        if name not in vendor_id:
            cur.execute("INSERT INTO vendor (vendor_name) VALUES (?)", (name,))
            vendor_id[name] = cur.lastrowid
        return vendor_id[name]

    n_prod = 0
    for ridx, row in enumerate(rows[1:], start=2):     # row 1 = header
        cells = {k: clean(row[i]) for k, i in COLS.items()}
        if not cells["program_no"] and not cells["program_name"]:
            continue                                    # skip fully-empty rows
        n_prod += 1

        cap_vendor = cells["_cap_vendor"]
        vid = get_vendor(cap_vendor) if cap_vendor and cap_vendor.upper() not in ("N/A",) else None

        prod = {k: cells[k] for k in COLS if not k.startswith("_")}
        prod["source_row"] = ridx
        prod["cap_vendor_id"] = vid
        prod["dp_mfr_site_raw"] = cells["_dp_mfr"]
        prod["dp_pkging_site_raw"] = cells["_dp_pkging"]

        keys = list(prod.keys())
        cur.execute(
            f"INSERT INTO product ({','.join(keys)}) VALUES ({','.join('?' * len(keys))})",
            [prod[k] for k in keys])
        pid = cur.lastrowid

        # provenance for each non-null Smartsheet scalar
        for f in PROV_FIELDS:
            if cells[f] is not None:
                loc = f"{ws.title}!{get_column_letter(COLS[f] + 1)}{ridx}"
                cur.execute(
                    "INSERT INTO field_provenance "
                    "(entity, entity_id, field_name, source_doc, source_locator, match_type, extraction_method) "
                    "VALUES ('product', ?, ?, 'Smartsheet', ?, 'Exact', 'auto')",
                    (pid, f, loc))

        # cap_color (1 per presentation in the sheet)
        if cells["_cap_color"]:
            cur.execute(
                "INSERT INTO cap_color (product_id, batch_type, cap_color_name, color_code, vendor_id) "
                "VALUES (?,?,?,?,?)",
                (pid, cells["batch_type"], cells["_cap_color"],
                 parse_color_code(cells["_cap_color"]), vid))

        # product_site bridge (multi-valued cells)
        for code in split_sites(cells["_dp_mfr"]):
            cur.execute("INSERT OR IGNORE INTO product_site VALUES (?,?,?)", (pid, get_site(code), "mfr"))
        for code in split_sites(cells["_dp_pkging"]):
            cur.execute("INSERT OR IGNORE INTO product_site VALUES (?,?,?)", (pid, get_site(code), "pkging"))

    con.commit()
    counts = {t: con.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
              for t in ("product", "site", "vendor", "cap_color", "product_site", "field_provenance")}
    con.close()
    print(f"Ingested from row 2..{len(rows)}")
    for t, c in counts.items():
        print(f"  {t}: {c}")


if __name__ == "__main__":
    main()
