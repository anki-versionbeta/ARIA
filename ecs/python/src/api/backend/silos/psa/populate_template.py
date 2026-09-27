"""Populate the PSA template for one program from the DB + template_field_map.
Smartsheet/derived fields are filled; everything else gets a typed placeholder.
Usage: populate(program_no='AGN-151586')."""
import copy, io, os, re, sqlite3
from docx import Document
from docx.shared import Inches
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import asset_store, paths

SITE_NAMES = {"AP16": "AbbVie US, AP16", "LU": "AbbVie DE, LU"}
SKIP_SITE = {"", "TBD", "TBC", "N/A", "TO BE ADDED"}
PDD_ADDR = "[PENDING: PDD - full site address]"
PDD_LOC = "[PENDING: PDD - site location]"
STRENGTH_REVIEW = "[PENDING: per-vial strength]"


def _codes(raw):
    if not raw:
        return []
    return [t.strip() for t in re.split(r"[\n,]+", str(raw))
            if t.strip() and t.strip().upper() not in SKIP_SITE
            and not t.strip().upper().startswith("N/A")]


def _family(form):
    if not form:
        return "TBD"
    return "Lyophilized drug products" if "lyo" in form.lower() else f"{form} drug products"


def _risk_box(decision):
    d = (decision or "").strip().lower()
    if d == "no":
        return "☐  Yes              ☒  No"
    if d == "yes":
        return "☒  Yes              ☐  No"
    return None


BOX_GLYPHS = ("☒", "☐")        # ☒ checked (U+2612), ☐ unchecked (U+2610)
_BOX_SPLIT = re.compile("([" + "".join(BOX_GLYPHS) + "])")

# Copied verbatim from the blank template's OWN ballot-box runs, which is why the boxes render exactly
# as they do on the approved form. Confirmed by unzipping psa/assets/PSA_template.docx: all four of its
# ☐ runs carry precisely this rFonts element. MS Gothic contains U+2610/U+2612; `w:hint="eastAsia"` is
# what tells Word to resolve these characters through the eastAsia font rather than the Latin one.
BOX_RFONTS = {
    "w:ascii": "MS Gothic",
    "w:eastAsia": "MS Gothic",
    "w:hAnsi": "MS Gothic",
    "w:cs": "Times New Roman",
    "w:hint": "eastAsia",
}


def _write_paragraph(paragraph, text):
    """Write `text` into `paragraph`, giving the ballot-box glyphs the template's own font.

    `cell.text = ...` creates a run with NO `w:rFonts`, so Word resolved the glyph's font from the style
    hierarchy — which for the Similarity Risk cell does not reach a font containing U+2610/U+2612, and the
    boxes rendered as tiny unreadable marks. The surrounding "Yes"/"No" text keeps the form's typeface;
    only the two glyph runs are pinned.

    Matching the template rather than inventing a font matters twice over: the output is formatted
    identically to the approved blank form, and it adds no dependency on a font that might be absent
    (the first attempt pinned "Segoe UI Symbol", which works on Windows but is not what the form uses).
    """
    for chunk in _BOX_SPLIT.split(str(text)):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        if chunk in BOX_GLYPHS:
            fonts = OxmlElement("w:rFonts")
            for attr, value in BOX_RFONTS.items():
                fonts.set(qn(attr), value)
            # rFonts must come first inside rPr.
            run._element.get_or_add_rPr().insert(0, fonts)


def _set_cell_text(cell, text):
    """Write text, honouring newlines as separate paragraphs (multi-paragraph comments).

    Goes through `_write_paragraph` rather than assigning `cell.text`, so every writer — the generic
    `template_field_map` loop, Part D's `_derive`, and `_fill_e_table` — gets the glyph fix for free.
    The plain characters are still what lands in the text, so `verify.py`'s `"☒" in ...` checks and the
    tests that rely on them keep working.
    """
    parts = str(text).split("\n")
    cell.text = ""
    _write_paragraph(cell.paragraphs[0], parts[0])
    for extra in parts[1:]:
        _write_paragraph(cell.add_paragraph(), extra)


def _embed_image(cell, source, width_in=1.9, max_height_in=1.4):
    """Embed the product photo scaled (aspect preserved) to fit within a modest box: wide enough to be
    clearly visible, but height-bounded so it does not balloon the Photo row enough to reflow the
    document and push a table under the repeating page header.

    `source` is a path OR the image bytes. Bytes are what the ARIA `report` stage has: it reads the photo
    back from the run's stored media, where the `fetch` stage put it, so nothing depends on a file being
    on this container's disk.
    """
    cell.text = ""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    pic = cell.paragraphs[0].add_run().add_picture(source, width=Inches(width_in))
    cap = Inches(max_height_in)
    if pic.height > cap:
        pic.width = int(pic.width * cap / pic.height)
        pic.height = int(cap)


def _page_break_para():
    """A bare <w:p> holding a page break, inserted between successive Part E site blocks."""
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def _fill_e_table(t, s):
    """Fill one Part E table (table-3 layout) for a single site assessment dict."""
    _set_cell_text(t.rows[3].cells[1], s.get("site", "TBD"))                 # E.1 AbbVie Site
    _set_cell_text(t.rows[4].cells[1], s.get("families", "TBD"))            # E.1 Product Families
    _set_cell_text(t.rows[5].cells[1], _risk_box(s.get("risk")) or "")      # E.1 Similarity Risk box
    _set_cell_text(t.rows[7].cells[0], s.get("comments", ""))              # E.2 Comments
    if s.get("qa_label"):                                                   # E.3 "Site QA – <site>"
        _set_cell_text(t.rows[11].cells[0], s["qa_label"])


def _render_part_e(doc, e_sites):
    """Render Part E as ONE block PER SITE. The template has a single Part E table (index 3); the
    first site fills it, and each additional site gets a page break + a deep-copied clone after it."""
    tbl = doc.tables[3]
    if not e_sites:
        e_sites = [{"site": "TBD", "families": "No co-located products identified in the catalogue.",
                    "risk": "No",
                    "comments": "No in-scope co-located products were found; no mix-up risk identified."}]
    blank = copy.deepcopy(tbl._element)          # pristine Part E structure, for cloning
    parent = tbl._parent
    _fill_e_table(tbl, e_sites[0])
    last = tbl._element
    for s in e_sites[1:]:
        pb = _page_break_para()
        last.addnext(pb)
        clone = copy.deepcopy(blank)
        pb.addnext(clone)
        _fill_e_table(Table(clone, parent), s)
        last = clone


def _derive(key, p, docattrs, placeholder):
    if key == "product_name_full":
        # INN from TPP if extracted, else the Smartsheet short name; append strength only when known
        base = p["inn_name"] or p["program_name"] or p["program_no"] or "TBD"
        return f"{base} ({p['per_vial_strength']})" if p["per_vial_strength"] else base
    if key == "mfr_site_expand":
        if docattrs.get("b_manufacturing_site"):        # assessment-confirmed (matches signed form, multi-line)
            return docattrs["b_manufacturing_site"]
        if p["mfr_site_address"]:                       # full address extracted from the uploaded document
            return p["mfr_site_address"]
        raw = (p["dp_mfr_site_raw"] or "").replace("\n", ", ").strip()
        return raw if raw else "TBD"                    # no address found → site code(s) only, no placeholder
    if key == "pkging_site_expand":
        codes = _codes(p["dp_pkging_site_raw"])
        return "\n".join(SITE_NAMES.get(c, c) for c in codes) if codes else "TBD"   # one per line
    if key == "product_families_latestage":
        if docattrs.get("d1_product_families"):
            return docattrs["d1_product_families"]
        fam = _family(p["form"])
        return f"{fam} in late-stage development" if fam != "TBD" else "TBD"
    if key == "product_families":
        return docattrs.get("e1_product_families") or _family(p["form"])
    if key == "abbvie_site":
        if docattrs.get("e1_abbvie_site"):              # assessment-provided (matches signed form)
            return docattrs["e1_abbvie_site"]
        codes = _codes(p["dp_pkging_site_raw"])
        base = " / ".join(codes) if codes else "TBD"
        if p["site_location"]:                          # PDD-extracted location
            return f"{base} {p['site_location']}" if codes else p["site_location"]
        return (base + f"  {PDD_LOC}") if codes else "TBD"
    if key == "site_qa_function":
        codes = _codes(p["dp_pkging_site_raw"])
        return f"Site QA - {codes[0]}" if codes else "Site QA - TBD"
    if key == "pd_director":
        return docattrs.get("pd_director_name") or placeholder   # TPP approver via doc_attribute
    if key in ("risk_d1", "risk_e1"):                            # analyst determination (assessment file)
        attr = "d1_similarity_risk" if key == "risk_d1" else "e1_similarity_risk"
        return _risk_box(docattrs.get(attr)) or placeholder
    if key == "d2_comments":
        return docattrs.get("d2_comments") or placeholder
    if key == "e2_comments":
        return docattrs.get("e2_comments") or placeholder
    if key == "photo":
        return docattrs.get("product_photo") or ""   # Smartsheet image path, else blank (no placeholder)
    return "TBD"


def _resolve(fm, p, docattrs):
    src = fm["source"]
    if src == "smartsheet":
        v = p[fm["resolver"]]
        return v if v not in (None, "") else "TBD"
    if src == "derived":
        return _derive(fm["resolver"], p, docattrs, fm["placeholder_tag"])
    return fm["placeholder_tag"]            # pdd | tpp | analysis | execution


def populate(program_no="AGN-151586", product_id=None, photo=None):
    # product_id: when a program code spans several presentation rows, the caller (workflow) picks the
    # row that best matches the uploaded document and passes its product_id here so the right one is used.
    # photo: raw image bytes to embed instead of resolving a path from the database. The ARIA `report`
    # stage passes what `fetch` stored as run media; every other caller leaves it None and the
    # Smartsheet-ingested file on disk is used exactly as before.
    con = sqlite3.connect(paths.db_path())
    con.row_factory = sqlite3.Row
    if product_id is not None:
        rows = con.execute("SELECT * FROM product WHERE product_id=?", (product_id,)).fetchall()
    else:
        rows = con.execute("SELECT * FROM product WHERE program_no=?", (program_no,)).fetchall()
    if not rows:
        raise SystemExit(f"No product with program_no={program_no!r}")
    if product_id is None and len(rows) > 1:
        print(f"NOTE: {len(rows)} presentations for {program_no}; using source_row {rows[0]['source_row']} "
              f"(pass product_id to select a specific presentation)")
    p = rows[0]
    # PDD/TPP-extracted values not stored as product columns (e.g. the approver).
    # product_photo now only ever comes from the Smartsheet (PDD photo extraction is disabled).
    docattrs = {r["attribute_key"]: r["value_text"] for r in con.execute(
        "SELECT da.attribute_key, da.value_text FROM doc_attribute da "
        "JOIN doc_source ds ON da.doc_id = ds.doc_id WHERE ds.product_id = ?", (p["product_id"],))}
    fields = con.execute("SELECT * FROM template_field_map ORDER BY field_id").fetchall()
    con.close()

    # Parts D & E auto-fill (draft). Part D is a single block driven by template_field_map via docattrs;
    # Part E is rendered dynamically, one block per site. setdefault → an analyst assessment record
    # (already in docattrs from ingest_assessment) always wins over the auto-fill.
    from . import risk        # imported here: risk imports this module's siblings, not this module
    risk_form = risk.form_fields(paths.db_path(), p["program_no"], p["product_id"])
    for k in ("d1_product_families", "d1_similarity_risk", "d2_comments"):
        if risk_form.get(k) is not None:
            docattrs.setdefault(k, risk_form[k])

    doc = Document(io.BytesIO(asset_store.read("PSA_template.docx")))
    filled, deferred = [], []
    for fm in fields:
        value = _resolve(fm, p, docattrs)
        cell = doc.tables[fm["table_index"]].rows[fm["row_index"]].cells[fm["value_col"]]
        is_photo = fm["resolver"] == "photo"
        # photo path may be absolute (Smartsheet image saved under ASSETS) or root-relative — resolve both
        img_abs = ((value if os.path.isabs(value) else os.path.join(paths.root(), value))
                   if is_photo and value else None)
        if is_photo and photo:
            # Explicit bytes win: under ARIA there is no ingested file to point at.
            _embed_image(cell, photo)
            disp = f"<image: {len(photo)} bytes from the run snapshot>"
        elif is_photo and img_abs and "[PENDING" not in value and os.path.exists(img_abs):
            _embed_image(cell, img_abs)          # modest fixed box (see _embed_image)
            disp = f"<image: {value}>"
        else:
            _set_cell_text(cell, value)            # newlines → paragraphs (comments)
            disp = value if len(value) <= 80 else value[:77] + "..."
        # classify by the actual value so derived fields that still resolve to a
        # placeholder (e.g. strength held for review) are listed as deferred
        (deferred if "[PENDING" in value else filled).append((fm["field_label"], disp))

    # ---- Part E: one block per site (analyst record → single block; else risk-engine per-site) ----
    analyst_e_keys = ("e1_abbvie_site", "e1_product_families", "e1_similarity_risk", "e2_comments")
    if any(docattrs.get(k) for k in analyst_e_keys):
        e_sites = [{"site": docattrs.get("e1_abbvie_site") or _derive("abbvie_site", p, docattrs, ""),
                    "families": docattrs.get("e1_product_families") or _family(p["form"]),
                    "risk": docattrs.get("e1_similarity_risk") or "No",
                    "comments": docattrs.get("e2_comments") or "", "qa_label": None}]
    else:
        e_sites = risk_form.get("_e_sites") or []
    _render_part_e(doc, e_sites)
    for s in e_sites:
        filled.append((f"E: {s.get('site', 'site')}", f"risk={s.get('risk')}"))

    os.makedirs(paths.output_dir(), exist_ok=True)
    out = os.path.join(paths.output_dir(), f"{program_no}_PSA_generated.docx")
    doc.save(out)

    print(f"Saved {out}\n--- FILLED (Smartsheet/derived) ---")
    for lbl, v in filled:
        print(f"  {lbl}: {v}")
    print("--- DEFERRED (typed placeholders) ---")
    for lbl, v in deferred:
        print(f"  {lbl}: {v}")


DEFAULT_PROGRAM = "AGN-151586"


def available_products():
    """Distinct selectable products (those with a program number)."""
    con = sqlite3.connect(paths.db_path())
    rows = con.execute(
        "SELECT program_no, MAX(program_name) FROM product "
        "WHERE program_no IS NOT NULL AND TRIM(program_no) <> '' "
        "GROUP BY program_no ORDER BY program_no").fetchall()
    con.close()
    return rows


def resolve_program(requested=None):
    """Decide which product to generate. If `requested` is a valid code, use it;
    otherwise show a numbered menu and let the user pick (blank = BoNT/E default)."""
    products = available_products()
    codes = [r[0] for r in products]

    if requested:
        if requested in codes:
            return requested
        print(f"Product {requested!r} is not in the database — choose from the list below.")

    print("\nAvailable products:")
    for i, (code, name) in enumerate(products, 1):
        marker = "  (default)" if code == DEFAULT_PROGRAM else ""
        print(f"  {i:2}. {code:<16} {name or ''}{marker}")
    while True:
        sel = input(f"\nEnter a number or product code "
                    f"(blank = {DEFAULT_PROGRAM}): ").strip()
        if not sel:
            return DEFAULT_PROGRAM
        if sel in codes:
            return sel
        if sel.isdigit() and 1 <= int(sel) <= len(products):
            return products[int(sel) - 1][0]
        print("  Not a valid choice — try again.")


def main():
    import sys
    try:
        sys.stdout.reconfigure(encoding="utf-8")   # avoid cp1252 crashes printing the ☒/☐ risk-box glyphs
    except Exception:
        pass
    requested = sys.argv[1] if len(sys.argv) > 1 else None
    populate(resolve_program(requested))


if __name__ == "__main__":
    main()
