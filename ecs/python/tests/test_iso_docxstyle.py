"""The DOCX table primitives (`silos/iso/docxstyle.py`).

This module had no tests at all, which is awkward for the one place in the ISO port where
every constant is load-bearing: the column widths have to total the printable width of A4,
and the two shades are what tell a header row apart from a section divider. A drifted
twip here produces a report that still opens cleanly and still looks wrong.

Everything is exercised against real `python-docx` objects rather than mocks, because the
whole purpose of these functions is the OOXML they emit — a mock would assert that we
called a setter, not that the document says `w:tblLayout type="fixed"`. Reading the XML
back is the only assertion that would survive python-docx changing its API.

Two behaviours are worth naming because they are deliberate and easy to "fix" into a bug:
`_set_picture_alt_text` and `_add_image_to_cell` both swallow exceptions on purpose, so a
single unrenderable image degrades to a note in the cell instead of losing the report.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage

from api.backend.da_platform.settings import BACKEND_ROOT
from api.backend.da_platform.silo_registry import _load_module

ISO_DIR = BACKEND_ROOT / "silos" / "iso"

pytestmark = pytest.mark.skipif(
    not (ISO_DIR / "docxstyle.py").is_file(), reason="the ISO silo is not present"
)

EMU_PER_INCH = 914400

# The DPI the capture path uses, mirrored from the module. A PNG this many pixels wide is
# one inch wide in the finished document, which is what makes the scaling assertions
# readable instead of arithmetic.
CAP_DPI = 200


@pytest.fixture(scope="module")
def docxstyle():
    _load_module("iso", ISO_DIR / "silo.py")
    return _load_module("iso", ISO_DIR / "docxstyle.py", name="docxstyle")


@pytest.fixture
def cell():
    """A single real table cell, which is all these helpers ever operate on."""
    return Document().add_table(rows=1, cols=1).cell(0, 0)


def png_bytes(width_px: int, height_px: int) -> bytes:
    """A PNG of an exact pixel size, so the DPI-based scaling has something to divide."""
    buffer = io.BytesIO()
    PILImage.new("RGB", (width_px, height_px), color=(200, 30, 30)).save(
        buffer, format="PNG"
    )
    return buffer.getvalue()


def picture_width_emu(run) -> int:
    """The width Word will actually render, read back out of the drawing extent."""
    extents = run._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    )
    assert extents, "the run holds no picture"
    return int(extents[-1].get("cx"))


def added_picture_run(docxstyle, **kwargs):
    """Put one image in a fresh cell and hand back the cell plus the run holding it."""
    target = Document().add_table(rows=1, cols=1).cell(0, 0)
    docxstyle._add_image_to_cell(target, **kwargs)
    runs = [run for para in target.paragraphs for run in para.runs]
    return target, runs


# ── the constants ────────────────────────────────────────────────────────────


def test_the_two_column_widths_total_the_printable_width_of_a4(docxstyle):
    """2.54 cm margins on A4 leave ~6.27", and the table is sized to fill exactly that.

    Asserted because the module docstring calls it non-adjustable: a wider total silently
    pushes the description column off the page in print.
    """
    total_inches = sum(docxstyle._COL_W) / 1440
    assert total_inches == pytest.approx(6.27, abs=0.01)
    # Section number narrow, description wide -- swapping them would still total 6.27".
    assert docxstyle._COL_W[0] < docxstyle._COL_W[1]


def test_the_header_and_divider_shades_stay_distinguishable(docxstyle):
    """Both greys, but a reviewer has to be able to tell the two row kinds apart."""
    assert docxstyle._HDR_SHADE != docxstyle._SEC_SHADE
    assert docxstyle._HDR_SHADE == "C0C0C0"
    assert docxstyle._SEC_SHADE == "D9D9D9"


# ── _set_picture_alt_text ────────────────────────────────────────────────────


def test_alt_text_is_written_to_the_pictures_descr_attribute(docxstyle):
    """Formula images carry their LaTeX here, which is what makes them searchable."""
    _, runs = added_picture_run(
        docxstyle, img_bytes=png_bytes(200, 200), alt_text=r"\frac{a}{b}"
    )
    doc_prs = runs[0]._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    )
    assert doc_prs[-1].get("descr") == r"\frac{a}{b}"
    assert doc_prs[-1].get("title") == r"\frac{a}{b}"


def test_a_long_latex_source_is_truncated_in_the_title_but_kept_whole_in_descr(docxstyle):
    """`title` is a short label in Word's UI; the real alt text must not be lossy."""
    long_latex = "x" * 200
    _, runs = added_picture_run(
        docxstyle, img_bytes=png_bytes(200, 200), alt_text=long_latex
    )
    doc_prs = runs[0]._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    )
    assert doc_prs[-1].get("descr") == long_latex
    assert doc_prs[-1].get("title") == "x" * 120 + "…"


def test_empty_alt_text_leaves_the_picture_untouched(docxstyle):
    """The caller passes '' for ordinary figures, which must not stamp an empty descr."""
    _, runs = added_picture_run(docxstyle, img_bytes=png_bytes(200, 200), alt_text="")
    doc_prs = runs[0]._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    )
    assert doc_prs[-1].get("descr") is None


def test_setting_alt_text_on_a_run_with_no_picture_is_a_no_op(docxstyle):
    """Reached when an image failed to render; it must not raise on top of that."""
    run = Document().add_paragraph().add_run("just words")
    docxstyle._set_picture_alt_text(run, "anything")  # must not raise


def test_an_empty_alt_text_returns_before_touching_the_element(docxstyle):
    """Called directly because `_add_image_to_cell` guards this branch for itself.

    The guard is what lets every caller pass `alt_text` unconditionally instead of each
    one repeating the check.
    """

    class Exploding:
        @property
        def _element(self):
            raise AssertionError("an empty alt text must not inspect the run")

    docxstyle._set_picture_alt_text(Exploding(), "")
    docxstyle._set_picture_alt_text(Exploding(), None)


def test_a_picture_whose_docPr_is_missing_is_left_alone(docxstyle):
    """A drawing with no `docPr` has nowhere to put alt text, so it is skipped."""
    _, runs = added_picture_run(docxstyle, img_bytes=png_bytes(200, 200))
    run = runs[0]
    WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    inline = run._element.findall(f".//{{{WP}}}inline")[-1]
    doc_pr = inline.find(f"{{{WP}}}docPr")
    inline.remove(doc_pr)

    docxstyle._set_picture_alt_text(run, "some latex")  # must not raise

    assert inline.find(f"{{{WP}}}docPr") is None


def test_an_unexpected_element_failure_is_swallowed(docxstyle):
    """Deliberately broad: alt text is a nicety, and losing it must not lose the report."""

    class Broken:
        class _element:
            @staticmethod
            def findall(_path):
                raise RuntimeError("lxml said no")

    docxstyle._set_picture_alt_text(Broken(), "latex")  # must not raise


def test_control_characters_are_stripped_from_caption_and_alt_text(docxstyle):
    """Both can come from the vision model, and a raw \\x0b crashes the serializer.

    Saving the document is the real assertion: a control character that survives here
    raises only at write time, which is far from the code that let it through.
    """
    target, runs = added_picture_run(
        docxstyle,
        img_bytes=png_bytes(200, 200),
        caption="Figure\x0b1 \x07caption",
        alt_text="alt\x0btext",
    )
    doc_prs = runs[0]._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    )
    assert doc_prs[-1].get("descr") == "alttext"
    assert "\x0b" not in target.text and "\x07" not in target.text
    assert "Figure1 caption" in target.text


# ── _add_image_to_cell scaling ───────────────────────────────────────────────


def test_an_image_narrower_than_the_column_keeps_its_natural_size(docxstyle):
    """400 px at 200 DPI is 2", which fits the 4.5" default, so nothing is scaled."""
    _, runs = added_picture_run(docxstyle, img_bytes=png_bytes(400, 400))
    assert picture_width_emu(runs[0]) == pytest.approx(2 * EMU_PER_INCH, rel=0.01)


def test_an_image_wider_than_the_column_is_scaled_down_to_fit_it(docxstyle):
    """2000 px is 10" at capture DPI; the column allows 4.5", so it must come down."""
    _, runs = added_picture_run(docxstyle, img_bytes=png_bytes(2000, 400))
    assert picture_width_emu(runs[0]) == pytest.approx(4.5 * EMU_PER_INCH, rel=0.01)


def test_a_narrow_but_very_tall_image_is_clamped_by_page_height_not_column_width(
    docxstyle,
):
    """The second clamp exists so a full-page capture is not taller than the page.

    200x2000 px is 1" wide and 10" tall. Width already fits, so only the 6.5" height cap
    applies, and the width has to shrink with it or the aspect ratio would break.
    """
    _, runs = added_picture_run(docxstyle, img_bytes=png_bytes(200, 2000))
    expected = 6.5 / 10 * EMU_PER_INCH
    assert picture_width_emu(runs[0]) == pytest.approx(expected, rel=0.02)


def test_max_width_is_honoured_when_the_caller_narrows_the_column(docxstyle):
    """Nested cells pass a smaller max_w_in, which has to win over the 4.5" default."""
    _, runs = added_picture_run(
        docxstyle, img_bytes=png_bytes(2000, 400), max_w_in=2.0
    )
    assert picture_width_emu(runs[0]) == pytest.approx(2.0 * EMU_PER_INCH, rel=0.01)


# ── _add_image_to_cell captions and failure ──────────────────────────────────


def test_a_caption_is_added_centred_italic_and_small(docxstyle):
    target, _ = added_picture_run(
        docxstyle, img_bytes=png_bytes(200, 200), caption="Figure 1 — Overview"
    )
    caption_paragraph = target.paragraphs[-1]
    assert caption_paragraph.text == "Figure 1 — Overview"
    assert caption_paragraph.runs[0].font.italic is True
    assert caption_paragraph.runs[0].font.size == Pt(8)


def test_no_caption_means_no_extra_paragraph(docxstyle):
    """Most figures have no caption, and an empty paragraph would show as a blank line."""
    with_caption, _ = added_picture_run(
        docxstyle, img_bytes=png_bytes(200, 200), caption="Figure 1"
    )
    without_caption, _ = added_picture_run(docxstyle, img_bytes=png_bytes(200, 200))
    assert len(without_caption.paragraphs) == len(with_caption.paragraphs) - 1


def test_an_unreadable_image_leaves_a_note_in_the_cell_instead_of_raising(docxstyle):
    """One bad image must not lose the whole report -- and a reviewer needs to find it."""
    target = Document().add_table(rows=1, cols=1).cell(0, 0)
    docxstyle._add_image_to_cell(target, b"this is not a png", caption="Figure 4")
    assert "rendering failed" in target.text
    assert "Figure 4" in target.text


def test_an_unreadable_image_without_a_caption_still_says_which_cell_failed(docxstyle):
    target = Document().add_table(rows=1, cols=1).cell(0, 0)
    docxstyle._add_image_to_cell(target, b"nonsense")
    assert "Image" in target.text and "rendering failed" in target.text


# ── cell and row formatting ──────────────────────────────────────────────────


def test_a_bordered_cell_gets_all_four_single_edges(docxstyle, cell):
    """Three-sided borders are the classic symptom of a loop bound edited by hand."""
    docxstyle._cell_border(cell, color="FF0000", sz="12")
    borders = cell._tc.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    sides = [child.tag.split("}")[-1] for child in borders]
    assert sides == ["top", "left", "bottom", "right"]
    for child in borders:
        assert child.get(qn("w:val")) == "single"
        assert child.get(qn("w:color")) == "FF0000"
        assert child.get(qn("w:sz")) == "12"


def test_shading_a_cell_sets_a_clear_pattern_over_the_requested_fill(docxstyle, cell):
    """`val=clear` with an explicit fill is what Word reads as a solid background."""
    docxstyle._cell_shade(cell, docxstyle._HDR_SHADE)
    shade = cell._tc.find(qn("w:tcPr")).find(qn("w:shd"))
    assert shade.get(qn("w:fill")) == "C0C0C0"
    assert shade.get(qn("w:val")) == "clear"


def test_cell_width_is_written_in_twips_as_an_absolute_measurement(docxstyle, cell):
    """`type=dxa` is the difference between a fixed width and a percentage."""
    docxstyle._cell_width(cell, docxstyle._COL_W[1])
    width = cell._tc.find(qn("w:tcPr")).findall(qn("w:tcW"))[-1]
    assert width.get(qn("w:w")) == str(docxstyle._COL_W[1])
    assert width.get(qn("w:type")) == "dxa"


def test_a_float_width_is_coerced_to_an_integer_twip(docxstyle, cell):
    """OOXML rejects a decimal here, and the arithmetic upstream produces floats."""
    docxstyle._cell_width(cell, 1234.7)
    width = cell._tc.find(qn("w:tcPr")).findall(qn("w:tcW"))[-1]
    assert width.get(qn("w:w")) == "1234"


def test_setting_a_width_appends_a_second_tcW_rather_than_replacing_the_first(
    docxstyle, cell
):
    """Characterisation, not endorsement: this is a real defect, pinned so it is visible.

    python-docx already writes a `w:tcW` when the table is created, and `_cell_width`
    appends instead of replacing, so the cell ends up with two. `CT_TcPr` allows at most
    one, and a lenient reader takes the first -- which is python-docx's autofit value, not
    the width this module was asked for. The report still looks right today only because
    `_apply_fixed_cols` sets the widths a second way, through `tblGrid` with a fixed
    layout. Change this to replace rather than append and the assertion below should be
    inverted.
    """
    docxstyle._cell_width(cell, 2160)
    widths = cell._tc.find(qn("w:tcPr")).findall(qn("w:tcW"))
    assert len(widths) == 2
    assert widths[-1].get(qn("w:w")) == "2160"
    assert widths[0].get(qn("w:w")) != "2160"


def test_shading_a_cell_twice_stacks_two_shd_elements(docxstyle, cell):
    """Same defect shape as the width one, and the reason it has not bitten: nothing in
    the ISO path shades a cell twice. Pinned so that if something starts to, the
    duplicate is a test failure rather than a subtly wrong colour.
    """
    docxstyle._cell_shade(cell, docxstyle._HDR_SHADE)
    docxstyle._cell_shade(cell, docxstyle._SEC_SHADE)
    shades = cell._tc.find(qn("w:tcPr")).findall(qn("w:shd"))
    assert len(shades) == 2
    assert shades[0].get(qn("w:fill")) == docxstyle._HDR_SHADE


def test_a_row_minimum_height_is_a_floor_not_a_fixed_height(docxstyle):
    """`hRule=atLeast` lets a tall image grow the row; `exactly` would crop it."""
    row = Document().add_table(rows=1, cols=1).rows[0]
    docxstyle._row_min_height(row, 500)
    height = row._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
    assert height.get(qn("w:val")) == "500"
    assert height.get(qn("w:hRule")) == "atLeast"


def test_cell_font_applies_to_every_run_in_every_paragraph(docxstyle, cell):
    """A cell holds one run per emphasis change, so styling only the first is a bug."""
    cell.paragraphs[0].add_run("first ")
    cell.paragraphs[0].add_run("second")
    cell.add_paragraph().add_run("third")

    docxstyle._cell_font(cell, bold=True, italic=True, size=Pt(11))

    runs = [run for para in cell.paragraphs for run in para.runs]
    assert len(runs) == 3
    for run in runs:
        assert run.font.bold is True
        assert run.font.italic is True
        assert run.font.size == Pt(11)


def test_cell_font_defaults_to_nine_point_and_not_bold(docxstyle, cell):
    cell.paragraphs[0].add_run("body text")
    docxstyle._cell_font(cell)
    run = cell.paragraphs[0].runs[0]
    assert run.font.size == docxstyle._FONT_SIZE == Pt(9)
    assert run.font.bold is False
    assert run.font.italic is False


# ── _apply_fixed_cols ────────────────────────────────────────────────────────


def test_fixed_layout_and_a_grid_are_stamped_onto_a_real_table(docxstyle):
    """Without `tblLayout=fixed` Word re-flows the columns and ignores the widths."""
    table = Document().add_table(rows=1, cols=2)
    docxstyle._apply_fixed_cols(table._tbl, docxstyle._COL_W)

    layout = table._tbl.find(qn("w:tblPr")).find(qn("w:tblLayout"))
    assert layout.get(qn("w:type")) == "fixed"

    grid = table._tbl.findall(qn("w:tblGrid"))[0]
    assert [col.get(qn("w:w")) for col in grid] == [str(w) for w in docxstyle._COL_W]


def test_a_table_with_no_properties_element_gets_one_created(docxstyle):
    """python-docx supplies tblPr, but a hand-built or round-tripped tbl may not."""
    bare_table = OxmlElement("w:tbl")
    docxstyle._apply_fixed_cols(bare_table, [1440, 2880])

    tblPr = bare_table.find(qn("w:tblPr"))
    assert tblPr is not None
    assert tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    # tblPr first, then the grid: Word rejects the reverse order.
    assert list(bare_table)[0] is tblPr
    assert list(bare_table)[1].tag == qn("w:tblGrid")


def test_grid_column_widths_are_coerced_to_integers(docxstyle):
    """The callers compute widths by ratio, so floats reach this function."""
    bare_table = OxmlElement("w:tbl")
    docxstyle._apply_fixed_cols(bare_table, [1440.9, 2880.2])
    grid = bare_table.find(qn("w:tblGrid"))
    assert [col.get(qn("w:w")) for col in grid] == ["1440", "2880"]


# ── the whole module, end to end ─────────────────────────────────────────────


def test_a_fully_styled_table_still_saves_as_a_valid_document(docxstyle):
    """Each helper appends to the same tcPr, so together they can emit invalid OOXML.

    Saving is the only check that catches an ordering or namespace mistake; every
    individual assertion above would still pass.
    """
    document = Document()
    table = document.add_table(rows=2, cols=2)
    docxstyle._apply_fixed_cols(table._tbl, docxstyle._COL_W)
    docxstyle._row_min_height(table.rows[0], 400)

    for index, cell_ in enumerate(table.rows[0].cells):
        cell_.text = f"Header {index}"
        docxstyle._cell_border(cell_)
        docxstyle._cell_shade(cell_, docxstyle._HDR_SHADE)
        docxstyle._cell_width(cell_, docxstyle._COL_W[index])
        docxstyle._cell_font(cell_, bold=True)

    docxstyle._add_image_to_cell(
        table.rows[1].cells[1],
        png_bytes(600, 400),
        caption="Figure 1 — Overview",
        alt_text=r"E = mc^2",
    )

    buffer = io.BytesIO()
    document.save(buffer)
    assert buffer.tell() > 0

    reopened = Document(io.BytesIO(buffer.getvalue()))
    assert reopened.tables[0].rows[0].cells[0].text == "Header 0"
    assert "Figure 1 — Overview" in reopened.tables[0].rows[1].cells[1].text
    assert len(reopened.inline_shapes) == 1
