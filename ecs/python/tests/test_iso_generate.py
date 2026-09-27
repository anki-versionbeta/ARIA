"""Assembling the assessment DOCX (`silos/iso/generate.py`).

This module is the last thing that runs before bytes leave the silo, and it was at 0%
coverage. Nothing downstream of it is tested either, so a mistake here — the wrong title
in the wrong place, a lost image, a filename built from the wrong string — reaches the user
unchallenged.

Three things are worth stating about how these tests are built.

**The two titles are asserted separately, and against different sources.** The port's own
docstring warns that the title printed *inside* the document comes from the PDF's metadata
while the RAG-derived title names the *file* only, and that conflating them is exactly the
silent drift the port exists to prevent. A test that only checked "a title appears" would
pass with them swapped, so each is pinned to its own origin.

**`prescan_media` and `extract_all_rows` are replaced, not driven.** Both are separately
tested, both rasterise, and both are slow. Replacing them turns the arity ladder in
`generate_docx` into something that can actually be exercised at all six of its lengths —
which matters, because the source itself flags that ladder as a hazard that would absorb an
arity mistake into a subtly wrong document instead of raising.

**The output is read back out of the DOCX, not out of a mock.** `generate_docx` returns
bytes; every structural assertion re-opens those bytes with `python-docx` and inspects the
real table, so the shading, the widths and the `keepWithNext` are checked as Word will see
them.
"""

from __future__ import annotations

import io

import fitz
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from PIL import Image as PILImage

from api.backend.da_platform.settings import BACKEND_ROOT
from api.backend.da_platform.silo_registry import _load_module
from tests.iso_fakes import FakeLlm

ISO_DIR = BACKEND_ROOT / "silos" / "iso"

pytestmark = pytest.mark.skipif(
    not (ISO_DIR / "generate.py").is_file(), reason="the ISO silo is not present"
)


@pytest.fixture(scope="module")
def gen():
    _load_module("iso", ISO_DIR / "silo.py")
    return _load_module("iso", ISO_DIR / "generate.py", name="generate")


# ── test doubles ─────────────────────────────────────────────────────────────


class FakeWorkspace:
    """`generate_docx` reads exactly one attribute off the workspace."""

    def __init__(self, pdf_path):
        self.pdf_path = str(pdf_path)


class FakeCtx:
    """The stage context: a progress callback and the two capabilities."""

    def __init__(self, textract="TEXTRACT", llm=None):
        self.textract = textract
        self.llm = llm if llm is not None else FakeLlm()
        self.progress_calls = []

    def progress(self, message, pct=None):
        self.progress_calls.append((message, pct))


def build_pdf(tmp_path, *, name="ISO 10943-2023.pdf", title=None, pages=1):
    """A real PDF, because `generate_docx` opens it with fitz to read its metadata.

    The metadata title is the value the document's own heading is built from, so it has to
    come from a genuine PDF trailer rather than a stub.
    """
    doc = fitz.open()
    for _ in range(pages):
        doc.new_page(width=595, height=842)
    if title is not None:
        doc.set_metadata({"title": title})
    path = tmp_path / name
    doc.save(str(path))
    doc.close()
    return path


def png_bytes(width_px=400, height_px=200):
    """A real PNG, since `_add_image_to_cell` opens it with PIL to size it."""
    buffer = io.BytesIO()
    PILImage.new("RGB", (width_px, height_px), color=(10, 90, 200)).save(
        buffer, format="PNG"
    )
    return buffer.getvalue()


TOC = [
    {"title": "7 Requirements", "page": 1, "level": 1},
    {"title": "7.1 General", "page": 2, "level": 2},
    {"title": "8.2 Marking", "page": 3, "level": 2},
]


def content_row(section_num="7.1", text="Prose.", images=None):
    return {
        "type": "content",
        "section_num": section_num,
        "text": text,
        "level": 2,
        "inline_images": images or [],
    }


def header_row(section_num="7.1", text="7.1 General"):
    return {
        "type": "header",
        "section_num": section_num,
        "text": text,
        "level": 2,
        "inline_images": [],
    }


@pytest.fixture
def stubbed(gen, monkeypatch):
    """Replace the two heavy collaborators and record what they were handed.

    Returns the recorder dict; tests mutate `recorder["prescan"]` to change the arity of
    the pre-scan result and `recorder["rows"]` to change the rows that get rendered.
    """
    recorder = {
        "prescan": ({}, {}, {}, [], {}, {}, {}, {}),
        "rows": [],
        "prescan_kwargs": None,
        "prescan_args": None,
        "rows_kwargs": None,
        "rows_args": None,
    }

    def fake_prescan(*args, **kwargs):
        recorder["prescan_args"] = args
        recorder["prescan_kwargs"] = kwargs
        return recorder["prescan"]

    def fake_rows(*args, **kwargs):
        recorder["rows_args"] = args
        recorder["rows_kwargs"] = kwargs
        return recorder["rows"]

    monkeypatch.setattr(gen, "prescan_media", fake_prescan)
    monkeypatch.setattr(gen, "extract_all_rows", fake_rows)
    return recorder


def rendered(gen, tmp_path, stubbed, **kwargs):
    """Build the DOCX and hand back the parsed document plus its single table."""
    ws = FakeWorkspace(kwargs.pop("pdf_path", None) or build_pdf(tmp_path))
    payload = gen.generate_docx(
        ws, kwargs.pop("toc", TOC), kwargs.pop("start_idx", 0),
        kwargs.pop("end_idx", 2), **kwargs
    )
    assert isinstance(payload, bytes)
    doc = Document(io.BytesIO(payload))
    return doc, doc.tables[0]


def all_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


# ── the two titles ───────────────────────────────────────────────────────────


def test_the_title_printed_inside_the_document_comes_from_the_pdf_metadata(
    gen, tmp_path, stubbed
):
    """Not from the RAG-derived title, which names the file. The source calls this out
    explicitly, so it is asserted explicitly."""
    pdf = build_pdf(tmp_path, name="whatever.pdf", title="ISO 10943:2023 Ophthalmic")
    doc, _tbl = rendered(gen, tmp_path, stubbed, pdf_path=pdf)
    assert "ISO 10943:2023 Ophthalmic" in all_text(doc)


def test_the_document_title_falls_back_to_the_filename_without_its_extension(
    gen, tmp_path, stubbed
):
    """Most ISO PDFs ship with an empty metadata title, so the fallback is the live path."""
    pdf = build_pdf(tmp_path, name="ISO 10943-2023.pdf", title="")
    doc, _tbl = rendered(gen, tmp_path, stubbed, pdf_path=pdf)
    assert "ISO 10943-2023" in all_text(doc)


def test_a_whitespace_only_metadata_title_is_treated_as_absent(gen, tmp_path, stubbed):
    pdf = build_pdf(tmp_path, name="fallback.pdf", title="   ")
    doc, _tbl = rendered(gen, tmp_path, stubbed, pdf_path=pdf)
    assert "fallback" in all_text(doc)


def test_a_control_byte_in_the_pdf_title_aborts_the_whole_report(gen, tmp_path, stubbed):
    """Characterisation of a live defect (generate.py:184, and the same at 196).

    `doc_title` is read straight from the PDF metadata and is passed through `_xml_safe`
    at generate.py:158 for the title run — but the Purpose/Overview paragraph and the
    "Assessment of ..." heading interpolate the *raw* value. Real ISO PDFs do carry stray
    C0 bytes (the module's own comment cites 0x08 in ISO 13485:2016), and lxml rejects
    them, so a single invisible byte in the metadata takes down the entire build instead
    of being stripped. Pinned here rather than fixed; the fix is one `_xml_safe` call at
    each of the two interpolation sites.
    """
    pdf = build_pdf(tmp_path, name="ctrl.pdf", title="ISO\x08 10943")
    with pytest.raises(ValueError, match="XML compatible"):
        rendered(gen, tmp_path, stubbed, pdf_path=pdf)


def test_the_title_run_itself_does_survive_a_control_byte(gen, tmp_path, stubbed):
    """The one call site that is guarded. Kept alongside the failure above so the
    asymmetry is on the record rather than looking like an oversight in the test."""
    safe = gen._xml_safe("ISO\x08 10943")
    assert safe == "ISO 10943"


# ── the fixed front matter ───────────────────────────────────────────────────


def test_the_front_matter_carries_the_report_subtitle_and_both_guidance_notes(
    gen, tmp_path, stubbed
):
    """These strings are the deliverable's identity; a port that drops one is wrong even
    though nothing would raise."""
    doc, _tbl = rendered(gen, tmp_path, stubbed)
    text = all_text(doc)
    assert "Applicability Assessment Report" in text
    assert "Project Information" in text
    assert "<< Add a description of the specific project here. >>" in text
    assert "Purpose / Overview" in text
    assert "NOTE: Where differences exist between British English" in text
    assert "— END OF DOCUMENT —" in text


def test_the_purpose_paragraph_names_the_first_and_last_selected_section_numbers(
    gen, tmp_path, stubbed
):
    """The numbers come from the ToC entries at the selected indices, not from the rows,
    so an off-by-one in the slice would be invisible in the table but visible here."""
    doc, _tbl = rendered(gen, tmp_path, stubbed, start_idx=1, end_idx=2)
    assert "sections 7.1 through 8.2" in all_text(doc)


def test_the_guidance_placeholder_is_blue_italic_so_it_is_obvious_it_must_be_removed(
    gen, tmp_path, stubbed
):
    doc, _tbl = rendered(gen, tmp_path, stubbed)
    placeholder = next(
        p for p in doc.paragraphs if p.text.startswith("<< Add a description")
    )
    assert placeholder.runs[0].font.italic is True
    assert placeholder.runs[0].font.color.rgb == RGBColor(0x00, 0x70, 0xC0)


def test_a_template_with_no_title_style_does_not_stop_the_report(gen, tmp_path, stubbed,
                                                                 monkeypatch):
    """The `try/except pass` around the style assignment. python-docx's default template
    always has 'Title', but the styles collection is template-dependent, and losing the
    styling of one paragraph is not worth losing the report over — so the swallow is
    correct here even though a bare `except Exception` usually is not.
    """
    real_document = gen.Document

    def document_without_a_title_style():
        doc = real_document()
        styles = doc.styles.element
        for element in styles.findall(qn("w:style")):
            if element.get(qn("w:styleId")) == "Title":
                styles.remove(element)
        return doc

    monkeypatch.setattr(gen, "Document", document_without_a_title_style)
    doc, _tbl = rendered(gen, tmp_path, stubbed)
    assert "Applicability Assessment Report" in all_text(doc)


def test_the_page_is_set_up_as_a4_with_one_inch_margins(gen, tmp_path, stubbed):
    """The report is circulated in Europe; Word's default Letter size would reflow it."""
    doc, _tbl = rendered(gen, tmp_path, stubbed)
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 2) == 2.54


# ── the assessment table ─────────────────────────────────────────────────────


def test_the_table_opens_with_a_shaded_bold_two_column_header(gen, tmp_path, stubbed):
    doc, tbl = rendered(gen, tmp_path, stubbed)
    assert len(tbl.columns) == 2
    header = tbl.rows[0]
    assert [c.text for c in header.cells] == ["Section\nNumber", "Description"]
    # `cell.text = ''` leaves an empty run behind, so the labelled run is the last one.
    assert all(c.paragraphs[0].runs[-1].font.bold for c in header.cells)
    shade = header.cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd"))
    assert shade is not None, "the header row is unshaded"


def test_a_content_row_puts_the_section_number_beside_its_text(gen, tmp_path, stubbed):
    stubbed["rows"] = [content_row("7.1 a)", "The device shall be marked.")]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    assert [c.text for c in tbl.rows[1].cells] == [
        "7.1 a)",
        "The device shall be marked.",
    ]


def test_a_header_row_is_shaded_differently_from_a_content_row(gen, tmp_path, stubbed):
    """The section divider has to read as a divider, which is the only thing separating
    one clause's rows from the next in a table hundreds of rows long."""
    stubbed["rows"] = [header_row(), content_row()]
    _doc, tbl = rendered(gen, tmp_path, stubbed)

    def fill(cell):
        pr = cell._tc.find(qn("w:tcPr"))
        shd = None if pr is None else pr.find(qn("w:shd"))
        return None if shd is None else shd.get(qn("w:fill"))

    assert fill(tbl.rows[1].cells[0]) is not None
    assert fill(tbl.rows[2].cells[0]) is None
    assert tbl.rows[1].cells[1].paragraphs[0].runs[-1].font.bold is True


def test_a_reference_in_the_description_is_highlighted(gen, tmp_path, stubbed):
    """The highlight is how a reviewer finds the figure the sentence is talking about."""
    stubbed["rows"] = [content_row(text="as shown in Figure 4 of this clause")]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    highlighted = [
        r.text for r in tbl.rows[1].cells[1].paragraphs[0].runs
        if r.font.highlight_color is not None
    ]
    assert highlighted == ["Figure 4"]


def test_an_inline_image_is_placed_in_the_description_cell(gen, tmp_path, stubbed):
    stubbed["rows"] = [
        content_row(
            text="Figure 4 — Test rig",
            images=[{"type": "figure", "data": png_bytes(), "caption": ""}],
        )
    ]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    blips = tbl.rows[1].cells[1]._tc.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )
    assert len(blips) == 1


def test_a_row_with_an_image_keeps_its_text_on_the_same_page_as_the_image(
    gen, tmp_path, stubbed
):
    """Without `keepWithNext` Word splits a caption from its figure across a page break,
    which reads as a figure with no caption and a caption with no figure."""
    stubbed["rows"] = [
        content_row(images=[{"type": "table", "data": png_bytes(), "caption": ""}])
    ]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    para = tbl.rows[1].cells[1].paragraphs[0]
    assert para._p.find(qn("w:pPr")).find(qn("w:keepWithNext")) is not None


def test_a_row_with_no_images_is_not_given_a_keep_with_next(gen, tmp_path, stubbed):
    """Applied to every row it would defeat pagination entirely."""
    stubbed["rows"] = [content_row()]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    pPr = tbl.rows[1].cells[1].paragraphs[0]._p.find(qn("w:pPr"))
    assert pPr is None or pPr.find(qn("w:keepWithNext")) is None


def test_an_image_entry_carrying_no_data_is_skipped_rather_than_crashing(
    gen, tmp_path, stubbed
):
    """`_build_row_images` can emit a reference-only entry, and one missing PNG must not
    abort a report that is otherwise complete."""
    stubbed["rows"] = [
        content_row(images=[{"type": "figure", "number": 9, "caption": "Figure 9"}])
    ]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    assert tbl.rows[1].cells[1].text == "Prose."


def test_a_formula_image_carries_its_latex_through_as_the_alt_text(gen, tmp_path, stubbed):
    """The LaTeX is the only searchable form of the equation once it is a raster."""
    stubbed["rows"] = [
        content_row(
            images=[{"type": "formula", "data": png_bytes(), "caption": "(7)",
                     "latex": "x = \\mu + k s"}]
        )
    ]
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    descriptions = [
        el.get("descr")
        for el in tbl.rows[1].cells[1]._tc.iter()
        if el.get("descr") is not None
    ]
    assert "x = \\mu + k s" in descriptions


def test_an_empty_row_list_still_produces_a_valid_document_with_just_the_header(
    gen, tmp_path, stubbed
):
    """A selection that matched no content must not produce a corrupt file."""
    stubbed["rows"] = []
    _doc, tbl = rendered(gen, tmp_path, stubbed)
    assert len(tbl.rows) == 1


# ── what gets handed to the collaborators ────────────────────────────────────


def test_the_prescan_is_given_the_pdf_path_the_toc_and_the_selected_index_bounds(
    gen, tmp_path, stubbed
):
    """The bounds are what stop the formula detector rasterising the whole standard."""
    pdf = build_pdf(tmp_path)
    rendered(gen, tmp_path, stubbed, pdf_path=pdf, start_idx=1, end_idx=2)
    assert stubbed["prescan_args"] == (str(pdf),)
    kwargs = stubbed["prescan_kwargs"]
    assert kwargs["toc"] is TOC
    assert (kwargs["toc_start_idx"], kwargs["toc_end_idx"]) == (1, 2)


def test_the_vision_page_texts_are_keyed_by_page_index_for_the_row_extractor(
    gen, tmp_path, stubbed
):
    llm_pages = [{"page_idx": 4, "text": "Page five text."},
                 {"page_idx": 5, "text": "Page six text."}]
    rendered(gen, tmp_path, stubbed, llm_pages=llm_pages)
    assert stubbed["rows_kwargs"]["llm_texts"] == {
        4: "Page five text.", 5: "Page six text."
    }


def test_a_vision_page_that_produced_no_text_is_left_out_of_the_lookup(
    gen, tmp_path, stubbed
):
    """A page whose vision call failed must fall back to fitz, and it only does that when
    the page is absent from the map — an empty string would suppress the whole page."""
    llm_pages = [None, {"page_idx": 1, "text": ""}, {"page_idx": 2},
                 {"page_idx": 3, "text": "Kept."}]
    rendered(gen, tmp_path, stubbed, llm_pages=llm_pages)
    assert stubbed["rows_kwargs"]["llm_texts"] == {3: "Kept."}


def test_without_vision_pages_the_row_extractor_is_told_there_is_no_llm_text_at_all(
    gen, tmp_path, stubbed
):
    """`None` and `{}` are different to `_extract_body`: only `None` means "every page
    uses fitz"."""
    rendered(gen, tmp_path, stubbed, llm_pages=[])
    assert stubbed["rows_kwargs"]["llm_texts"] is None


def test_the_capabilities_are_passed_through_to_the_prescan(gen, tmp_path, stubbed):
    """The prescan is the only place in the build that spends money."""
    llm = FakeLlm()
    rendered(gen, tmp_path, stubbed, textract="TEXTRACT", llm=llm)
    assert stubbed["prescan_kwargs"]["textract"] == "TEXTRACT"
    assert stubbed["prescan_kwargs"]["llm"] is llm


# ── the pre-scan arity ladder ────────────────────────────────────────────────


def test_the_full_eight_element_prescan_result_is_unpacked_in_order(
    gen, tmp_path, stubbed
):
    """The live shape. Ordering matters and is not self-describing: swapping the two rect
    maps would suppress body text under figures and leak table cells, with nothing raising.
    """
    stubbed["prescan"] = (
        {1: b"F"}, {"3": [b"T"]}, {2: b"Q"}, [(0, {})],
        {0: [(1, 2, 3, 4)]}, {1: [(5, 6, 7, 8)]}, {2: ["x1"]}, {"table:3": 9},
    )
    rendered(gen, tmp_path, stubbed)
    args = stubbed["rows_args"]
    kwargs = stubbed["rows_kwargs"]
    assert args[4:7] == ({1: b"F"}, {"3": [b"T"]}, {2: b"Q"})
    assert kwargs["unnumbered_formulas"] == [(0, {})]
    assert kwargs["table_rects_by_page"] == {0: [(1, 2, 3, 4)]}
    assert kwargs["fig_rects_by_page"] == {1: [(5, 6, 7, 8)]}
    assert kwargs["unlabeled_tables_by_page"] == {2: ["x1"]}
    assert kwargs["media_first_page"] == {"table:3": 9}


@pytest.mark.parametrize(
    "result, absent",
    [
        pytest.param(
            ({}, {}, {}, [], {"t": 1}, {"u": 1}, {"m": 1}),
            ["fig_rects_by_page"],
            id="seven-elements-predate-figure-rectangles",
        ),
        pytest.param(
            ({}, {}, {}, [], {"t": 1}, {"u": 1}),
            ["fig_rects_by_page", "media_first_page"],
            id="six-elements-predate-media-page-tracking",
        ),
        pytest.param(
            ({}, {}, {}, [], {"t": 1}),
            ["fig_rects_by_page", "media_first_page", "unlabeled_tables_by_page"],
            id="five-elements-predate-unlabeled-tables",
        ),
        pytest.param(
            ({}, {}, {}, []),
            ["fig_rects_by_page", "media_first_page", "unlabeled_tables_by_page",
             "table_rects_by_page"],
            id="four-elements-predate-table-rectangles",
        ),
    ],
)
def test_a_shorter_prescan_result_defaults_the_arguments_it_cannot_supply(
    gen, tmp_path, stubbed, result, absent
):
    """Characterisation of a hazard the source names itself (generate.py:88-91).

    Only lengths 3 and 8 are reachable against the current `prescan_media`; the middle
    rungs are dead code kept for the port's fidelity. The hazard is that they are *silent*:
    an arity mistake in a future `prescan_media` lands on one of these rungs and produces a
    document with no figure suppression and no page-range gating, rather than raising. These
    tests pin what each rung actually defaults so the blast radius is documented.
    """
    stubbed["prescan"] = result
    rendered(gen, tmp_path, stubbed)
    kwargs = stubbed["rows_kwargs"]
    for name in absent:
        assert kwargs[name] == {}, name
    if "table_rects_by_page" not in absent:
        assert kwargs["table_rects_by_page"] == {"t": 1}


def test_a_three_element_prescan_result_is_the_no_toc_legacy_shape(gen, tmp_path, stubbed):
    """The documented backward-compatible return for callers that pass no ToC."""
    stubbed["prescan"] = ({1: b"F"}, {}, {})
    rendered(gen, tmp_path, stubbed)
    kwargs = stubbed["rows_kwargs"]
    assert kwargs["unnumbered_formulas"] == []
    assert kwargs["table_rects_by_page"] == {}
    assert kwargs["fig_rects_by_page"] == {}
    assert kwargs["unlabeled_tables_by_page"] == {}
    assert kwargs["media_first_page"] == {}


# ── build_document: the stage-facing wrapper ─────────────────────────────────


@pytest.fixture
def wrapped(gen, monkeypatch):
    """Replace `generate_docx` so the wrapper's own logic is what is under test."""
    calls = {}

    def fake_generate_docx(ws, toc, start_idx, end_idx, llm_pages=None,
                           textract=None, llm=None):
        calls["ws"] = ws
        calls["toc"] = toc
        calls["idx"] = (start_idx, end_idx)
        calls["llm_pages"] = llm_pages
        calls["textract"] = textract
        calls["llm"] = llm
        return b"DOCXBYTES"

    monkeypatch.setattr(gen, "generate_docx", fake_generate_docx)
    return calls


def call_build(gen, *, doc_title="ISO 10943:2023", use_vision=False, ctx=None,
               total_pages=42):
    ctx = ctx or FakeCtx()
    return ctx, gen.build_document(
        ctx,
        FakeWorkspace("ignored.pdf"),
        toc=TOC,
        doc_title=doc_title,
        start_idx=0,
        end_idx=2,
        use_vision=use_vision,
        llm_pages=None,
        total_pages=total_pages,
    )


def test_the_wrapper_returns_the_docx_bytes_under_the_word_content_type(gen, wrapped):
    _ctx, result = call_build(gen)
    assert result["docx"] == b"DOCXBYTES"
    assert result["content_type"] == gen.DOCX_CONTENT_TYPE
    assert result["content_type"].endswith("wordprocessingml.document")


def test_the_wrapper_reports_no_editable_sections(gen, wrapped):
    """Deliberate: the app being replaced has no ISO section editor, so producing them
    would be an addition rather than a port."""
    _ctx, result = call_build(gen)
    assert result["sections"] is None


def test_the_output_filename_is_built_from_the_rag_derived_title(gen, wrapped):
    """This is the *other* title — the one that must not reach the document heading."""
    _ctx, result = call_build(gen, doc_title="ISO 10943:2023 Ophthalmic implants")
    assert result["filename"] == "ISO_109432023_Ophthalmic_implants_Assessment.docx"


def test_punctuation_is_stripped_from_the_filename_so_it_is_safe_on_any_filesystem(
    gen, wrapped
):
    _ctx, result = call_build(gen, doc_title='A/B:C*D?"E<F>G|H')
    assert result["filename"] == "ABCDEFGH_Assessment.docx"


def test_a_long_title_is_truncated_to_forty_characters_before_the_suffix(gen, wrapped):
    _ctx, result = call_build(gen, doc_title="x" * 100)
    assert result["filename"] == "x" * 40 + "_Assessment.docx"


def test_a_title_that_sanitises_to_nothing_falls_back_to_a_generic_filename(gen, wrapped):
    """A download with an empty name is worse than a generic one."""
    for title in ("", None, "///:::"):
        _ctx, result = call_build(gen, doc_title=title)
        assert result["filename"] == "Assessment_Report.docx", title


def test_the_wrapper_reports_progress_at_eighty_five_percent(gen, wrapped):
    """The pipeline's progress bar is monotonic; a wrong pct makes it jump backwards."""
    ctx, _result = call_build(gen)
    assert ctx.progress_calls == [("Building the assessment", 85)]


def test_textract_is_only_offered_to_the_build_when_the_vision_path_is_selected(
    gen, wrapped
):
    """Textract is charged per page, so the capability is withheld rather than the call
    being skipped somewhere deeper where a later edit could reach it anyway."""
    call_build(gen, use_vision=False)
    assert wrapped["textract"] is None
    call_build(gen, use_vision=True)
    assert wrapped["textract"] == "TEXTRACT"


def test_the_llm_capability_is_always_passed_through(gen, wrapped):
    """Unlike Textract it is needed on both paths, for the formula triage."""
    ctx = FakeCtx(textract=None, llm=FakeLlm())
    call_build(gen, ctx=ctx, use_vision=False)
    assert wrapped["llm"] is ctx.llm


def test_the_workspace_and_selected_indices_are_forwarded_unchanged(gen, wrapped):
    call_build(gen)
    assert wrapped["idx"] == (0, 2)
    assert wrapped["toc"] is TOC
