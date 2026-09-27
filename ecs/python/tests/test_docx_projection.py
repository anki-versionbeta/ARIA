"""Self-tests for the structural docx projection.

These matter more than they look. A projector that quietly finds nothing would make every
golden comparison pass vacuously, which is a worse outcome than having no golden test at
all — so the first test here asserts the projector actually sees content.
"""

from __future__ import annotations

import io

from tests.docx_projection import compare_projections, project_docx


def sample_docx(
    *,
    title: str = "A Standard",
    cell: str = "Applicable",
    tables: int = 1,
    second_heading: str = "Purpose / Overview",
) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(title, style="Title")
    document.add_heading("Project Information", level=1)
    document.add_heading(second_heading, level=1)
    for _ in range(tables):
        table = document.add_table(rows=2, cols=4)
        table.rows[0].cells[0].text = "Clause"
        table.rows[1].cells[0].text = "4.1"
        table.rows[1].cells[1].text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_the_projector_actually_finds_content():
    projection = project_docx(sample_docx())

    assert projection.headings == (
        ("title", "A Standard"),
        ("heading1", "Project Information"),
        ("heading1", "Purpose / Overview"),
    )
    assert len(projection.tables) == 1
    assert projection.tables[0].grid_columns == 4
    assert projection.tables[0].row_cell_counts == (4, 4)
    assert "Clause" in projection.tables[0].cell_text


def test_two_saves_of_the_same_content_project_equal():
    """The reason this module exists: two builds of identical content are not
    byte-identical, but they must project identically."""
    first, second = sample_docx(), sample_docx()

    assert not compare_projections(project_docx(first), project_docx(second)).hard


def test_a_changed_heading_is_a_hard_difference():
    diff = compare_projections(
        project_docx(sample_docx(second_heading="Purpose / Overview")),
        project_docx(sample_docx(second_heading="Purpose and Scope")),
    )

    assert diff.hard
    assert any("heading 2" in entry for entry in diff.hard)


def test_the_title_paragraph_text_is_exempt():
    """The in-document title comes from PDF metadata, which the port cannot reproduce when
    the metadata is absent — so its text is reported but never fails."""
    diff = compare_projections(
        project_docx(sample_docx(title="A Standard")),
        project_docx(sample_docx(title="Something Else Entirely")),
    )

    assert not diff.hard
    assert any("title" in entry for entry in diff.soft)


def test_a_missing_table_is_a_hard_difference():
    diff = compare_projections(
        project_docx(sample_docx(tables=2)),
        project_docx(sample_docx(tables=1)),
    )
    assert diff.hard
    assert any("table count" in entry for entry in diff.hard)


def test_cell_text_drift_beyond_the_tolerance_is_hard():
    diff = compare_projections(
        project_docx(sample_docx(cell="Applicable")),
        project_docx(
            sample_docx(cell="Not applicable - see 4.3 and Annex B for the rationale")
        ),
    )
    assert diff.hard
    assert any("similarity" in entry for entry in diff.hard)


def test_identical_documents_report_no_differences_at_all():
    diff = compare_projections(project_docx(sample_docx()), project_docx(sample_docx()))
    assert not diff.hard
    assert not diff.soft
    assert diff.report().startswith("0 structural difference(s)")
