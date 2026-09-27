"""The cap-recommendation .docx: the chip grids, the graded risk tables and the fixed layout.

`tests/test_psa_render.py` stops at the two pure helpers — `_rgb01` and `_chip_png` — so everything
that turns a recommendation into a document is untested there, and it is the half a reviewer signs:
`_fix_layout` / `_no_borders` (the four table flags that stop Word re-sizing the grid), `_chip_grid`
(the wrapping tile grid, its hatch and badge variants) and `build` itself.

The tests read the produced .docx back rather than asserting on intermediate state, because the
regressions this module documents were all invisible in the calling code:

  * a swatch that rendered as **one merged full-width band** instead of separate squares, because the
    grid (`w:gridCol`) still said "text width ÷ N" while only the cells had been narrowed;
  * a 10-item recommended list rendered as a 6-column table followed by a 4-column one, with visibly
    different chip widths, because the final short chunk sized itself with `len(chunk)`;
  * a hex-less cap drawn as flat #EEEEEE, which reads as a real grey cap rather than "unknown".

So the assertions are on the XML the reader's Word will act on: `w:tblLayout`, `w:tblW`, the grid
columns, the column count of every chunk, and the `wp:extent` of each embedded chip — plus the bytes
of the chip itself, compared against `_chip_png`, which is how "this chip is the 45° hatch" is
checked without rasterising anything.

**Nothing here touches the network.** The catalogues come from `tests/psa_fixtures.py`, which loads a
Smartsheet-shaped dict through the real ingest path, and `offline` below forces `ingest_source="xlsx"`
so `ingest_smartsheet_api.available()` is False — that is what makes `engines.recommendation`'s
default `refresh_first=True` skip the live re-ingest. `urllib.request.urlopen` is then replaced with a
failing stub, so a lost guard fails the test rather than reaching Smartsheet.

Two things are stubbed, both because the real code cannot be driven to them from outside. Every row of
the shipped `cap_palette.csv` is off-the-shelf, so `off_the_shelf` is overwritten on a real
recommendation to reach the "· not stock" suffix and the deliberate `is False` (not falsy) test beside
it. And `build` is called directly with a blank `program_label` for the heading fallback, which
`engines.export_recommendation` never produces.

One characterisation worth reading before trusting the caption: with the default "every supplier"
recommendation the header line renders **"Seal manufacturer: None"**. That looks wrong — see
`test_a_recommendation_that_spans_every_supplier_names_the_seal_manufacturer_as_none`.
"""

from __future__ import annotations

import dataclasses
import os

import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from api.backend.da_platform.settings import BACKEND_ROOT
from api.backend.da_platform.silo_registry import _load_module
from tests.psa_fixtures import (SUBJECT, SUBJECT_ROW, catalogue, con,  # noqa: F401
                                psa_config, row, sheet, snapshot_bytes)

SILO_DIR = BACKEND_ROOT / "silos" / "psa"

pytestmark = pytest.mark.skipif(
    not (SILO_DIR / "silo.py").is_file(), reason="the PSA silo is not present"
)


@pytest.fixture(scope="module")
def psa():
    _load_module("psa", SILO_DIR / "silo.py")
    from da_silos.psa import cap_export, config, engines, paths, snapshot

    return {"export": cap_export, "engines": engines, "config": config,
            "paths": paths, "snapshot": snapshot}


@pytest.fixture(autouse=True)
def offline(psa, psa_config, monkeypatch):
    """Forbid the live sheet, then forbid the network outright.

    `psa_config` supplies a token and a sheet id — it has to, because other PSA tests exercise the
    credential path — so `ingest_api.available()` would be True and `recommendation()`'s default
    `refresh_first=True` would re-ingest from Smartsheet. `ingest_source="xlsx"` is the real switch
    that turns that off (`config.PsaConfig.smartsheet_live`), and the `urlopen` stub is the backstop:
    it turns a lost guard into a failed test instead of a REST read.
    """
    psa["config"].set_config(dataclasses.replace(psa_config, ingest_source="xlsx"))
    import urllib.request

    monkeypatch.setattr(urllib.request, "urlopen",
                        lambda *a, **k: pytest.fail("the export reached the network"))


@pytest.fixture
def rebuild(psa):
    """Factory: build a catalogue of exactly the given presentations.

    Each presentation is the keyword set `psa_fixtures.row` takes, so a test states only the cap and
    vial it cares about. The first one is the subject.
    """
    def build(*presentations):
        from da_silos.psa.ingest_smartsheet import COLS

        rows = [row(COLS, n, **spec) for n, spec in enumerate(presentations)]
        return psa["snapshot"].rebuild(snapshot_bytes(sheet(rows)))

    return build


@pytest.fixture
def exported(psa):
    """Run the real export path and return the opened document."""
    def run(program=SUBJECT, presentation=SUBJECT_ROW, vendor=None):
        path, message = psa["engines"].export_recommendation(program, presentation, vendor)
        assert path, message
        return Document(path)

    return run


@pytest.fixture
def blank_cell():
    """A one-cell table in a real document, to write a tile line into."""
    doc = Document()
    return doc.add_table(rows=1, cols=1).cell(0, 0)


# -- reading a produced document -----------------------------------------------

CURRENT = "Current cap"
SELECTED = "Already selected"
RECOMMENDED = "Recommended"
RISKS = "Highest similarity risks"
DISCOURAGED = "Discouraged"


def sections(doc):
    """The document split at its level-2 headings, in body order.

    Positional rather than by index: how many grid tables a section holds depends on how many chips
    are in it, so `doc.tables[4]` silently starts reading a different section as soon as the
    recommended list changes length.
    """
    found = [("", {"paragraphs": [], "tables": []})]
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            if paragraph.style.name == "Heading 2":
                found.append((paragraph.text, {"paragraphs": [], "tables": []}))
            else:
                found[-1][1]["paragraphs"].append(paragraph)
        elif child.tag == qn("w:tbl"):
            found[-1][1]["tables"].append(Table(child, doc))
    return found


def section(doc, phrase):
    """The one section whose heading contains `phrase`."""
    hits = [body for heading, body in sections(doc) if phrase in heading]
    assert len(hits) == 1, f"expected exactly one {phrase!r} section, found {len(hits)}"
    return hits[0]


def headings(doc):
    return [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]


def texts(table):
    """Every cell of a grid table, as rows of strings."""
    return [[cell.text for cell in table_row.cells] for table_row in table.rows]


def names(table):
    """The name line (row 1) of a chip grid, padding cells dropped."""
    return [text for text in texts(table)[1] if text]


def chip_widths(table):
    """The absolute EMU width of every chip picture in a grid table."""
    return [int(e.get("cx")) for e in table._tbl.iter(qn("wp:extent"))]


def bullets(doc):
    return [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]


def images(doc):
    return {part.blob for part in doc.part.package.image_parts}


def tail_run(table, column):
    """The last run of the tail line (row 3) — the ΔE, or the graded risk."""
    return table.cell(3, column).paragraphs[-1].runs[-1]


# -- pinning the table so Word cannot re-size the grid -------------------------


@pytest.fixture
def pinned(psa):
    """A raw 2-column table put through both table helpers, in build order."""
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    grid_before = [column.width for column in table.columns]
    psa["export"]._no_borders(table)
    psa["export"]._fix_layout(table, 2)
    return table, grid_before


def test_the_table_layout_is_switched_from_auto_to_fixed(psa, pinned):
    """Under the AUTO algorithm Word treats every stated width as a hint, so the chip column
    re-sizes itself to whatever the label needs and the swatches stop being square."""
    table, _before = pinned

    assert table.autofit is False
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    assert layout is not None and layout.get(qn("w:type")) == "fixed"


def test_the_table_width_is_stated_in_twips_rather_than_left_automatic(psa, pinned):
    """`add_table` emits `type="auto" w="0"`, which asks Word to work the width out — and a
    computed width is what re-introduces the auto layout the fixed flag just turned off."""
    table, _before = pinned

    width = table._tbl.tblPr.find(qn("w:tblW"))
    assert width.get(qn("w:type")) == "dxa"
    assert int(width.get(qn("w:w"))) == int(round(psa["export"]._COL_IN * 2 * 1440))


def test_the_grid_columns_are_narrowed_from_the_full_text_width(psa, pinned):
    """This is the one auto layout was overriding: `add_table` sizes `w:gridCol` to the whole text
    width divided by the column count, and under auto layout the GRID wins over the cells — which
    is how the swatch used to render as one merged full-width band."""
    table, before = pinned

    assert before == [Inches(3.0), Inches(3.0)], "add_table splits the 6in text width"
    assert [column.width for column in table.columns] == [Inches(psa["export"]._COL_IN)] * 2


def test_every_cell_states_the_same_width_as_its_column(psa, pinned):
    """Setting only the cells was the original bug — asking for narrow cells inside a 6in grid. Both
    have to agree, or Word has a reason to reflow."""
    table, _before = pinned

    assert [cell.width for cell in table.rows[0].cells] == \
        [Inches(psa["export"]._COL_IN)] * 2


def test_the_border_switch_is_inserted_before_its_schema_successors(psa, pinned):
    """`CT_TblPr` enforces child order, so an APPENDED `w:tblBorders` lands after `w:tblLook` and
    the document is schema-invalid — Word then refuses to open it at all."""
    table, _before = pinned

    order = [child.tag for child in table._tbl.tblPr]
    assert order.index(qn("w:tblBorders")) < order.index(qn("w:tblLayout"))
    assert order.index(qn("w:tblBorders")) < order.index(qn("w:tblLook"))


def test_every_edge_of_the_grid_is_switched_off(psa, pinned):
    """A tile grid with gridlines reads as a table of data rather than as the web's swatch cards,
    and the default table style supplies them unless all six edges say otherwise."""
    table, _before = pinned

    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert [child.get(qn("w:val")) for child in borders] == ["none"] * 6
    assert {child.tag for child in borders} == {
        qn(f"w:{edge}") for edge in ("top", "left", "bottom", "right", "insideH", "insideV")}


# -- one line of tile text -----------------------------------------------------


def test_a_tile_line_is_left_aligned_and_tightly_spaced(psa, blank_cell):
    """Zeroed spacing is what makes the four stacked cells read as one compact tile; with Normal's
    spacing they read as a loose column and the name drifts away from its chip."""
    paragraph = psa["export"]._line(blank_cell, "Blue 6043")

    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.paragraph_format.space_before == Pt(0)
    assert paragraph.paragraph_format.space_after == Pt(0)
    assert paragraph.paragraph_format.line_spacing == 1.0


def test_an_empty_line_still_resets_the_paragraph(psa, blank_cell):
    """The padding cells of a short final row go through here with no text. If they kept Normal's
    spacing the last row of a grid would sit lower than the rows above it."""
    paragraph = psa["export"]._line(blank_cell, "")

    assert paragraph.text == "", "no text was asked for, so none may appear"
    assert paragraph.paragraph_format.space_after == Pt(0)
    assert paragraph.paragraph_format.line_spacing == 1.0


def test_writing_a_line_replaces_whatever_the_cell_held(psa, blank_cell):
    """A cell is written four times as the grid is built, so an appending writer would produce
    'Datwyler' followed by the next tile's supplier in the same cell."""
    psa["export"]._line(blank_cell, "first")

    psa["export"]._line(blank_cell, "second")

    assert blank_cell.text == "second"


def test_a_line_carries_its_size_weight_and_colour(psa, blank_cell):
    """8pt is what lets a 1.5in tile hold 'Magenta 2063C (L9320)' without wrapping, and the colour
    is how the reserved first choice and the risk grade are distinguished at a glance."""
    from docx.shared import RGBColor

    paragraph = psa["export"]._line(blank_cell, "1st choice", size=8, bold=True,
                                   color=RGBColor(0x33, 0x87, 0x00))

    run = paragraph.runs[-1]
    assert run.text == "1st choice"
    assert (run.font.size, run.bold, str(run.font.color.rgb)) == (Pt(8), True, "338700")


# -- the delta-E line, which mirrors a JS template literal ---------------------


@pytest.mark.parametrize("value,rendered", [
    (96.6, "ΔE 96.6"),
    (49.0, "ΔE 49"),                     # plain str() would wrongly render "ΔE 49.0"
    (0, "ΔE 0"),
    (0.0, "ΔE 0"),
])
def test_a_distance_is_formatted_the_way_the_web_formats_it(psa, value, rendered):
    """The export mirrors the on-screen view, so a reviewer comparing the two must not see 49.0 in
    one and 49 in the other and wonder which number the decision was made on."""
    assert psa["export"]._de(value) == rendered


def test_no_distance_at_all_reads_as_free(psa):
    """None means no co-located cap had a comparable hex. 'ΔE 0' there would claim a perfect
    match against something that was never measured."""
    assert psa["export"]._de(None) == "free"


def test_an_unformattable_distance_is_shown_rather_than_raised(psa):
    """`build` runs on every export, and one odd value must not lose the whole document."""
    assert psa["export"]._de("n/a") == "ΔE n/a"


# -- the wrapping chip grid ----------------------------------------------------


def test_an_empty_grid_adds_nothing_at_all(psa):
    """A 4x4 table of blank cells, or the spacer paragraph after it, would print as a band of white
    space under a heading that has no chips to show."""
    doc = Document()
    before = len(doc.paragraphs)

    psa["export"]._chip_grid(doc, [])

    assert doc.tables == []
    assert len(doc.paragraphs) == before


def test_the_grid_wraps_after_four_chips(psa, catalogue, exported):
    """Four 1.5in tiles is exactly the default template's 6.0in text width, so a fifth chip on the
    same row would push the grid into the margin."""
    grids = section(exported(), RECOMMENDED)["tables"]

    assert [len(names(table)) for table in grids] == [4, 4, 2], "ten recommendations"


def test_a_short_final_row_still_has_four_columns(psa, catalogue, exported):
    """Using `len(chunk)` is what produced a 6-column table followed by a 4-column one, with
    visibly different chip widths — the tiles of the last row have to line up with the rows above."""
    grids = section(exported(), RECOMMENDED)["tables"]

    assert [len(table.columns) for table in grids] == [psa["export"]._CHIPS_PER_ROW] * 3
    assert all(len(table.rows) == 4 for table in grids), "chip + name + supplier + tail"


def test_the_padding_cells_of_a_short_row_hold_nothing(psa, catalogue, exported):
    """They exist only to hold the column open. A stray label or chip in one would read as a
    recommendation that was never made."""
    last = section(exported(), RECOMMENDED)["tables"][-1]

    assert texts(last)[1][2:] == ["", ""]
    assert len(chip_widths(last)) == 2, "no chip in a padding column"


def test_each_tile_carries_its_name_supplier_and_distance(psa, catalogue, exported):
    """The three lines under the chip are the whole content of the web tile; a missing supplier line
    would leave the reader unable to tell a Datwyler blue from a West one."""
    first = section(exported(), RECOMMENDED)["tables"][0]

    assert texts(first)[1] == ["Black 6006", "Black 7607", "Blue 6043", "Dark Blue 4020"]
    assert texts(first)[2] == ["Datwyler", "West", "Datwyler", "West"]
    assert texts(first)[3] == ["ΔE 89.3", "ΔE 89.3", "ΔE 88.1", "ΔE 88.1"]


def test_a_chip_is_an_inline_picture_and_never_cell_shading(psa, catalogue, exported):
    """The headline decision of the module. Shading fills the whole cell, so the swatch is always as
    wide as the column its LABEL needs — which is how this export once rendered one merged
    full-width band instead of separate squares."""
    first = section(exported(), RECOMMENDED)["tables"][0]

    assert len(chip_widths(first)) == 4, "one inline picture per tile"
    assert "w:shd" not in first._tbl.xml, "no cell shading anywhere in the grid"


def test_a_chip_keeps_its_absolute_size_whatever_the_column_does(psa, catalogue, exported):
    """`wp:extent` is absolute EMU, which is what makes the bug unable to regress: the chip stays a
    small square even if a later edit disturbs a table flag and the column widens."""
    doc = exported()
    recommended = section(doc, RECOMMENDED)["tables"][0]
    risks = section(doc, RISKS)["tables"][0]

    assert set(chip_widths(recommended)) == {Inches(psa["export"]._CHIP_IN_LARGE)}
    assert set(chip_widths(risks)) == {Inches(psa["export"]._CHIP_IN)}


def test_the_recommended_grid_uses_the_larger_of_the_two_chips(psa, catalogue, exported):
    """The web shows the recommended swatches `large` and the rest `small`, and the export mirrors
    the on-screen view — so the reader's eye lands on the same section it does on screen."""
    doc = exported()

    assert psa["export"]._CHIP_IN_LARGE > psa["export"]._CHIP_IN
    assert set(chip_widths(section(doc, RECOMMENDED)["tables"][0])) != \
        set(chip_widths(section(doc, DISCOURAGED)["tables"][0]))


def test_the_chip_row_is_at_least_as_tall_as_the_chip(psa, catalogue, exported):
    """AT_LEAST, never EXACTLY: an exact row rule CLIPS an inserted picture, so the swatch would
    render as a sliver of its colour."""
    first_row = section(exported(), RECOMMENDED)["tables"][0].rows[0]

    assert first_row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
    assert Inches(psa["export"]._CHIP_IN_LARGE) <= first_row.height \
        <= Inches(psa["export"]._CHIP_IN_LARGE + 0.1)


# -- the 45 degree hatch: a colour with no solid hex ---------------------------


def test_a_hex_less_recommendation_embeds_the_hatched_chip(psa, rebuild, exported):
    """'Transparent 6001' is translucent and carries no hex. Cell shading could only make it flat
    #EEEEEE, which reads as a real grey cap; the hatch reads as "no solid colour".

    Comparing the embedded bytes against `_chip_png` is what makes this specific — a chip that had
    silently become a solid fill would still be a valid PNG of the right size.
    """
    rebuild(dict(program=SUBJECT, cap=""))
    doc = exported()

    assert names(section(doc, RECOMMENDED)["tables"][0])[0] == "Transparent 6001"
    assert psa["export"]._chip_png("", psa["export"]._CHIP_IN_LARGE) in images(doc)


def test_a_hex_less_cap_already_in_use_embeds_the_hatch_at_the_small_size(psa, rebuild, exported):
    """The same honesty at the other chip size: a comparator whose cap has no swatch must not be
    drawn as a grey cap that nobody chose."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"),
            dict(program="ABBV-400", cap="Transparent 6001"))
    doc = exported()

    assert texts(section(doc, RISKS)["tables"][0])[1][0] == "Transparent"
    assert psa["export"]._chip_png("", psa["export"]._CHIP_IN) in images(doc)


def test_a_white_cap_is_embedded_as_its_own_bordered_chip(psa, catalogue, exported):
    """The other case shading cannot express: a white fill is invisible against white paper, which
    is why every chip carries a 1px container border."""
    doc = exported()

    assert "White" in str(texts(section(doc, RISKS)["tables"][0]))
    assert psa["export"]._chip_png("#FFFFFF", psa["export"]._CHIP_IN) in images(doc)


# -- the document body ---------------------------------------------------------


def test_the_document_is_headed_by_the_program_label(psa, catalogue, exported):
    """A file named after a row number is not identifiable on its own; the heading is what ties the
    export to a program in a review pack."""
    doc = exported()

    assert doc.paragraphs[0].text == f"{SUBJECT} (Product 0)"
    assert doc.paragraphs[0].style.name == "Heading 1"


def test_a_document_with_no_program_label_still_has_a_title(psa, catalogue, tmp_path):
    """`build` is a public entry point and the label is whatever the caller resolved, so a blank one
    must not produce an untitled governance document."""
    result, program_label, presentation_label = psa["engines"].recommendation(
        SUBJECT, SUBJECT_ROW, None)
    path = psa["export"].build(str(tmp_path / "untitled.docx"), result, "", presentation_label)

    assert Document(path).paragraphs[0].text == "Cap Colour Recommendation"


def test_the_presentation_is_named_in_bold_under_the_heading(psa, catalogue, exported):
    """A program has several presentations and the recommendation is per-presentation, so a reader
    has to be able to see which one this is without opening the database."""
    bar = exported().paragraphs[1]

    assert bar.text == "Presentation: 2R  |  Commercial  |  100 mg"
    assert bar.runs[0].bold is True


def test_the_header_line_states_the_vial_state_sites_and_neighbour_count(psa, catalogue, exported):
    """These four facts are the scope of the whole answer: change the vial size, the state or the
    site list and a different set of caps is in contention."""
    caption = exported().paragraphs[2].text

    assert "Vial size: 2R" in caption
    assert "State: Liquid" in caption
    assert "Manufacturing site(s): AP16 (AP16)" in caption
    assert "2 same-state co-located product(s)" in caption


def test_a_missing_state_is_dashed_rather_than_left_blank(psa, rebuild, exported):
    """'State:    ·' reads as a rendering fault. An em dash says the sheet does not record one."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043", modality=""),
            dict(program="ABBV-400", cap="Red 6070", modality=""))

    assert "State: —" in exported().paragraphs[2].text


def test_a_recommendation_that_spans_every_supplier_names_the_seal_manufacturer_as_none(
        psa, catalogue, exported):
    """Characterisation, and it LOOKS WRONG. `vendor=None` is the default and means "every supplier
    was considered"; the caption interpolates it straight into the text, so the header of a default
    export reads 'Seal manufacturer: None'. `res.get("vendor", "")` cannot help — the key is
    present and its value is None — so a reader is told the seal manufacturer is nothing at all.
    Reported rather than worked around, because the source is frozen.
    """
    assert "Seal manufacturer: None" in exported().paragraphs[2].text


def test_naming_a_supplier_puts_that_supplier_in_the_header(psa, catalogue, exported):
    """The other half of the case above: when the caller DID restrict to a catalogue, the header
    reports it correctly — which is why the None above is a formatting gap and not a data one."""
    assert "Seal manufacturer: Datwyler" in \
        exported(vendor="Datwyler").paragraphs[2].text


def test_the_document_is_written_where_the_configuration_points(psa, catalogue):
    """The silo carries no paths of its own. A hard-coded output directory would write a GxP
    document into the source tree of whatever deployment imported it."""
    path, message = psa["engines"].export_recommendation(SUBJECT, SUBJECT_ROW, None)

    assert message == "Recommendation exported."
    assert os.path.dirname(path) == psa["paths"].output_dir()
    assert os.path.getsize(path) > 0


# -- the current cap -----------------------------------------------------------


def test_an_already_chosen_cap_gets_its_own_section(psa, catalogue, exported):
    """A presentation that has already been decided is a different question from one still being
    chosen, and the heading is what tells the reader which they are looking at."""
    doc = exported()
    body = section(doc, SELECTED)

    assert names(body["tables"][0]) == ["Blue 6043"]
    assert texts(body["tables"][0])[3] == [""] * 4, "the heading already says it is current"


def test_an_unresolved_cap_is_shown_as_the_current_cap_rather_than_a_selection(
        psa, rebuild, exported):
    """'TBD' is not a decision. Listing it under "already selected" would tell a reviewer the cap
    question was closed when it is exactly what is open."""
    rebuild(dict(program=SUBJECT, cap="TBD"), dict(program="ABBV-400", cap="Red 6070"))
    doc = exported()

    assert SELECTED not in " ".join(headings(doc))
    body = section(doc, CURRENT)
    assert names(body["tables"][0]) == ["TBD"]
    assert texts(body["tables"][0])[3][0] == "current"


def test_a_presentation_with_no_cap_at_all_says_so_in_words(psa, rebuild, exported):
    """An empty section under "Current cap" is indistinguishable from a rendering failure, and the
    assessor needs to know the sheet is blank rather than that the export broke."""
    rebuild(dict(program=SUBJECT, cap=""), dict(program="ABBV-400", cap=""))
    body = section(exported(), CURRENT)

    assert body["tables"] == [], "nothing to draw a chip for"
    assert [p.text for p in body["paragraphs"] if p.text] == ["none recorded"]


def test_a_presentation_with_no_manufacturing_site_stops_after_the_current_cap(
        psa, rebuild, exported):
    """Cap co-location is a property of a manufacturing line, so with no site there is nothing to
    recommend against. Printing empty Recommended / Risks / Discouraged sections would look like an
    answer; the note says what to go and fix instead."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043", site="TBD"))
    doc = exported()

    assert headings(doc) == [f"{SUBJECT} (Product 0)", "✓ Already selected (current cap)"]
    assert len(doc.tables) == 1
    assert "no manufacturing site recorded" in doc.paragraphs[-1].text


# -- the reserved first choice -------------------------------------------------


def test_the_reserved_first_choice_is_badged_above_its_distance(psa, rebuild, exported):
    """`recommend_program_presentations` hands each presentation a distinct #1 pick, and the badge
    is the only place the document says which colour that was — without it the top-left tile looks
    like any other candidate."""
    rebuild(dict(program=SUBJECT, cap="TBD"), dict(program="ABBV-400", cap="Red 6070"))
    grid = section(exported(), RECOMMENDED)["tables"][0]

    tail = grid.cell(3, 0)
    assert [p.text for p in tail.paragraphs] == ["1st choice · unique", "ΔE 119.8"]
    assert (tail.paragraphs[0].runs[-1].bold, str(tail.paragraphs[0].runs[-1].font.color.rgb)) \
        == (True, "338700")


def test_only_the_reserved_choice_has_a_green_bold_name(psa, rebuild, exported):
    """The highlight is what makes the pick findable. Applied to every tile it would say nothing;
    applied to none, a reader has to count columns to find the recommendation."""
    rebuild(dict(program=SUBJECT, cap="TBD"), dict(program="ABBV-400", cap="Red 6070"))
    grid = section(exported(), RECOMMENDED)["tables"][0]

    first = grid.cell(1, 0).paragraphs[0].runs[-1]
    second = grid.cell(1, 1).paragraphs[0].runs[-1]
    assert (first.bold, str(first.font.color.rgb)) == (True, "338700")
    assert (second.bold, second.font.color.rgb) == (False, None)


def test_a_presentation_that_already_has_a_cap_badges_nothing(psa, catalogue, exported):
    """Its colour is reserved rather than re-picked, so there is no "1st choice" to claim — and
    claiming one would contradict the "already selected" section above it."""
    grid = section(exported(), RECOMMENDED)["tables"][0]

    assert len(grid.cell(3, 0).paragraphs) == 1
    assert "1st choice" not in grid._tbl.xml


def test_a_supplier_line_is_muted_rather_than_black(psa, catalogue, exported):
    """The tile's hierarchy is name, then supplier, then distance. A full-strength supplier line
    competes with the colour name for the reader's attention."""
    supplier = section(exported(), RECOMMENDED)["tables"][0].cell(2, 0).paragraphs[0].runs[-1]

    assert str(supplier.font.color.rgb) == str(psa["export"]._MUTED)


# -- stock status, which the shipped palette cannot currently express ----------


def test_a_colour_the_supplier_does_not_stock_is_marked_not_stock(psa, catalogue, tmp_path):
    """Deck slide 20 prefers off-the-shelf colours, so a non-stock suggestion has to carry the
    tooling cost on its face — it is shown rather than hidden, and ranked below.

    `off_the_shelf` is overwritten on a real recommendation because every row of the shipped
    `cap_palette.csv` is currently off-the-shelf, so the engine cannot produce this shape.
    """
    result, program_label, presentation_label = psa["engines"].recommendation(
        SUBJECT, SUBJECT_ROW, None)
    result["recommended"][0]["off_the_shelf"] = False

    path = psa["export"].build(str(tmp_path / "stock.docx"), result, program_label,
                               presentation_label)

    grid = section(Document(path), RECOMMENDED)["tables"][0]
    assert texts(grid)[2][0] == "Datwyler · not stock"


def test_a_colour_of_unknown_stock_status_is_not_marked_not_stock(psa, catalogue, tmp_path):
    """`is False` on purpose, not a falsy test: None means the catalogue does not say yet, and
    printing "not stock" there would invent a procurement problem."""
    result, program_label, presentation_label = psa["engines"].recommendation(
        SUBJECT, SUBJECT_ROW, None)
    result["recommended"][0]["off_the_shelf"] = None

    path = psa["export"].build(str(tmp_path / "unknown.docx"), result, program_label,
                               presentation_label)

    grid = section(Document(path), RECOMMENDED)["tables"][0]
    assert texts(grid)[2][0] == "Datwyler"


def test_a_discouraged_colour_that_is_not_stocked_says_so_too(psa, catalogue, tmp_path):
    """The suffix is built independently in the discouraged loop, so it can drift out of step with
    the recommended one — and a reader comparing the two sections would see the same colour
    described differently."""
    result, program_label, presentation_label = psa["engines"].recommendation(
        SUBJECT, SUBJECT_ROW, None)
    result["discouraged"][0]["off_the_shelf"] = False

    path = psa["export"].build(str(tmp_path / "disc.docx"), result, program_label,
                               presentation_label)

    grid = section(Document(path), DISCOURAGED)["tables"][0]
    assert texts(grid)[2][0] == "Datwyler · not stock"


def test_no_free_colour_is_reported_rather_than_left_as_an_empty_grid(psa, catalogue, exported):
    """Restricting to a supplier the catalogue does not know leaves nothing to recommend. A blank
    section reads as a broken export; the sentence tells the assessor to review."""
    doc = exported(vendor="Nobody Ltd")
    body = section(doc, RECOMMENDED)

    assert body["tables"] == []
    assert [p.text for p in body["paragraphs"] if p.text] == \
        ["No off-the-shelf colour is free — review."]


# -- the caps already in use, graded ------------------------------------------


@pytest.fixture
def graded(rebuild, exported):
    """A catalogue producing one comparator of each grade, worst first.

    High needs the same vial AND a near-identical shade, so the comparator carries the subject's own
    cap; Medium the same vial only; Low neither.
    """
    rebuild(dict(program=SUBJECT, cap="Blue 6043"),
            dict(program="ABBV-400", cap="Blue 6043"),
            dict(program="ABBV-401", cap="Red 6070"),
            dict(program="ABBV-402", cap="White 6003", vial="10R"))
    return exported()


def test_the_caps_in_use_are_listed_worst_first(psa, graded):
    """An assessor reads the top of the list, so a High buried under two Lows is a missed
    look-alike pair."""
    grid = section(graded, RISKS)["tables"][0]

    assert texts(grid)[3][:3] == ["High similarity risk", "Medium similarity risk",
                                  "Low similarity risk"]


@pytest.mark.parametrize("column,grade", [(0, "High"), (1, "Medium"), (2, "Low")])
def test_each_grade_carries_the_design_systems_own_risk_colour(psa, graded, column, grade):
    """The export mirrors the on-screen view, so a reviewer reading the file sees the same three
    colours the assessor saw — a red High that printed grey would lose the whole signal."""
    run = tail_run(section(graded, RISKS)["tables"][0], column)

    assert run.text == f"{grade} similarity risk"
    assert str(run.font.color.rgb) == str(psa["export"]._RISK_RGB[grade])
    assert run.bold is True, "the grade is the one bold tail in the grid"


def test_a_graded_cap_names_the_canonical_colour_above_its_sheet_text(psa, graded):
    """Two lines because they answer different questions: 'Red' is what the eye sees on the line,
    'Red 6070' is what to order. Collapsing them loses one or the other."""
    grid = section(graded, RISKS)["tables"][0]

    assert texts(grid)[1][:3] == ["Blue", "Red", "White"]
    assert texts(grid)[2][:3] == ["Blue 6043", "Red 6070 · custom", "White 6003"]


def test_a_custom_cap_is_flagged_beside_its_sheet_text(psa, rebuild, exported):
    """'Red 6070' does not tie to the palette, so it is a custom cap — the BoNT/E magenta case. A
    custom colour carries tooling work the assessor has to know about."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"), dict(program="ABBV-400", cap="Red 6070"))
    grid = section(exported(), RISKS)["tables"][0]

    assert texts(grid)[2][0] == "Red 6070 · custom"


def test_a_custom_cap_named_only_by_its_colour_is_flagged_on_its_own(psa, rebuild, exported):
    """When the sheet says just 'Magenta' there is no second line to qualify, and repeating
    'Magenta · custom' under 'Magenta' would read as two different caps."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"), dict(program="ABBV-400", cap="Magenta"))
    grid = section(exported(), RISKS)["tables"][0]

    assert texts(grid)[1][0] == "Magenta"
    assert texts(grid)[2][0] == "custom"


def test_a_stock_cap_named_only_by_its_colour_has_no_second_line(psa, rebuild, exported):
    """Nothing to add: the colour word IS the sheet's text and the colour is off-the-shelf. A
    repeated line here would be noise in every tile."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"), dict(program="ABBV-400", cap="Green"))
    grid = section(exported(), RISKS)["tables"][0]

    assert texts(grid)[1][0] == "Green"
    assert texts(grid)[2][0] == ""


def test_every_graded_cap_gets_a_bullet_behind_its_grade(psa, graded):
    """The grade is a word in a tile; the bullet is the evidence for it. Without the metadata an
    assessor cannot defend the decision, and G-5 requires the audit trail to say why."""
    lines = [line for line in bullets(graded) if line.startswith(("High", "Medium", "Low"))]

    assert len(lines) == 3
    assert "ABBV-400 (Product 1)" in lines[0]
    assert "vial 2R  (same as this presentation)" in lines[0]
    assert "cap Blue 6043" in lines[0]
    assert "state Liquid" in lines[0]
    assert "ΔE 0.0 to the current cap" in lines[0]


def test_a_comparator_in_a_different_vial_is_not_claimed_to_share_one(psa, graded):
    """The vial is one of the two drivers of the grade, so saying "same as this presentation" about
    a 10R would misstate the reason a Low is a Low."""
    low = next(line for line in bullets(graded) if line.startswith("Low"))

    assert "vial 10R" in low
    assert "same as this presentation" not in low


def test_a_comparator_with_no_swatch_says_the_shade_was_not_compared(psa, rebuild, exported):
    """The honesty rule the whole cap layer runs on: a translucent cap has no shade to measure, and
    a printed 'ΔE 0' there would read as a perfect colour match."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"),
            dict(program="ABBV-400", cap="Transparent 6001"))
    line = next(line for line in bullets(exported()) if line.startswith("Medium"))

    assert "·  shade not compared" in line
    assert "ΔE" not in line.split(".")[0], "no distance may be implied before the reason"


def test_the_caption_explains_what_makes_a_cap_a_risk(psa, catalogue, exported):
    """The two drivers are not obvious from a grade alone, and this section is the one an assessor
    is asked to justify."""
    caption = " ".join(p.text for p in section(exported(), RISKS)["paragraphs"])

    assert "same vial size" in caption
    assert "similar shade" in caption


def test_no_co_located_cap_is_stated_rather_than_left_blank(psa, rebuild, exported):
    """"Nothing to compare against" is a real and reassuring answer. An empty section is not: it
    looks like the comparison was never run."""
    rebuild(dict(program=SUBJECT, cap="Blue 6043"), dict(program="ABBV-400", cap=""))
    body = section(exported(), RISKS)

    assert body["tables"] == []
    assert [p.text for p in body["paragraphs"] if p.text] == \
        ["No cap is recorded on any co-located, same-state presentation."]


# -- the discouraged colours --------------------------------------------------


def test_a_blocked_colour_appears_as_a_chip_and_as_a_reason(psa, catalogue, exported):
    """The chip is recognisable at a glance and the reason is a sentence that cannot fit a 1.5in
    tile, so the section needs both — a grid alone would not say WHY the colour is out."""
    doc = exported()
    grid = section(doc, DISCOURAGED)["tables"][0]

    assert names(grid) == ["Red 6055", "Red 6036", "Red 3767"]
    reasons = [line for line in bullets(doc) if line.startswith("Red 6055")]
    assert reasons == ["Red 6055 — Red is already used at vial size 2R at a shared "
                       "manufacturing site (same colour × vial size)."]


def test_every_discouraged_colour_gets_its_own_reason(psa, catalogue, exported):
    """One shared reason at the top would not survive a second rule being added, and an assessor
    reading a single bullet cannot tell which of three colours it refers to."""
    doc = exported()
    grid = section(doc, DISCOURAGED)["tables"][0]

    reasons = [line for line in bullets(doc) if " is already used at vial size " in line]
    assert len(reasons) == len(names(grid))


def test_nothing_discouraged_reads_as_none(psa, rebuild, exported):
    """A heading with nothing under it invites the reader to wonder whether the section failed.

    This is the LAST section, so the closing note lands in it too — hence the first paragraph
    rather than the whole list.
    """
    rebuild(dict(program=SUBJECT, cap="Blue 6043"), dict(program="ABBV-400", cap=""))
    body = section(exported(), DISCOURAGED)

    assert body["tables"] == []
    assert [p.text for p in body["paragraphs"] if p.text][0] == "None."


# -- saving the file ----------------------------------------------------------


def test_the_note_is_the_last_thing_on_the_page(psa, catalogue, exported):
    """It says the recommendation is provisional and that an SME approves the final colour. A
    caveat that is not on the document cannot be relied on to have been read."""
    last = exported().paragraphs[-1]

    assert last.text.startswith("Provisional, explainable recommendation")
    assert last.runs[-1].italic is True, "the caveat is styled as a caption, not as body text"


def test_the_saved_file_reopens_as_a_document(psa, catalogue, tmp_path):
    """`build` returns the path it was given, and the caller hands that straight to a download.
    A path returned for a file that was never written would 404 at the end of a run.
    """
    result, program_label, presentation_label = psa["engines"].recommendation(
        SUBJECT, SUBJECT_ROW, None)
    target = str(tmp_path / "reopen.docx")

    path = psa["export"].build(target, result, program_label, presentation_label)

    assert path == target
    assert len(Document(path).tables) == 6, "current cap + 3 recommended chunks + risks + blocked"
