"""ATR PDF/DOCX rendering — §1 Scope and §2 Summary are omitted when empty.

Business feedback: an empty "Summary and Conclusion" (no Data-Warehouse entry) should
not render as a bare "—" chapter; the same applies to Scope. Remaining sections must
renumber so the sequence has no gap. These render the real files and read the text back.
"""

from __future__ import annotations

import pytest

from api.backend.da_platform.settings import BACKEND_ROOT
from api.backend.da_platform.silo_registry import _load_module

SILO_DIR = BACKEND_ROOT / "silos" / "mfg_atr"

pytestmark = pytest.mark.skipif(
    not (SILO_DIR / "atr_pdf.py").is_file(), reason="the ATR pipeline is not present"
)


@pytest.fixture(scope="module")
def mods():
    _load_module("mfg_atr", SILO_DIR / "silo.py")
    rep = _load_module("mfg_atr", SILO_DIR / "atr_report.py", name="atr_report")
    pdf = _load_module("mfg_atr", SILO_DIR / "atr_pdf.py", name="atr_pdf")
    docx = _load_module("mfg_atr", SILO_DIR / "atr_docx.py", name="atr_docx")
    return rep, pdf, docx


def _make(rep, scope, summary):
    return rep.AtrReport(
        request_id_short="CMC-1", request_id_display="CMC-1", request_id_query="Q",
        header={"cmc_request_id": "1", "project_code": "PC", "batch_sample_id": "B1"},
        methods=[rep.MethodRow("Sub-Visible Particles - MFI", "ARM-19-01273", "ELN1")],
        sample_columns=["S1"], results=[], scope=scope, summary_conclusion=summary,
        flags=[], is_stability=True,
        batch_matrices=[rep.BatchMatrix(
            batch_id="B1", lead_headers=["Sample"], group_headers=[("MFI", 1)],
            sub_headers=["x"], rows=[["S1", "1"]], actual_flags=[[False, False]])],
    )


def _pdf_text(path):
    from pypdf import PdfReader
    return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)


def _docx_text(path):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_scope_and_summary_keep_classic_numbering_when_present(mods, tmp_path):
    rep, pdf, docx = mods
    r = _make(rep, "The scope.", "The conclusion.")
    for text in (_pdf_text(pdf.generate_atr_pdf(r, tmp_path / "a.pdf")),
                 _docx_text(docx.generate_atr_docx(r, tmp_path / "a.docx"))):
        assert "1 SCOPE" in text
        assert "2 SUMMARY AND CONCLUSION" in text
        assert "3 ANALYTICAL METHODS" in text
        assert "4 RESULTS" in text
        assert "3.1" in text and "4.1" in text


def test_empty_scope_and_summary_are_omitted_and_sections_renumber(mods, tmp_path):
    rep, pdf, docx = mods
    r = _make(rep, "", "")
    for text in (_pdf_text(pdf.generate_atr_pdf(r, tmp_path / "b.pdf")),
                 _docx_text(docx.generate_atr_docx(r, tmp_path / "b.docx"))):
        assert "SCOPE" not in text
        assert "SUMMARY AND CONCLUSION" not in text
        assert "—" not in text and "&mdash;" not in text
        assert "1 ANALYTICAL METHODS" in text
        assert "2 RESULTS" in text
        assert "1.1" in text and "2.1" in text


def test_only_summary_present_renumbers_scope_away(mods, tmp_path):
    rep, pdf, _docx = mods
    text = _pdf_text(pdf.generate_atr_pdf(_make(rep, "", "Only a conclusion."),
                                          tmp_path / "c.pdf"))
    assert "SCOPE" not in text
    assert "1 SUMMARY AND CONCLUSION" in text
    assert "2 ANALYTICAL METHODS" in text
    assert "3 RESULTS" in text


def _report_with_matrix(rep, bm, *, is_stability):
    return rep.AtrReport(
        request_id_short="C", request_id_display="C", request_id_query="Q",
        header={"cmc_request_id": "1", "project_code": "P", "batch_sample_id": "B"},
        methods=[rep.MethodRow("MFI", "M", "E")], sample_columns=[], results=[],
        scope="", summary_conclusion="", flags=[], is_stability=is_stability,
        batch_matrices=[bm])


def test_wide_dev_sample_results_table_is_transposed(mods, tmp_path):
    """More measurements than samples -> rows=measurements, columns=samples, so the
    'Test'/'Parameter' header appears and the sample is a column (business feedback)."""
    rep, pdf, _docx = mods
    bm = rep.BatchMatrix(
        batch_id="B1", lead_headers=["Sample"],
        group_headers=[("MFI", 2), ("SEC", 1)],
        sub_headers=["HMW [%]", "LMW [%]", "Main Peak [%]"],
        rows=[["Control", "0.2", "1.7", "98.1"]],
        actual_flags=[[False, False, False, False]])
    text = _pdf_text(pdf.generate_atr_pdf(_report_with_matrix(rep, bm, is_stability=False),
                                          tmp_path / "wide.pdf"))
    assert "Parameter" in text            # transposed header (only present when transposed)
    assert "HMW [%]" in text and "Main Peak [%]" in text
    assert "Control" in text


def test_stability_results_table_is_never_transposed(mods, tmp_path):
    """Stability ATRs keep samples/timepoints as rows even when measurements outnumber
    them, so the transposed 'Parameter' header is absent and the samples stay as rows."""
    rep, pdf, _docx = mods
    bm = rep.BatchMatrix(
        batch_id="B1", lead_headers=["Sample"], group_headers=[("MFI", 3)],
        sub_headers=["a [%]", "b [%]", "c [%]"],            # 3 measurements > 2 samples
        rows=[["S1", "1", "2", "3"], ["S2", "4", "5", "6"]],
        actual_flags=[[False, False, False, False]] * 2)
    text = _pdf_text(pdf.generate_atr_pdf(_report_with_matrix(rep, bm, is_stability=True),
                                          tmp_path / "tall.pdf"))
    assert "Parameter" not in text        # gated off for stability, so not transposed
    assert "S1" in text and "S2" in text


def test_wide_docx_results_table_is_transposed(mods, tmp_path):
    """The DOCX transposes wide dev-sample tables too, so Word matches the PDF."""
    rep, _pdf, docx = mods
    bm = rep.BatchMatrix(
        batch_id="B1", lead_headers=["Sample"],
        group_headers=[("MFI", 2), ("SEC", 1)],
        sub_headers=["HMW [%]", "LMW [%]", "Main Peak [%]"],
        rows=[["Control", "0.2", "1.7", "98.1"]],
        actual_flags=[[False, False, False, False]])
    text = _docx_text(docx.generate_atr_docx(_report_with_matrix(rep, bm, is_stability=False),
                                             tmp_path / "wide.docx"))
    assert "Parameter" in text
    assert "HMW [%]" in text and "Control" in text


def test_docx_editable_sections_are_bordered_tables(mods, tmp_path):
    """Regulatory + HQC render as bordered 'Table Grid' tables (lines/borders) matching the
    PDF, not plain paragraphs."""
    import docx as _docx
    rep, pdf, docx = mods
    d = pdf.ManualDefaults(regulatory="Yes", spec_id="ATR-1", hqc="meets", remarks_na=True)
    path = docx.generate_atr_docx(_make(rep, "", ""), tmp_path / "edit.docx", defaults=d)
    doc = _docx.Document(str(path))

    def find(text):
        return [t for t in doc.tables
                if any(text in c.text for r in t.rows for c in r.cells)]

    reg = find("Regulatory report intended")
    hqc = find("HQC Assessment")
    assert reg and reg[0].style.name == "Table Grid", "Regulatory must be a bordered table"
    assert hqc and hqc[0].style.name == "Table Grid", "HQC must be a bordered table"
    hqc_text = "\n".join(c.text for r in hqc[0].rows for c in r.cells)
    assert "Specification ID:" in hqc_text
    assert "Results meet acceptance criteria" in hqc_text
    assert "Remarks:" in hqc_text


def _ink(page, rect):
    """Count dark pixels inside a widget rect — a proxy for 'is this box drawn?'."""
    import fitz
    pad = 1.0
    clip = fitz.Rect(rect.x0 - pad, rect.y0 - pad, rect.x1 + pad, rect.y1 + pad)
    pix = page.get_pixmap(dpi=300, clip=clip, annots=True)
    s, ch = pix.samples, pix.n
    return sum(1 for i in range(0, len(s), ch) if s[i] < 128)


def test_pdf_checkboxes_are_visible_whether_or_not_they_are_set(mods, tmp_path):
    """Feedback: unchecked boxes must stay visible, and set ones carry a cross.

    Renders the real PDF, then rasterises the checked 'Yes' box and the unchecked 'No'
    box of the Regulatory radio and asserts both are drawn (the unchecked box is not
    blank) and the set one has more ink (the cross on top of the box outline).
    """
    import fitz
    rep, pdf, _docx = mods
    d = pdf.ManualDefaults(regulatory="Yes", spec_id="N/A", hqc="meets", remarks_na=True)
    path = pdf.generate_atr_pdf(_make(rep, "", ""), tmp_path / "cb.pdf", defaults=d)
    page = fitz.open(str(path))[0]

    reg = [w for w in page.widgets() if w.field_name == "regulatory"]
    assert len(reg) == 2, "Regulatory should have a Yes and a No box"

    def on_value(w):
        return next(s for s in w.button_states()["normal"] if s != "Off")

    checked = [w for w in reg if w.field_value == on_value(w)]
    unchecked = [w for w in reg if w.field_value != on_value(w)]
    assert len(checked) == 1 and len(unchecked) == 1

    ink_unchecked = _ink(page, unchecked[0].rect)
    ink_checked = _ink(page, checked[0].rect)
    assert ink_unchecked > 0, "the unchecked box must still be drawn (visible outline)"
    assert ink_checked > ink_unchecked, "the set box must add a cross over the outline"


def test_pdf_has_no_actual_value_amber_legend(mods, tmp_path):
    """Actual-Value use is no longer shown in the PDF (no amber legend); it moved to a UI
    remark. A matrix with an Actual-sourced cell must still render without the legend."""
    rep, pdf, _docx = mods
    bm = rep.BatchMatrix(
        batch_id="B1", lead_headers=["Sample"], group_headers=[("MFI", 2)],
        sub_headers=["HMW [%]", "Main Peak [%]"],
        rows=[["Control", "0.2", "98.1"]],
        actual_flags=[[False, True, False]])          # HMW came from Actual
    text = _pdf_text(pdf.generate_atr_pdf(_report_with_matrix(rep, bm, is_stability=False),
                                          tmp_path / "amber.pdf"))
    assert "Amber" not in text and "amber" not in text
    assert "populated from the" not in text and "Reported Value was unavailable" not in text


def test_finalized_pdf_does_not_set_needappearances(mods, tmp_path):
    """NeedAppearances makes each viewer regenerate its own appearance; Chrome/PDFium then
    drops the checkbox box outline (checkmark only, nothing when unchecked). The finalized
    record must leave it off so every viewer uses the embedded box appearances."""
    from api.backend.da_platform.silo_registry import _load_module
    from pypdf import PdfReader

    rep, pdf, _docx = mods
    finalize = _load_module("mfg_atr", SILO_DIR / "finalize.py", name="finalize")
    d = pdf.ManualDefaults(regulatory="Yes", hqc="meets", remarks_na=True)
    clean = pdf.generate_atr_pdf(_make(rep, "", ""), tmp_path / "clean.pdf", defaults=d)
    locked = finalize.finalize_pdf(clean, tmp_path / "locked.pdf")

    reader = PdfReader(str(locked))
    if reader.is_encrypted:
        reader.decrypt("")
    acro = reader.trailer["/Root"].get("/AcroForm")
    need = acro.get_object().get("/NeedAppearances") if acro is not None else None
    assert not need, "finalized PDF must not force NeedAppearances (boxes vanish in Chrome/PDFium)"
