"""Producing the assessment document: `populate_template.populate()`, end to end.

`test_psa_render.py` covers the small pure helpers (`_codes`, `_family`, `_risk_box`,
`_set_cell_text`). This file covers the thing they exist to serve — the populator itself — by
running it against a real catalogue and reading the produced `.docx` back with
`psa_fixtures.docx_text()`. Nothing here asserts on intermediate state: the deliverable is a
document an assessor signs, so the assertions are about what that document contains.

**Nothing here touches the network.** The catalogue comes from `tests/psa_fixtures.py`, which
loads a canned Smartsheet payload through the real ingest path, and every call passes
`refresh_first=False` (or calls `populate()` directly), so no live sheet is ever read. The
photo is a PNG built in-process with PIL rather than fetched.

The behaviours pinned here, and why each one matters on a GxP form:

  * **Part E is one block per manufacturing site** (QPP11-04-001-G004 §3). The template ships
    a single Part E table at index 3; `_render_part_e` fills it for the first site and
    deep-copies a page break + clone for each further one. A product at two sites must
    therefore produce one more table than a product at one — a populator that filled only the
    first would silently ask one site to sign for both.
  * **The photo cell is an image or blank, never prose.** `verify.py` accepts blank; what it
    rejects is text pretending to be a photo. `populate(photo=<bytes>)` is the ARIA path,
    where the image exists only as run media and there is no file on disk to point at.
  * **A missing source leaves a typed `[PENDING: ...]` tag**, not a blank and not a guess. The
    tag names which document supplies the value, which is what makes a deferred field
    actionable rather than merely absent.
  * **The product-family label distinguishes liquid from lyophilised**, because the comparator
    universe for Parts D and E is scoped by dosage-form family.
"""

from __future__ import annotations

import io
import os
import sqlite3

import pytest

from api.backend.da_platform.settings import BACKEND_ROOT
from api.backend.da_platform.silo_registry import _load_module

from tests.psa_fixtures import (  # noqa: F401 — fixtures must be imported BY NAME to be visible
    SUBJECT,
    SUBJECT_ROW,
    catalogue,
    con,
    docx_text,
    generated_report,
    psa_config,
    row,
    sheet,
    snapshot_bytes,
    subject_product_id,
)

SILO_DIR = BACKEND_ROOT / "silos" / "psa"

pytestmark = pytest.mark.skipif(
    not (SILO_DIR / "silo.py").is_file(), reason="the PSA silo is not present"
)


@pytest.fixture(scope="module")
def psa():
    _load_module("psa", SILO_DIR / "silo.py")
    from da_silos.psa import paths, populate_template

    return {"template": populate_template, "paths": paths}


def png_bytes(size=(120, 90), colour=(200, 30, 30)):
    """A real PNG, built in-process. python-docx reads the image header to size the shape, so a
    stub of arbitrary bytes would fail inside `add_picture` rather than exercise the embed."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()


def output_path(psa, program=SUBJECT):
    return os.path.join(psa["paths"].output_dir(), f"{program}_PSA_generated.docx")


def open_output(psa, program=SUBJECT):
    from docx import Document

    return Document(output_path(psa, program))


def product_id_for(con, program):
    return subject_product_id(con, program)


# ── the document gets produced at all ─────────────────────────────────────────


def test_populating_the_subject_writes_a_document_where_the_caller_will_look_for_it(psa, con):
    """`workflow.process_documents` reports `output_path` by reconstructing this exact name and
    then stats it. A populator that saved anywhere else would report `output_exists=False` on a
    run that in fact succeeded."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert os.path.exists(output_path(psa))


def test_the_produced_document_carries_the_four_parts_of_the_form(psa, con):
    """Parts A/B, C, D and E are addressed by table index by both the populator and
    `verify.py`. Losing one would move every later index and misfile every value."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert len(open_output(psa).tables) >= 4


def test_an_unknown_program_is_refused_rather_than_producing_an_empty_form(psa, catalogue):
    """A blank assessment form carrying a program code nobody can find is worse than no
    document: it looks like a considered "nothing to report"."""
    with pytest.raises(SystemExit):
        psa["template"].populate("ABBV-NOT-IN-THE-SHEET")


def test_a_program_code_with_no_presentation_selected_still_produces_a_document(psa, catalogue):
    """The CLI path passes no `product_id`. The populator then takes the first row and says so
    on stdout, so a multi-presentation product is not silently narrowed without a trace."""
    psa["template"].populate(SUBJECT)

    assert os.path.exists(output_path(psa))


# ── Part A: the Smartsheet-sourced cells ──────────────────────────────────────


@pytest.mark.parametrize(
    "expected",
    ["NDC-000", "Liquid", "2R", "Vial", "Colourless"],
    ids=["list_number", "form", "size", "shape", "colour"],
)
def test_every_smartsheet_backed_part_a_cell_reaches_the_document(psa, con, expected):
    """These are `source='smartsheet'` rows in `template_field_map`, read straight off the
    product row. A resolver naming a column that does not exist would raise; one naming the
    WRONG column produces a plausible-looking form describing another presentation."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert expected in docx_text(output_path(psa))


def test_the_product_name_cell_falls_back_to_the_smartsheet_short_name(psa, con):
    """`product_name_full` prefers the TPP's INN, which the fixture catalogue has no document
    for. Without the fallback Part A's headline cell would read "TBD" on every product whose
    TPP has not been uploaded — which is most of them."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    part_a = open_output(psa).tables[0]

    assert part_a.rows[3].cells[1].text.strip() == "Product 0"


def test_a_known_per_vial_strength_is_appended_to_the_product_name(psa, con):
    """The signed baselines write the strength beside the name. It is only appended when known,
    so the cell never reads "Product 0 (None)"."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET per_vial_strength='100 mg/vial' WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert "Product 0 (100 mg/vial)" in open_output(psa).tables[0].rows[3].cells[1].text


def test_an_extracted_inn_wins_over_the_smartsheet_short_name(psa, con):
    """The INN is what a regulator recognises; the Smartsheet's "Product 0" is an internal
    label. When the TPP has been read, the form must show the former."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET inn_name='Etentamig' WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert open_output(psa).tables[0].rows[3].cells[1].text.strip() == "Etentamig"


def test_an_empty_smartsheet_cell_becomes_tbd_rather_than_a_blank(psa, con):
    """`verify.py` fails a blank Part A cell. TBD is the honest answer — "we looked and the
    sheet does not say" — where an empty cell reads as an answered question."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET marking=NULL WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert open_output(psa).tables[0].rows[8].cells[1].text.strip() == "TBD"


# ── Part A: the photo cell ────────────────────────────────────────────────────


def test_explicit_photo_bytes_are_embedded_as_exactly_one_image(psa, con):
    """The ARIA `report` stage has the photo only as run media — there is no ingested file on
    this container's disk — so bytes are the only way in. Exactly one shape, because a second
    embed per run would grow the document on every regeneration."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT),
                             photo=png_bytes())

    assert len(open_output(psa).inline_shapes) == 1


def test_an_embedded_photo_leaves_no_text_in_the_photo_cell(psa, con):
    """`verify.py`'s rule is: blank, or an image, or a `[PENDING` tag — never prose. Text left
    beside an embedded image would trip the "placeholder pretending to be a photo" check."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT),
                             photo=png_bytes())

    photo_cell = open_output(psa).tables[0].rows[10].cells[1]

    assert photo_cell.text.strip() == ""


def test_a_subject_with_no_photo_gets_a_blank_cell_and_no_image(psa, con):
    """A presentation with no photo is normal. `_derive` returns "" for the photo key rather
    than a placeholder, and `verify.py` accepts a blank cell — so a missing image is recorded,
    not treated as a defect."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    doc = open_output(psa)

    assert len(doc.inline_shapes) == 0
    assert doc.tables[0].rows[10].cells[1].text.strip() == ""


def test_a_tall_photo_is_height_bounded_so_it_cannot_reflow_the_form(psa, con):
    """`_embed_image` caps the height at 1.4in with the aspect preserved. An unbounded portrait
    image balloons the Photo row far enough to push a later table under the repeating page
    header, which is a layout defect on a form that gets printed and signed."""
    from docx.shared import Inches

    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT),
                             photo=png_bytes(size=(100, 600)))

    shape = open_output(psa).inline_shapes[0]

    assert shape.height <= Inches(1.4)
    assert shape.width < Inches(1.9), "a capped image is narrowed to keep its aspect ratio"


def test_a_wide_photo_keeps_the_nominal_width(psa, con):
    """The counterpart: a landscape image is already inside the height box, so it is embedded
    at the full 1.9in and stays clearly visible."""
    from docx.shared import Inches

    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT),
                             photo=png_bytes(size=(600, 100)))

    assert open_output(psa).inline_shapes[0].width == Inches(1.9)


def test_a_database_photo_path_that_does_not_exist_is_not_embedded(psa, con):
    """A stale `product_photo` attribute — the image was ingested on another container, or has
    since been cleaned up — must fall back to writing the value rather than raising inside
    python-docx and failing the whole report."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'PDD')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) "
                "VALUES (?, 'product_photo', 'Output_Files/assets/gone.png')", (doc_id,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert len(open_output(psa).inline_shapes) == 0


def test_a_real_photo_file_on_disk_is_embedded(psa, con):
    """The non-ARIA path: the Smartsheet ingest saved the image under the output assets dir and
    recorded a repo-relative path. Resolved against `paths.root()` and embedded from disk."""
    pid = product_id_for(con, SUBJECT)
    image_dir = os.path.join(psa["paths"].output_dir(), "assets")
    os.makedirs(image_dir, exist_ok=True)
    image_path = os.path.join(image_dir, "vial.png")
    with open(image_path, "wb") as handle:
        handle.write(png_bytes())
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'PDD')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'product_photo', ?)", (doc_id, image_path))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert len(open_output(psa).inline_shapes) == 1


def test_explicit_bytes_win_over_a_photo_file_on_disk(psa, con):
    """Under ARIA the run's own media is authoritative: the report must describe the catalogue
    the run captured, not whatever image happens to be left on this container."""
    pid = product_id_for(con, SUBJECT)
    image_dir = os.path.join(psa["paths"].output_dir(), "assets")
    os.makedirs(image_dir, exist_ok=True)
    disk_image = os.path.join(image_dir, "stale.png")
    with open(disk_image, "wb") as handle:
        handle.write(png_bytes(size=(40, 40)))
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'PDD')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'product_photo', ?)", (doc_id, disk_image))
    con.commit()
    run_media = png_bytes(size=(600, 100))

    psa["template"].populate(SUBJECT, product_id=pid, photo=run_media)

    from docx.shared import Inches

    shape = open_output(psa).inline_shapes[0]
    # The run-media image is 6:1, so it embeds at the full nominal width; the 1:1 disk image
    # would have been height-capped and therefore narrower.
    assert shape.width == Inches(1.9)


# ── Part B: the site cells ────────────────────────────────────────────────────


def test_the_manufacturing_site_cell_falls_back_to_the_site_code(psa, con):
    """No PDD has been uploaded, so there is no full address. The site CODE is a real answer;
    a `[PENDING` tag here would defer a field the Smartsheet already answers."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert open_output(psa).tables[0].rows[14].cells[0].text.strip() == "AP16"


def test_an_extracted_full_address_replaces_the_site_code(psa, con):
    """Part B is meant to carry the site's postal address. Once the PDD has been read, the form
    shows it rather than the internal code."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET mfr_site_address=? WHERE product_id=?",
                ("AbbVie Inc., 1 North Waukegan Road, North Chicago, IL", pid))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert "North Waukegan Road" in open_output(psa).tables[0].rows[14].cells[0].text


def test_an_analyst_confirmed_site_wins_over_both(psa, con):
    """A signed assessment record is a human determination. It must override the extracted
    address, which is a machine reading of a document."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET mfr_site_address='an extracted address' WHERE product_id=?",
                (pid,))
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'other')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'b_manufacturing_site', 'AbbVie Ireland NL B.V., Manorhamilton Road, Sligo')",
                (doc_id,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    text = open_output(psa).tables[0].rows[14].cells[0].text
    assert "Sligo" in text
    assert "an extracted address" not in text


def test_the_packaging_site_cell_shows_the_readable_site_name(psa, con):
    """A site code means nothing to an external reviewer, so `SITE_NAMES` expands it. AP16 is
    the US site."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert "AbbVie US, AP16" in open_output(psa).tables[0].rows[14].cells[2].text


def test_a_second_site_is_expanded_onto_its_own_line(psa, con):
    """One site per line: run together, an address flows into the next site's name and a
    reviewer cannot tell how many sites were assessed."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET dp_pkging_site_raw='AP16, LU' WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    cell = open_output(psa).tables[0].rows[14].cells[2]

    assert "AbbVie US, AP16" in cell.text
    assert "AbbVie DE, LU" in cell.text
    assert cell.text.count("\n") == 1, "one site per line, not run together"


def test_an_unknown_site_code_is_shown_verbatim(psa, con):
    """`SITE_NAMES` is a small dictionary of the two sites in scope today. A code it does not
    know must reach the form as itself rather than be dropped."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET dp_pkging_site_raw='ZZ99' WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert "ZZ99" in open_output(psa).tables[0].rows[14].cells[2].text


def test_a_placeholder_only_site_cell_becomes_tbd(psa, con):
    """"TBD" in the sheet is not a site. Expanded literally it would produce a Part B naming a
    site that does not exist."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET dp_pkging_site_raw='TBD' WHERE product_id=?", (pid,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert open_output(psa).tables[0].rows[14].cells[2].text.strip() == "TBD"


# ── Part C: the execution-time signature cells ────────────────────────────────


def test_the_part_c_signature_cells_are_left_blank(psa, con):
    """These are `source='execution'` with an empty placeholder. A generator that pre-filled a
    quality approver's name and signature would be forging an approval — which is why
    `verify.py` fails a NON-blank Part C."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    part_c = open_output(psa).tables[1]

    assert part_c.rows[3].cells[1].text.strip() == ""
    assert part_c.rows[3].cells[2].text.strip() == ""


# ── Part D: the pipeline similarity assessment ────────────────────────────────


def test_part_d_names_the_liquid_family_for_a_liquid_product(psa, con):
    """Parts D and E are scoped to the subject's dosage-form family. Naming the wrong family
    would document a comparison against products that were never assessed."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert "Liquid" in open_output(psa).tables[2].rows[3].cells[1].text


def test_a_lyophilised_product_gets_the_lyophilised_family_label(psa, con):
    """The fixture's ABBV-404 is a Lyo Powder. Its Part D must not describe it as a liquid —
    lyophilised and liquid vials are separate comparator universes."""
    psa["template"].populate("ABBV-404", product_id=product_id_for(con, "ABBV-404"))

    families = open_output(psa, "ABBV-404").tables[2].rows[3].cells[1].text

    assert "yophili" in families, families
    assert "Liquid" not in families


def test_exactly_one_similarity_risk_box_is_ticked_in_part_d(psa, con):
    """`verify.py` looks for a `☒`. Two ticks would mean the form asserts both answers at
    once, which no assessor could sign."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    risk_cell = open_output(psa).tables[2].rows[4].cells[1]

    assert risk_cell.text.count("☒") == 1
    assert risk_cell.text.count("☐") == 1


def test_the_risk_boxes_carry_the_templates_own_font(psa, con):
    """Copied verbatim from the blank template's `☐` runs. Without the eastAsia hint Word
    resolves U+2612 through the Latin font and the tick renders as a missing-character box on
    the approved form."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    xml = open_output(psa).tables[2].rows[4].cells[1]._tc.xml

    assert "MS Gothic" in xml
    assert 'w:hint="eastAsia"' in xml


def test_the_d2_comments_cell_carries_the_draft_narrative(psa, con):
    """The auto-fill is explicitly labelled a draft, so an SME reading the form knows it was
    machine-written and must be reviewed before signature."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert "draft" in open_output(psa).tables[2].rows[6].cells[0].text.lower()


def test_an_analyst_record_overrides_the_auto_filled_part_d(psa, con):
    """`setdefault` — a human determination always wins over the risk engine's draft. The
    reverse would overwrite a signed assessment with a generated guess."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'other')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    for key, value in [("d1_product_families", "Liquid vial products, as assessed by the SME."),
                       ("d1_similarity_risk", "Yes"),
                       ("d2_comments", "The analyst's own conclusion.")]:
        con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) "
                    "VALUES (?,?,?)", (doc_id, key, value))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    part_d = open_output(psa).tables[2]
    assert part_d.rows[3].cells[1].text.strip() == "Liquid vial products, as assessed by the SME."
    assert part_d.rows[6].cells[0].text.strip() == "The analyst's own conclusion."
    box = part_d.rows[4].cells[1].text
    assert box.index("☒") < box.index("☐"), "an analyst 'Yes' must tick Yes"


def test_a_multi_paragraph_analyst_comment_keeps_its_paragraphs(psa, con):
    """Comments are written as separate paragraphs rather than one run with newlines in it,
    because Word renders an embedded newline as a single line and the assessor's structure is
    lost."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'other')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'d2_comments', 'First conclusion.\nSecond conclusion.')", (doc_id,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    cell = open_output(psa).tables[2].rows[6].cells[0]

    assert len(cell.paragraphs) == 2


def test_the_d3_approval_cells_are_left_blank(psa, con):
    """D.3 is an execution-time signature, not driven by the field map at all. It must survive
    population untouched."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    part_d = open_output(psa).tables[2]

    assert part_d.rows[10].cells[1].text.strip() == ""
    assert part_d.rows[10].cells[2].text.strip() == ""


# ── Part E: one block per manufacturing site ──────────────────────────────────


def test_a_product_at_one_site_produces_a_single_part_e_block(psa, con):
    """The baseline for the cloning test below. One site, one signature block."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert len(open_output(psa).tables) == 4


def test_a_product_at_two_sites_clones_part_e(psa, psa_config):
    """QPP11-04-001-G004 §3: "each AbbVie manufacturing site should complete a separate Part
    E". One block for two sites would ask AP16's QA to sign for LU's product as well."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    two_site = row(COLS, 0, program=SUBJECT, site="AP16\nLU")
    snapshot.rebuild(snapshot_bytes(sheet([two_site])))

    psa["template"].populate(SUBJECT)

    assert len(open_output(psa).tables) == 5, "the template's single Part E, plus one clone"


def test_each_cloned_part_e_block_names_its_own_site(psa, psa_config):
    """A clone that carried the first site's name would be indistinguishable from the original
    and both would be signed by the same site's QA."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    e_blocks = open_output(psa).tables[3:]
    named = [block.rows[3].cells[1].text for block in e_blocks]

    assert any("AP16" in name for name in named)
    assert any("LU" in name for name in named)


def test_each_part_e_block_carries_its_own_ticked_risk_box(psa, psa_config):
    """`verify.py` checks every Part E block from index 3 onward. A clone whose risk cell was
    left as the template had it would ship an unanswered similarity question."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    for block in open_output(psa).tables[3:]:
        assert block.rows[5].cells[1].text.count("☒") == 1


def test_each_part_e_block_addresses_its_own_site_qa_function(psa, psa_config):
    """E.3 names the function that signs. The template ships the generic "Site QA – Site
    Name"; a block that kept it would not tell anyone whose signature is required."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    labels = [block.rows[11].cells[0].text for block in open_output(psa).tables[3:]]

    assert "Site QA - AP16" in labels
    assert "Site QA - LU" in labels


def test_the_cloned_block_is_separated_by_a_page_break(psa, psa_config):
    """Successive Part E blocks are separate signature pages. Run together on one page, a
    signature could be read as applying to the block above it."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    from docx import Document

    xml = Document(output_path(psa)).element.body.xml

    assert 'w:type="page"' in xml


def test_the_part_e_blocks_keep_the_templates_structure(psa, psa_config):
    """The clone is a deep copy of the PRISTINE table, taken before the first site is written
    — so a second site's block has the same twelve rows and no trace of the first site's
    values. A shallow copy would share elements with the original."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    blocks = open_output(psa).tables[3:]

    assert [len(b.rows) for b in blocks] == [12, 12]
    assert "LU" not in blocks[0].rows[3].cells[1].text
    assert "AP16" not in blocks[1].rows[3].cells[1].text


def test_part_e_is_scoped_to_manufacturing_sites_only(psa, psa_config):
    """Business decision 2026-07-23 (QPP §3): packaging sites appear in the app preview but
    are never written to the report's Part E. A product manufactured at one site and packaged
    at another must still produce exactly one block."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    split = row(COLS, 0, program=SUBJECT)
    for c in split["cells"]:
        if c["columnId"] == 100 + COLS["_dp_pkging"]:
            c["value"] = "LU"
    snapshot.rebuild(snapshot_bytes(sheet([split])))

    psa["template"].populate(SUBJECT)

    doc = open_output(psa)
    assert len(doc.tables) == 4, "one manufacturing site → one Part E block"
    assert "AP16" in doc.tables[3].rows[3].cells[1].text


def test_an_analyst_part_e_record_collapses_to_one_block(psa, con):
    """When a human has written Part E, the form shows what they wrote — a single block —
    rather than a per-site auto-fill contradicting it."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'other')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    for key, value in [("e1_abbvie_site", "AbbVie US, AP16 (North Chicago)"),
                       ("e1_product_families", "Liquid vials, per the SME."),
                       ("e1_similarity_risk", "Yes"),
                       ("e2_comments", "Assessed on site; no mix-up risk.")]:
        con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) "
                    "VALUES (?,?,?)", (doc_id, key, value))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    doc = open_output(psa)
    assert len(doc.tables) == 4
    part_e = doc.tables[3]
    assert part_e.rows[3].cells[1].text.strip() == "AbbVie US, AP16 (North Chicago)"
    assert part_e.rows[7].cells[0].text.strip() == "Assessed on site; no mix-up risk."
    box = part_e.rows[5].cells[1].text
    assert box.index("☒") < box.index("☐"), "an analyst 'Yes' must tick Yes"


def test_an_analyst_part_e_record_leaves_the_qa_function_as_the_template_had_it(psa, con):
    """`_fill_e_table` only writes E.3 when a `qa_label` is supplied, and the analyst path
    supplies None. Writing a fabricated site name into a signature line would be worse than
    leaving the form's own generic label for the analyst to complete."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'other')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'e1_abbvie_site', 'AbbVie US, AP16')", (doc_id,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    assert "Site QA" in open_output(psa).tables[3].rows[11].cells[0].text
    assert "AP16" not in open_output(psa).tables[3].rows[11].cells[0].text


def test_a_product_with_no_site_still_gets_a_part_e_block(psa, psa_config):
    """`verify.py` fails a document with no Part E. The fallback block says plainly that no
    co-located product was found, which is a reportable answer; an absent Part E is not."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="TBD")])))

    psa["template"].populate(SUBJECT)

    part_e = open_output(psa).tables[3]

    assert part_e.rows[3].cells[1].text.strip() == "TBD"
    assert "no mix-up risk" in part_e.rows[7].cells[0].text.lower()


def test_the_e3_signature_cells_stay_blank_in_every_block(psa, psa_config):
    """E.3's Name and Signature are execution-time. `_fill_e_table` writes only the Function
    column, so the two signature cells of every block — original and clone — stay empty."""
    from da_silos.psa.ingest_smartsheet import COLS
    from da_silos.psa import snapshot

    snapshot.rebuild(snapshot_bytes(sheet([row(COLS, 0, program=SUBJECT, site="AP16\nLU")])))

    psa["template"].populate(SUBJECT)

    for block in open_output(psa).tables[3:]:
        assert block.rows[11].cells[1].text.strip() == ""
        assert block.rows[11].cells[2].text.strip() == ""


# ── the deferred fields ───────────────────────────────────────────────────────


def test_a_field_no_source_supplies_carries_a_typed_pending_tag(psa, con, monkeypatch):
    """A bare blank tells the assessor nothing and a bare "TBD" barely more. The tag names
    which document supplies the value, which is what makes a deferred field actionable.

    `risk.form_fields` is documented as never raising — it returns `{}` on any failure, and D/E
    then fall back to derived-or-blank. That is the branch pinned here: an unavailable risk
    analysis must leave the two Part D cells explicitly outstanding rather than blank.
    """
    from da_silos.psa import risk

    monkeypatch.setattr(risk, "form_fields", lambda *a, **k: {})

    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    part_d = open_output(psa).tables[2]
    assert part_d.rows[4].cells[1].text.strip() == "[PENDING: similarity risk]"
    assert part_d.rows[6].cells[0].text.strip() == "[PENDING: D.2 comments]"


def test_an_unavailable_risk_analysis_still_leaves_a_part_e_block_to_sign(psa, con,
                                                                         monkeypatch):
    """The other half of the same failure. With no per-site analysis `_render_part_e` writes its
    fallback block, so `verify.py`'s "at least one Part E" rule is satisfied and the document
    says plainly that nothing was found — rather than omitting the part."""
    from da_silos.psa import risk

    monkeypatch.setattr(risk, "form_fields", lambda *a, **k: {})

    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    doc = open_output(psa)
    assert len(doc.tables) == 4
    assert "no mix-up risk" in doc.tables[3].rows[7].cells[0].text.lower()


def test_nothing_is_deferred_when_the_risk_engine_answers(psa, con):
    """The counterpart, and the reason the branch above needed forcing: against a populated
    catalogue every mapped field resolves, so a generated report carries NO `[PENDING` tag at
    all. A tag appearing here would mean a field regressed from answered to outstanding."""
    psa["template"].populate(SUBJECT, product_id=product_id_for(con, SUBJECT))

    assert "[PENDING" not in docx_text(output_path(psa))


def test_the_pending_tags_name_their_source_document(psa):
    """`[PENDING: PDD - ...]` vs `[PENDING: per-vial strength]`: the reader has to know whether
    to chase a document or a person."""
    assert "PDD" in psa["template"].PDD_ADDR
    assert "PENDING" in psa["template"].PDD_LOC
    assert "PENDING" in psa["template"].STRENGTH_REVIEW


def test_a_placeholder_valued_photo_field_is_never_embedded_as_an_image(psa, con):
    """The photo resolver's placeholder is `[PENDING: PDD - product image]`. Treated as a path
    it would be joined onto the repo root and stat'd, and a file that happened to exist there
    would be embedded as the product's photo."""
    pid = product_id_for(con, SUBJECT)
    con.execute("INSERT INTO doc_source (product_id, doc_role) VALUES (?, 'PDD')", (pid,))
    doc_id = con.execute("SELECT MAX(doc_id) FROM doc_source").fetchone()[0]
    con.execute("INSERT INTO doc_attribute (doc_id, attribute_key, value_text) VALUES (?, "
                "'product_photo', '[PENDING: PDD - product image]')", (doc_id,))
    con.commit()

    psa["template"].populate(SUBJECT, product_id=pid)

    doc = open_output(psa)
    assert len(doc.inline_shapes) == 0
    assert "[PENDING" in doc.tables[0].rows[10].cells[1].text


# ── the resolver, exercised directly ──────────────────────────────────────────


def test_a_derived_key_the_resolver_does_not_know_is_marked_for_attention(psa, con):
    """The `_derive` fall-through. A field map row naming a resolver nobody implemented must
    produce a visible TBD rather than a blank cell that reads as an answered question."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    assert psa["template"]._derive("no_such_resolver", product, {}, "[PENDING: x]") == "TBD"


@pytest.mark.parametrize("source", ["pdd", "tpp", "analysis", "execution"])
def test_an_unintegrated_source_resolves_to_its_placeholder_tag(psa, con, source):
    """The data-driven promise in `schema.sql`: adding PDD/TPP later flips a field's `source`
    from a placeholder tag to a real resolver with NO populator change. Until then the tag is
    the value."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()
    field = {"source": source, "resolver": None, "placeholder_tag": "[PENDING: TPP]"}

    assert psa["template"]._resolve(field, product, {}) == "[PENDING: TPP]"


def test_a_smartsheet_source_reads_the_named_product_column(psa, con):
    """`_resolve` indexes the row by the resolver name, so the field map's resolver has to be a
    real `product` column — a typo would raise at report time, not at build time."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()
    field = {"source": "smartsheet", "resolver": "vial_container_size", "placeholder_tag": None}

    assert psa["template"]._resolve(field, product, {}) == "2R"


def test_the_site_qa_function_names_the_first_packaging_site(psa, con):
    """The single-block fallback for E.3's Function column."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    assert psa["template"]._derive("site_qa_function", product, {}, None) == "Site QA - AP16"


def test_the_site_qa_function_is_marked_tbd_when_there_is_no_site(psa, psa_config, con):
    """A signature line addressed to nobody is still a signature line — it must say TBD rather
    than "Site QA - "."""
    con.execute("UPDATE product SET dp_pkging_site_raw='N/A' WHERE program_no=?", (SUBJECT,))
    con.commit()
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    assert psa["template"]._derive("site_qa_function", product, {}, None) == "Site QA - TBD"


def test_the_abbvie_site_cell_flags_a_missing_pdd_location(psa, con):
    """The site code alone does not say where the site IS. The trailing tag records that the
    location is still to come from the PDD rather than leaving the reader to assume."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    derived = psa["template"]._derive("abbvie_site", product, {}, None)

    assert derived.startswith("AP16")
    assert psa["template"].PDD_LOC in derived


def test_an_extracted_site_location_replaces_the_pending_tag(psa, con):
    """Once the PDD has been read the tag must go, or a completed field would still read as
    outstanding."""
    pid = product_id_for(con, SUBJECT)
    con.execute("UPDATE product SET site_location='North Chicago, IL' WHERE product_id=?", (pid,))
    con.commit()
    product = con.execute("SELECT * FROM product WHERE product_id=?", (pid,)).fetchone()

    derived = psa["template"]._derive("abbvie_site", product, {}, None)

    assert derived == "AP16 North Chicago, IL"
    assert "PENDING" not in derived


def test_the_pd_director_is_deferred_until_a_tpp_names_one(psa, con):
    """A fabricated approver name on a governance form is the worst kind of plausible."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    assert psa["template"]._derive("pd_director", product, {}, "[PENDING: TPP]") \
        == "[PENDING: TPP]"
    assert psa["template"]._derive("pd_director", product, {"pd_director_name": "A. Person"},
                                  "[PENDING: TPP]") == "A. Person"


@pytest.mark.parametrize("key,attr", [("risk_d1", "d1_similarity_risk"),
                                      ("risk_e1", "e1_similarity_risk")])
def test_an_unrecognised_risk_decision_defers_rather_than_ticking_a_box(psa, con, key, attr):
    """`_risk_box` returns None for anything that is not yes/no, and the placeholder is what
    lands instead. A form must never show a ticked box that no assessment produced."""
    product = con.execute("SELECT * FROM product WHERE program_no=?", (SUBJECT,)).fetchone()

    resolved = psa["template"]._derive(key, product, {attr: "maybe"}, "[PENDING: risk]")

    assert resolved == "[PENDING: risk]"


# ── the whole report path, as the ARIA stage runs it ──────────────────────────


def test_the_report_path_produces_a_document_that_passes_its_own_verify(psa, catalogue):
    """The end-to-end contract. `generate_report` populates and then verifies; a report that
    fails its own self-check is still downloadable, but the run log would carry the failure."""
    result = generated_report()

    assert result["status"] == "ok", result.get("message")
    assert result["output_exists"] is True
    assert result["verify_ok"] is True, result["log"]


def test_the_report_path_embeds_the_photo_it_is_handed(psa, catalogue):
    """The ARIA `report` stage reads the photo back from the run's stored media and passes it
    through `generate_report`. One inline shape end to end is what proves the whole hand-off."""
    result = generated_report(photo=png_bytes())

    assert result["status"] == "ok", result.get("message")

    from docx import Document

    assert len(Document(result["output_path"]).inline_shapes) == 1


def test_the_report_path_names_the_presentation_it_used(psa, catalogue):
    """A program code can span several presentation rows. The report has to say which one it
    describes, or a reviewer cannot tell whether it is the right vial."""
    result = generated_report()

    assert result["presentation_used"] == SUBJECT_ROW


def test_the_report_reads_back_as_a_populated_assessment(psa, catalogue):
    """The single assertion that would catch a document saved before it was filled in."""
    result = generated_report()

    text = docx_text(result["output_path"])

    assert "Product 0" in text
    assert "AbbVie US, AP16" in text
    assert "☒" in text
